"""Separate field-migration jobs, transport, checkpoints and scoped exports."""
import asyncio
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
import csv
import hashlib
import json
from pathlib import Path
import shutil
import threading
from uuid import uuid4

import captionqa.caption_qa_agent as qa
from captionqa.caption_field_schema import VERSION, FIELDS, parse, normalize, recognizable, is_legacy_reference
from captionqa.caption_qa_id_list import detect_id_column
from captionqa.caption_qa_table_io import iter_table_records, table_columns, _schema


def atomic(path,value):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix(path.suffix+'.tmp');temp.write_text(json.dumps(value,ensure_ascii=False),encoding='utf-8');temp.replace(path)


def validate_options(options):
    for key in ('column','id_column','sheet'):
        if options.get(key) is not None and (not isinstance(options[key],str) or len(options[key])>200):raise ValueError('列名或工作表名称无效。')
    concurrency=options.get('concurrency',8)
    if type(concurrency) is not int or not 1<=concurrency<=10:raise ValueError('并行条数须为1–10。')
    mapping=options.get('task_mapping') or {}
    if not isinstance(mapping,dict):raise ValueError('任务映射必须为JSON对象。')
    for key,value in mapping.items():
        if not isinstance(key,str) or not key or not (isinstance(value,str) and value.strip() or
            isinstance(value,list) and value and all(isinstance(v,str) and v.strip() for v in value)):
            raise ValueError('每个旧任务类型必须对应非空新版标签或标签列表。')


def mark_legacy_reference(result):
    result['legacy_reference']=is_legacy_reference(result.get('original'))
    if result['legacy_reference']:
        result['tags']=list(dict.fromkeys(result.get('tags',[])+['旧任务类型为参考']))
    return result


def missing_fields(text):
    """Missing headings, not empty values; recognize old names before comparing."""
    parsed=parse(text)
    if not parsed['headers']:return []
    present={h['field'] for h in parsed['headers']}
    if '_task' in present:present.add('任务类型')
    return [name for name in FIELDS if name not in present]


def mark_source_fields(result):
    mark_legacy_reference(result)
    result['missing_fields']=missing_fields(result.get('original'))
    result['remaining_missing_fields']=missing_fields(result.get('caption'))
    if result['missing_fields']:
        result['tags']=list(dict.fromkeys(result.get('tags',[])+['原始缺少字段']))
    return result


def csv_safe(value):
    return "'"+value if value.startswith(('=','+','-','@')) else value


def write_missing_reports(output_dir,stem,results,sheet,suffix):
    selected=[(index,r) for index,r in sorted(results.items(),key=lambda item:int(item[0])) if r.get('missing_fields')]
    ids=list(dict.fromkeys(r.get('sample_id') for _,r in selected if r.get('sample_id')))
    id_path=output_dir/(stem+'_缺字段ID名单.txt')
    id_path.write_text(''.join(value+'\n' for value in ids),encoding='utf-8')
    report=output_dir/(stem+'_缺字段明细.csv')
    with report.open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.writer(stream);writer.writerow(['题目 ID','原始缺少字段','整理后仍缺少字段','原表行位置','工作表'])
        for index,r in selected:
            writer.writerow([csv_safe(v) for v in [r.get('sample_id',''),'、'.join(r['missing_fields']),
                '、'.join(r['remaining_missing_fields']),str(int(index)+1 if suffix in {'.csv','.xlsx'} else index),sheet or '']])
    return {'missing_field_ids':str(id_path),'missing_field_report':str(report)}


def detect_column(source,sheet=None,column=None):
    columns=table_columns(source,sheet_name=sheet)
    if column:
        if column not in columns:raise ValueError('指定的Caption列不存在。')
        return column
    ranked=[(qa._caption_column_priority(name),name) for name in columns]
    # Also accept the established two-row header convention.
    first=next(iter(iter_table_records(source,sheet_name=sheet)),(None,{},None))[1]
    ranked=[(max(score,qa._caption_column_priority(first.get(name))),name) if not qa._excluded_caption_column(name)
            else (0,name) for score,name in ranked]
    best=max((score for score,_ in ranked),default=0)
    choices=[name for score,name in ranked if score==best and score>0]
    if len(choices)!=1:raise ValueError('未找到唯一的Caption/人标列，请在字段规范设置中明确填写列名；不会检查机标或备注列。')
    return choices[0]


def select_id_column(source,sheet=None,column=None):
    if column:return detect_id_column(source,sheet_name=sheet,column=column)
    try:return detect_id_column(source,sheet_name=sheet)
    except ValueError:
        import re
        choices=[name for name in table_columns(source,sheet_name=sheet) if re.sub(r'\s+','',name).lower() in {'题目id','题目编号','题号'}]
        if len(choices)==1:return choices[0]
        raise ValueError('未找到唯一真实题目ID列，请在字段规范设置中填写ID列名。') from None


def select_table(source,column=None,sheet=None):
    if source.suffix!='.xlsx':return sheet,detect_column(source,sheet,column)
    from captionqa.caption_qa_table_io import _openpyxl
    book=_openpyxl().load_workbook(source,read_only=True,data_only=False)
    try:names=[sheet] if sheet else book.sheetnames
    finally:book.close()
    choices=[]
    for name in names:
        try:
            found=detect_column(source,name,column)
            if any(recognizable(row.get(found)) for _,row,error in iter_table_records(source,sheet_name=name) if not error):
                choices.append((name,found))
        except ValueError:
            if sheet:raise
    ordinary=[v for v in choices if not any(word in v[0] for word in ('示例','模板','说明','参考'))]
    choices=ordinary or choices
    if len(choices)!=1:raise ValueError('未找到唯一的数据工作表，请在字段规范设置中填写工作表名称。')
    return choices[0]


def write_copy(source,target,column,results,sheet=None):
    """Only change successful target cells; preserve IDs, headers and other data."""
    changes={int(index):r['caption'] for index,r in results.items() if r['status']=='converted' and r['caption']!=r['original']}
    suffix=source.suffix
    if not changes:shutil.copyfile(source,target);return
    if suffix=='.csv':
        with source.open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.reader(stream))
        col=_schema(source,sheet).keys.index(column)
        for index,value in changes.items():
            if col>=len(rows[index]):raise ValueError('目标单元格缺失，未扩展原表。')
            rows[index][col]=value
        with target.open('w',encoding='utf-8-sig',newline='') as stream:csv.writer(stream).writerows(rows)
    elif suffix=='.xlsx':
        from captionqa.caption_qa_review_io import patch_xlsx
        from openpyxl.utils import get_column_letter
        col=get_column_letter(_schema(source,sheet).keys.index(column)+1)
        patch_xlsx(source,target,{sheet:{col+str(index+1):value for index,value in changes.items()}})
    elif suffix=='.jsonl':
        with source.open(encoding='utf-8-sig',newline='') as stream:lines=stream.readlines()
        for index,value in changes.items():
            raw=json.loads(lines[index-1]);raw[column]=value
            ending='\r\n' if lines[index-1].endswith('\r\n') else '\n' if lines[index-1].endswith('\n') else ''
            lines[index-1]=json.dumps(raw,ensure_ascii=False)+ending
        with target.open('w',encoding='utf-8',newline='') as stream:stream.writelines(lines)
    elif suffix=='.json':
        value=json.loads(source.read_text(encoding='utf-8-sig'));items=value if isinstance(value,list) else [value]
        for index,caption in changes.items():
            if isinstance(items[index-1],str):items[index-1]=caption
            else:items[index-1][column]=caption
        target.write_text(json.dumps(items if isinstance(value,list) else items[0],ensure_ascii=False,indent=2),encoding='utf-8')
    else:raise ValueError('文件格式不支持。')


def export(source,output_dir,name,column,sheet,results,source_hash):
    if qa._source_sha256(source)!=source_hash:raise ValueError('上传副本已变化，拒绝回填。')
    for result in results.values():mark_source_fields(result)
    output_dir.mkdir(parents=True,exist_ok=True);stem=Path(name).stem[:45]+'_字段规范_'+uuid4().hex[:8]
    target=output_dir/(stem+source.suffix);temp=target.with_name('.'+target.name)
    try:
        write_copy(source,temp,column,results,sheet)
        if qa._source_sha256(source)!=source_hash:raise ValueError('导出期间原文件发生变化。')
        temp.replace(target)
    finally:temp.unlink(missing_ok=True)
    report=output_dir/(stem+'_ID报告.csv');labels={'converted':'字段已调整','already_standard':'字段已规范','needs_confirmation':'需确认','failed':'失败','skipped':'跳过'}
    with report.open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.writer(stream);writer.writerow(['题目 ID','修复标签','结果','说明','原表行位置','工作表','原始缺少字段','整理后仍缺少字段'])
        for index,result in sorted(results.items(),key=lambda item:int(item[0])):
            row=[result.get('sample_id',''),'｜'.join('【'+v+'】' for v in result.get('tags',[])),labels[result['status']],
                 '；'.join(result.get('notes',[])),str(int(index)+1 if source.suffix in {'.csv','.xlsx'} else index),sheet or '',
                 '、'.join(result['missing_fields']),'、'.join(result['remaining_missing_fields'])]
            writer.writerow([csv_safe(value) for value in row])
    paths={'corrected':str(target),'report':str(report)}
    for key,statuses in [('fixed_ids',{'converted'}),('review_ids',{'needs_confirmation'}),('failed_ids',{'failed'}),('legacy_reference_ids',None)]:
        path=output_dir/(stem+'_'+{'fixed_ids':'已修复ID','review_ids':'需确认ID','failed_ids':'失败ID','legacy_reference_ids':'旧任务类型参考ID名单'}[key]+'.txt')
        ids=list(dict.fromkeys(r.get('sample_id','') for r in results.values()
             if (r.get('legacy_reference') if statuses is None else r['status'] in statuses) and r.get('sample_id')))
        path.write_text(''.join(value+'\n' for value in ids),encoding='utf-8');paths[key]=str(path)
    audit=output_dir/(stem+'_修改记录.jsonl')
    with audit.open('w',encoding='utf-8') as stream:
        for index,result in sorted(results.items(),key=lambda item:int(item[0])):
            if result['status']=='converted':stream.write(json.dumps({'row_index':int(index),**result},ensure_ascii=False)+'\n')
    paths['audit']=str(audit);paths.update(write_missing_reports(output_dir,stem,results,sheet,source.suffix));return paths


class LazyClient:
    def __init__(self):self.client=None;self.chat=self;self.completions=self
    async def create(self,**options):
        if self.client is None:
            from openai import AsyncOpenAI
            from captionqa.run_mimo import read_saved_key, BASE_URL
            self.client=AsyncOpenAI(api_key=read_saved_key(),base_url=BASE_URL,timeout=180,max_retries=1)
        return await self.client.chat.completions.create(**options)
    async def close(self):
        if self.client:await self.client.close()


class FieldJobs:
    def __init__(self,work_dir,output_dir,client_factory=LazyClient):
        self.root=Path(work_dir);self.output_dir=Path(output_dir).resolve();self.client_factory=client_factory
        self.jobs={};self.lock=threading.RLock();self.executor=ThreadPoolExecutor(max_workers=1,thread_name_prefix='caption-fields')
        path=self.root/'session.json'
        if path.exists():
            try:
                for job in json.loads(path.read_text())['jobs']:
                    if not re_id(job.get('id','')) or not (self.root/job['id']/('input'+Path(job['name']).suffix.lower())).is_file():continue
                    if job['status'] in {'queued','running'}:job.update(status='failed',error='上次服务已停止，可重试未完成项。')
                    self.jobs[job['id']]=job
            except (OSError,ValueError,KeyError,TypeError):pass
        for job in self.jobs.values():self.restore_missing_reports(job)

    def restore_missing_reports(self,job):
        """Supplement finished histories from source-bound cache without model calls."""
        if job['status'] not in {'done','partial'}:return
        source=self.root/job['id']/('input'+Path(job['name']).suffix.lower())
        checkpoint=source.parent/'checkpoint.json'
        try:
            saved=json.loads(checkpoint.read_text())
            fingerprint=hashlib.sha256(json.dumps([job.get('schema_version') or 1,qa._source_sha256(source),
                job.get('selected_sheet'),job.get('selected_column'),job.get('id_column'),job.get('task_mapping',{})],sort_keys=True).encode()).hexdigest()
            if saved.get('fingerprint')!=fingerprint:return
            results=saved['results']
            for result in results.values():mark_source_fields(result)
            job['missing_field_count']=sum(bool(r['missing_fields']) for r in results.values())
            corrected=Path(job.get('outputs',{}).get('corrected',''))
            if corrected.parent.resolve()!=self.output_dir or not corrected.is_file():return
            job['outputs'].update(write_missing_reports(self.output_dir,corrected.stem,results,job.get('selected_sheet'),source.suffix))
        except (OSError,ValueError,KeyError,TypeError):return

    def save(self):atomic(self.root/'session.json',{'jobs':list(self.jobs.values())})
    def snapshots(self):
        with self.lock:return [{**job,'outdated_schema':job.get('schema_version')!=VERSION} for job in reversed(list(self.jobs.values()))]
    def create(self,name,stream,length,*,column=None,id_column=None,sheet=None,concurrency=8,task_mapping=None):
        suffix=Path(name).suffix.lower()
        if Path(name).name!=name or suffix not in {'.csv','.xlsx','.jsonl','.json'}:raise ValueError('请上传CSV、XLSX、JSONL或JSON文件。')
        if not 0<length<=100*1024*1024:raise ValueError('文件大小无效。')
        validate_options({'column':column,'id_column':id_column,'sheet':sheet,'concurrency':concurrency,'task_mapping':task_mapping})
        identifier=uuid4().hex;folder=self.root/identifier;folder.mkdir(parents=True)
        source=folder/('input'+suffix)
        try:
            with source.open('xb') as target:
                remaining=length
                while remaining:
                    chunk=stream.read(min(remaining,1024*1024))
                    if not chunk:raise ValueError('文件上传不完整。')
                    target.write(chunk);remaining-=len(chunk)
        except Exception:source.unlink(missing_ok=True);raise
        job={'id':identifier,'name':name,'status':'queued','column':column,'id_column':id_column,'sheet':sheet,
             'concurrency':concurrency,'task_mapping':task_mapping or {},'total':0,'processed':0,'counts':{},'error':'','outputs':{},'model_requests':0,'schema_version':VERSION}
        with self.lock:self.jobs[identifier]=job;self.save();self.executor.submit(self.process,identifier)
        return dict(job)

    def retry(self,identifier,**options):
        with self.lock:
            job=self.jobs[identifier]
            if job['status'] in {'queued','running'}:raise ValueError('字段转换正在进行。')
            if set(options)-{'column','id_column','sheet','task_mapping','concurrency'}:raise ValueError('重试参数无效。')
            validate_options({**job,**options})
            job.update(options);job.update(status='queued',error='',outputs={});self.save();self.executor.submit(self.process,identifier)
            return dict(job)

    def process(self,identifier):
        job=self.jobs[identifier];folder=self.root/identifier;source=folder/('input'+Path(job['name']).suffix.lower())
        try:
            with self.lock:job.update(status='running',error='',outputs={},schema_version=VERSION);self.save()
            sheet,column=select_table(source,job['column'],job['sheet'])
            id_column=select_id_column(source,sheet,job['id_column'])
            if source.suffix in {'.json','.jsonl'}:
                from captionqa.caption_qa_review_io import _unique_json
                if source.suffix=='.json':json.loads(source.read_text(encoding='utf-8-sig'),object_pairs_hook=_unique_json)
                else:
                    with source.open(encoding='utf-8-sig') as stream:
                        for line in stream:
                            if line.strip():json.loads(line,object_pairs_hook=_unique_json)
            rows=list(iter_table_records(source,sheet_name=sheet));source_hash=qa._source_sha256(source)
            fingerprint=hashlib.sha256(json.dumps([VERSION,source_hash,sheet,column,id_column,job['task_mapping']],sort_keys=True).encode()).hexdigest()
            checkpoint=folder/'checkpoint.json';saved=json.loads(checkpoint.read_text()) if checkpoint.exists() else {}
            results=saved.get('results',{}) if saved.get('fingerprint')==fingerprint else {}
            for index in list(results):
                if results[index]['status']=='failed':del results[index]
                else:mark_source_fields(results[index])
            duplicates=Counter(str(row.get(id_column)) for _,row,error in rows if not error and row.get(id_column) not in (None,''))
            with self.lock:job.update(total=len(rows),selected_column=column,selected_sheet=sheet,id_column=id_column,
                                      processed=len(results),counts=dict(Counter(r['status'] for r in results.values())),
                                      legacy_reference_count=sum(r.get('legacy_reference',False) for r in results.values()),
                                      missing_field_count=sum(bool(r.get('missing_fields')) for r in results.values()))
            async def run():
                from captionqa.run_mimo import MODEL
                client=self.client_factory();semaphore=asyncio.Semaphore(job['concurrency'])
                async def one(index,row,error):
                    key=str(index)
                    if key in results:return
                    async with semaphore:
                        text=row.get(column);value=row.get(id_column)
                        sample_id=str(value) if type(value) in (str,int) and str(value).strip() else ''
                        result={'original':text,'caption':text,'status':'needs_confirmation','tags':[],'notes':[],'model_requests':0}
                        if error:result.update(status='failed',notes=[error])
                        elif not isinstance(text,str) or not text.strip() or text.strip().lower() in {'caption','人标'}:result.update(status='skipped',notes=['空Caption或二级表头，未调用API。'])
                        elif not sample_id:result.update(notes=['原表缺少真实题目ID，保留原文。'])
                        elif duplicates[sample_id]>1:result.update(notes=['原表题目ID重复，保留原文以便确认。'])
                        elif source.suffix=='.xlsx' and text.startswith('='):result.update(notes=['Caption单元格是公式，保留原公式。'])
                        else:result=await normalize(text,client=client,model=MODEL,task_mapping=job['task_mapping'])
                        results[key]=mark_source_fields({**result,'sample_id':sample_id})
                        with self.lock:
                            job.update(processed=len(results),counts=dict(Counter(r['status'] for r in results.values())),model_requests=sum(r['model_requests'] for r in results.values()),
                                       legacy_reference_count=sum(r.get('legacy_reference',False) for r in results.values()),
                                       missing_field_count=sum(bool(r.get('missing_fields')) for r in results.values()))
                            atomic(checkpoint,{'fingerprint':fingerprint,'results':results});self.save()
                try:await asyncio.gather(*(one(*row) for row in rows))
                finally:await client.close()
            asyncio.run(run())
            outputs=export(source,self.output_dir,job['name'],column,sheet,results,source_hash)
            with self.lock:
                job.update(outputs=outputs,status='partial' if any(r['status'] in {'failed','needs_confirmation'} for r in results.values()) else 'done')
        except Exception as exc:
            with self.lock:job.update(status='failed',error=str(exc) if isinstance(exc,ValueError) else '字段转换失败（'+type(exc).__name__+'），上传原文保留，可重试。')
        finally:
            with self.lock:self.save()

    def output(self,identifier,kind):
        with self.lock:path=Path(self.jobs[identifier]['outputs'].get(kind,''))
        if path.resolve().parent!=self.output_dir or not path.is_file():raise ValueError('结果文件不存在。')
        return path
    def close(self):self.executor.shutdown(wait=True)


def re_id(value):return isinstance(value,str) and len(value)==32 and all(c in '0123456789abcdef' for c in value)
