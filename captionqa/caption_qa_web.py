#!/usr/bin/env python3
"""Local drag-and-drop UI. Run: .venv/bin/python agent.py serve --open"""
from __future__ import annotations

import argparse
import asyncio
import csv
import hmac
import json
import re
import secrets
import subprocess
import sys
import threading
import time
import uuid
import webbrowser
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlsplit
from zoneinfo import ZoneInfo

import captionqa.caption_qa_agent as qa
from captionqa.caption_qa_summary import aggregate_check_states, row_check_state
from captionqa.caption_qa_people import PeopleTracker, people_report_path
from captionqa.caption_qa_repair_table import export_repair_table, repair_report_path
from captionqa.run_mimo import BASE_URL, MODEL, CredentialError, read_saved_key

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = Path.home() / "Downloads"
SUPPORTED = {".xlsx", ".csv", ".jsonl", ".json"}
MAX_UPLOAD_BYTES = 100 * 1024 * 1024
DEFAULT_CONCURRENCY = 8
REQUEST_TIMEOUT = 60
CAPTION_TIMEOUT = 180
MAX_RETRIES = 2


@dataclass
class Job:
    id: str
    name: str
    source: Path
    output: Path
    cache: Path
    dry_run: bool
    column: str | None
    sheet: str | None
    concurrency: int
    status: str = "queued"
    total: int = 0
    error: str = ""
    offset: int = 0
    rows: dict = field(default_factory=dict)
    check_rows: dict = field(default_factory=dict)
    people_tracker: PeopleTracker = field(default_factory=PeopleTracker)
    sheet_choices: list = field(default_factory=list)
    selected_column: str | None = None
    caption_location: str = ""
    stale_results: bool = False
    id_column: str | None = None
    id_summary: dict = field(default_factory=dict)
    policy_version: int = qa.AUTO_FIX_POLICY_VERSION
    outdated_policy: bool = False
    repair_exporting: bool = False
    workflow_version: str = "legacy"
    outdated_instruction_rows: set = field(default_factory=set)
    recheck_of: str | None = None
    field_resolution_rows: dict = field(default_factory=dict)
    active_requests: dict = field(default_factory=dict)
    timed_out_rows: set = field(default_factory=set)

    def refresh(self):
        if self.stale_results or not self.cache.exists():
            return
        with self.cache.open("rb") as stream:
            stream.seek(self.offset)
            while True:
                start = stream.tell()
                line = stream.readline()
                if not line:
                    break
                if not line.endswith(b"\n"):
                    stream.seek(start)
                    break
                row = json.loads(line)
                if row.get('qa_checkpoint',{}).get('request_stats',{}).get('timed_out'):
                    self.timed_out_rows.add(row['row_index'])
                else:
                    self.timed_out_rows.discard(row['row_index'])
                caption_column=row.get('qa_checkpoint',{}).get('caption_field') or self.selected_column or self.column
                outdated=qa.instruction_parser_outdated(row.get(caption_column),
                    row.get('qa_report',{}).get('caption_parser_version',1),
                    semantic=row.get('qa_report',{}).get('semantic_qa_performed',False))
                if outdated:self.outdated_instruction_rows.add(row['row_index'])
                else:self.outdated_instruction_rows.discard(row['row_index'])
                notes=row.get('qa_report',{}).get('field_resolution_notes',[])
                if notes:
                    identifier=row.get(self.id_column) if self.id_column else next(
                        (row.get(k) for k in ('id','样本ID','题目ID','题目id','序列号') if row.get(k) is not None),None)
                    self.field_resolution_rows[row['row_index']]={'sample_id':identifier,'row_index':row['row_index'],'notes':notes}
                else:self.field_resolution_rows.pop(row['row_index'],None)
                from captionqa.caption_qa_id_list import has_content_problem
                status = row["qa_status"]
                if (row.get("qa_scope") == "input_error" or not row["qa_checkpoint"]["completed"]) and not has_content_problem(row):
                    status = "NOT_CHECKED"
                self.rows[row["row_index"]] = (status, row["qa_checkpoint"]["completed"])
                self.check_rows[row["row_index"]] = row_check_state(row)
                self.people_tracker.add(row)
                self.offset = stream.tell()

    def snapshot(self):
        self.refresh()
        counts = Counter(value[0] for value in self.rows.values())
        pending = sum(not value[1] for value in self.rows.values())
        from captionqa.caption_qa_review_store import read_review_summary
        review = read_review_summary(self.cache)
        saved_report = Path(self.id_summary.get('repair_report_path') or repair_report_path(self.output))
        if saved_report.parent != self.output.parent:
            saved_report = repair_report_path(self.output)
        return {"id": self.id, "name": self.name, "status": self.status, "total": self.total,
                "processed": len(self.rows), "counts": dict(counts), "pending_retry": pending,
                "request_timeout_seconds": REQUEST_TIMEOUT, "caption_timeout_seconds": CAPTION_TIMEOUT,
                "active_requests": [dict({k:v for k,v in item.items() if k != 'started'},
                    elapsed_seconds=int(max(0,time.monotonic()-item['started'])))
                    for item in self.active_requests.values()] if self.status == 'running' else [],
                "timed_out_count": len(self.timed_out_rows),
                "dry_run": self.dry_run, "error": self.error,
                "output_name": self.output.name, "output_path": str(self.output),
                "output_ready": self.output.is_file() and self.status in {"done", "partial"},
                "cache_path": str(self.cache), "sheet": self.sheet, "column": self.column,
                "concurrency": self.concurrency, "sheet_choices": self.sheet_choices,
                "selected_column": self.selected_column, "caption_location": self.caption_location,
                "stale_results": self.stale_results, "caption_scope_version": qa.CAPTION_SCOPE_VERSION,
                "id_column": self.id_column, "id_summary": self.id_summary,
                "check_report": aggregate_check_states(self.check_rows.values()),
                "people_report": self.people_tracker.snapshot(),
                "people_report_path": str(people_report_path(self.output)) if people_report_path(self.output).is_file() else "",
                "explanation_path": self.id_summary.get("explanation_path", ""),
                "repair_report_path": str(saved_report) if (
                    not self.repair_exporting and not self.id_summary.get("repair_report_error")
                    and saved_report.is_file()) else "",
                "repair_report_error": self.id_summary.get("repair_report_error", ""),
                "repair_exporting": self.repair_exporting,
                "workflow_version": self.workflow_version, "review_summary": review,
                "outdated_instruction_parser":bool(self.outdated_instruction_rows),
                "recheck_of":self.recheck_of,
                "field_resolution_notes":list(self.field_resolution_rows.values()),
                "review_csv_path": self.id_summary.get("review_csv_path", ""),
                "review_output": self.id_summary.get("review_output", {}),
                "auto_fix_policy_version": self.policy_version, "outdated_policy": self.outdated_policy,
                "policy_note": "旧规则历史结果；底线结论请按底线规则重新检测。新底线检查确定性格式、标签、内容符号配对和主体结构，不把模型疑似问题计为不合格。" if self.outdated_policy else ""}


class App:
    def __init__(self, output_dir=DEFAULT_OUTPUT_DIR, work_dir=None):
        self.output_dir = Path(output_dir).resolve()
        self.work_dir = Path(work_dir or ROOT / ".local" / "web_jobs").resolve()
        self.token = secrets.token_urlsafe(32)
        self.jobs: dict[str, Job] = {}
        self.lock = threading.RLock()
        self.executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="caption-qa")
        from captionqa.caption_field_jobs import FieldJobs
        self.field_jobs=FieldJobs(self.work_dir/'field_normalization',self.output_dir)
        self.restore_history()

    def restore_history(self):
        path = self.work_dir / "session.json"
        if not path.exists():
            return
        try:
            records = json.loads(path.read_text(encoding="utf-8"))["jobs"]
            for row in records:
                identifier, name = row["id"], row["name"]
                if len(identifier) != 32 or any(c not in "0123456789abcdef" for c in identifier):
                    continue
                source = self.work_dir / identifier / ("input" + Path(name).suffix.lower())
                output = Path(row["output_path"]).resolve()
                if not source.is_file() or output.parent != self.output_dir:
                    continue
                job = Job(identifier, name, source, output, output.with_suffix(".jsonl"), row["dry_run"],
                          row.get("column"), row.get("sheet"), row.get("concurrency", DEFAULT_CONCURRENCY),
                          status=row["status"], total=row["total"], error=row.get("error", ""))
                job.sheet_choices = row.get("sheet_choices", [])
                job.workflow_version = row.get("workflow_version", "legacy")
                job.recheck_of = row.get("recheck_of")
                saved_cache = Path(row.get("cache_path", str(job.cache))).resolve()
                if saved_cache.parent not in {self.output_dir, source.parent}:
                    continue
                job.cache = saved_cache
                job.id_column, job.id_summary = row.get("id_column"), row.get("id_summary", {})
                job.selected_column = row.get("selected_column")
                job.caption_location = row.get("caption_location", "")
                job.stale_results = row.get("stale_results", False)
                job.stale_results |= job.output.suffix != ".txt"
                if job.cache.exists() and job.cache.stat().st_size:
                    with job.cache.open(encoding="utf-8") as stream:
                        first = next((json.loads(line) for line in stream if line.strip()), {})
                    job.stale_results |= first.get("qa_checkpoint", {}).get("human_caption_filter") != qa.CAPTION_SCOPE_VERSION
                    job.policy_version = first.get("qa_checkpoint", {}).get("auto_fix_policy")
                    job.outdated_policy = job.policy_version != qa.AUTO_FIX_POLICY_VERSION
                if job.status in {"queued", "preparing", "running"}:
                    job.status, job.error = "failed", "上次服务已停止；可点击重试继续未完成记录。"
                if job.stale_results:
                    job.status, job.total = "failed", 0
                    job.selected_column, job.caption_location = None, ""
                    job.id_summary = {}
                    job.error = "旧结果已停用，避免跨列或标题格式误判混入名单。点击重试将仅检测 Caption，生成新的问题 ID 名单；旧文件保留。"
                if job.status == "failed" and not job.sheet and not job.cache.exists() and source.suffix == ".xlsx":
                    try:
                        self.choose_sheet(job)
                        job.error = f"已重新读取文件并识别工作表“{job.sheet}”，点击重试可继续。未调用 API。"
                    except ValueError as exc:
                        job.error = str(exc)
                self.jobs[identifier] = job
        except (OSError, ValueError, KeyError, TypeError):
            print("部分界面历史无法恢复，已保留原有结果和上传副本。", flush=True)

    def save_history(self):
        try:
            self.work_dir.mkdir(parents=True, exist_ok=True)
            path = self.work_dir / "session.json"
            temporary = path.with_suffix(".tmp")
            temporary.write_text(json.dumps({"jobs": [job.snapshot() for job in self.jobs.values()]},
                                             ensure_ascii=False), encoding="utf-8")
            temporary.replace(path)
        except OSError:
            pass  # Output/checkpoint files remain the source of truth.

    def choose_sheet(self, job):
        if job.source.suffix != ".xlsx" or job.sheet:
            return
        from captionqa.caption_qa_table_io import _openpyxl
        workbook = _openpyxl().load_workbook(job.source, read_only=True, data_only=False)
        choices = []
        try:
            names = workbook.sheetnames
        finally:
            workbook.close()
        ambiguous = []
        for name in names:
            try:
                column = job.column or qa.detect_caption_column(job.source, sheet_name=name)
            except ValueError as exc:
                if "多个 Caption" in str(exc):
                    ambiguous.append(name)
                continue
            count = sum(qa.is_human_caption(row.get(column))
                        for _, row, _ in qa._indexed_input_records(job.source, sheet_name=name))
            if count:
                choices.append({"name": name, "caption_rows": count, "column": column})
        with self.lock:
            job.sheet_choices = choices
            ordinary = [item for item in choices if not any(word in item["name"] for word in ("参考", "示例", "模板", "说明"))]
            candidates = ordinary or choices
            if len(candidates) == 1:
                job.sheet = candidates[0]["name"]
            elif candidates:
                raise ValueError("发现多个含 Caption 的工作表，请在此记录中选择工作表后重试。")
            else:
                if ambiguous:
                    raise ValueError(f"工作表 {ambiguous!r} 含多个 Caption 列，请在更多设置中明确指定工作表和列名。未调用 API。")
                raise ValueError("已检查全部工作表，未找到符合人标格式的 Caption。请检查列名及正文结构。未调用 API。")

    def config(self):
        try:
            read_saved_key()
            ready = True
        except (CredentialError, OSError, UnicodeError, subprocess.SubprocessError):
            ready = False
        return {"output_dir": str(self.output_dir), "key_ready": ready, "model": MODEL,
                "max_upload_mb": MAX_UPLOAD_BYTES // (1024 * 1024)}

    def create_job(self, name, stream, length, *, dry_run=False, column=None, sheet=None, concurrency=DEFAULT_CONCURRENCY, id_column=None, review_first=True, auto_fix=True):
        if not name or name != Path(name).name or any(c in name for c in ('\\', '\x00', '\n', '\r')):
            raise ValueError("文件名无效，请重新选择文件。")
        suffix = Path(name).suffix.lower()
        if suffix not in SUPPORTED:
            raise ValueError("支持 XLSX、CSV、JSONL、JSON 文件。")
        if not 0 < length <= MAX_UPLOAD_BYTES:
            raise ValueError("文件不能为空，且单个文件不能超过 100 MB。")
        if type(concurrency) is not int or not 1 <= concurrency <= 10:
            raise ValueError("并行条数须在 1–10 之间。")
        if type(dry_run) is not bool:
            raise ValueError("检查模式无效。")
        if type(review_first) is not bool or type(auto_fix) is not bool:
            raise ValueError("审核流程参数无效。")
        if any(value is not None and (not isinstance(value, str) or len(value) > 200) for value in (column, sheet, id_column)):
            raise ValueError("列名或工作表名无效。")
        identifier = uuid.uuid4().hex
        folder = self.work_dir / identifier
        folder.mkdir(parents=True)
        source = folder / ("input" + suffix)
        try:
            with source.open("xb") as target:
                remaining = length
                while remaining:
                    chunk = stream.read(min(1024 * 1024, remaining))
                    if not chunk:
                        raise ValueError("文件上传未完成，请重新拖入。")
                    target.write(chunk)
                    remaining -= len(chunk)
        except Exception:
            source.unlink(missing_ok=True)
            raise
        self.output_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y%m%d_%H%M%S")
        stem = Path(name).stem[:45]
        mode = "底线问题ID名单" if dry_run else "辅助筛查ID名单"
        output = self.output_dir / f"{stem}_{mode}_{stamp}_{identifier[:8]}.txt"
        job = Job(identifier, name, source, output, folder / "checkpoint.jsonl", dry_run,
                  column or None, sheet or None, concurrency)
        job.id_column = id_column or None
        if review_first:
            from captionqa.caption_qa_review import WORKFLOW, AUTO_WORKFLOW
            job.workflow_version = AUTO_WORKFLOW if auto_fix else WORKFLOW
        with self.lock:
            self.jobs[identifier] = job
            self.save_history()
            self.executor.submit(self.process, job)
            return job.snapshot()

    def process(self, job):
        try:
            with self.lock:
                job.status, job.error = "preparing", ""
            from captionqa.caption_qa_split_input import prepare_split_input
            split = prepare_split_input(job.source, sheet_name=job.sheet, id_column=job.id_column) if not job.column else None
            if not split:
                self.choose_sheet(job)
            input_source = split["path"] if split else job.source
            input_sheet = None if split else job.sheet
            selected_column = "caption" if split else (job.column or qa.detect_caption_column(input_source, sheet_name=input_sheet))
            from captionqa.caption_qa_table_io import table_columns
            columns = table_columns(input_source, sheet_name=input_sheet)
            if selected_column not in columns:
                raise ValueError("指定的 Caption 列不存在。未调用 API。")
            location = selected_column
            if split:
                location = f"{len(split['sheets'])} 个工作表的分列 Caption（基调v2、场景v2、主体v2、画面元素v2、总结v2、自然指令v2）"
                if not job.dry_run:
                    location += "；画面元素仅做本地格式检查并提供只读上下文，其余五项参与 AI 检查"
            elif job.source.suffix == ".xlsx":
                from openpyxl.utils import get_column_letter
                location = f"{job.sheet}!{get_column_letter(columns.index(selected_column) + 1)}列（Caption）"
            elif job.source.suffix == ".csv":
                location = f"CSV 第 {columns.index(selected_column) + 1} 列（Caption）"
            with self.lock:
                job.selected_column, job.caption_location = selected_column, location
            from captionqa.caption_qa_id_list import detect_id_column
            identifier_column = "id" if split else detect_id_column(input_source, sheet_name=input_sheet, column=job.id_column)
            total = sum(1 for _ in qa._indexed_input_records(input_source, sheet_name=input_sheet))
            if not total:
                raise ValueError("所选工作表没有数据行，请选择包含 Caption 的工作表。未调用 API。")
            key = None if job.dry_run else read_saved_key()
            with self.lock:
                job.total, job.status = total, "running"
            from captionqa.caption_qa_review import IdentificationAgent, SafeAutoFixAgent, WORKFLOW, AUTO_WORKFLOW, AUTO_WORKFLOWS
            agent_class = (SafeAutoFixAgent if job.workflow_version in AUTO_WORKFLOWS else
                           IdentificationAgent if job.workflow_version == WORKFLOW else qa.CaptionQAAgent)
            agent = agent_class(dry_run=job.dry_run, api_key=key,
                                      base_url=BASE_URL, model_name=MODEL, timeout=REQUEST_TIMEOUT,
                                      total_timeout=CAPTION_TIMEOUT, max_retries=MAX_RETRIES,
                                      **({'normalize_width':job.workflow_version==AUTO_WORKFLOW} if job.workflow_version in AUTO_WORKFLOWS else {}))
            def progress(event):
                with self.lock:
                    index=event['row_index']
                    if event['state']=='finished':
                        job.active_requests.pop(index,None)
                    else:
                        previous=job.active_requests.get(index,{})
                        job.active_requests[index]=dict(previous,**event,
                            started=previous.get('started',time.monotonic()-event['elapsed_seconds']))
            asyncio.run(qa.aprocess_batch(input_source, job.output, agent, column=selected_column,
                                         sheet_name=input_sheet, concurrency=job.concurrency,
                                         checkpoint_path=job.cache, id_column=identifier_column,
                                         on_progress=progress))
            with self.lock:
                job.refresh()
                job.status = "partial" if any(not value[1] for value in job.rows.values()) else "done"
                summary = job.cache.with_suffix(".summary.json")
                if summary.exists():
                    job.id_summary = json.loads(summary.read_text())
                    if not job.id_summary["complete"]:
                        job.status = "partial"
                        if job.id_summary["missing_problem_ids"]:
                            job.error = "部分问题行缺少样本 ID，名单不完整；请按检测说明补齐原表 ID 后重新添加文件。"
            if job.workflow_version in AUTO_WORKFLOWS:
                try:
                    from captionqa.caption_qa_review import approved_changes
                    with self.lock:
                        document = self._review_document(job)
                    if approved_changes(document):
                        self.apply_review(job.id, document['revision'], automatic=True)
                except (ValueError, OSError) as exc:
                    with self.lock:
                        job.id_summary['auto_correction_error'] = str(exc) or '自动修正副本保存失败，请重试生成。'
            try:
                self.export_repair_report(job.id)
            except ValueError:
                pass  # Export failure is separate from the completed detection.
        except (CredentialError, ValueError, csv.Error) as exc:
            with self.lock:
                job.status, job.error = "failed", str(exc)
        except Exception as exc:
            with self.lock:
                job.status, job.error = "failed", f"处理未完成（{type(exc).__name__}），请检查文件或稍后重试。"
        finally:
            with self.lock:
                job.active_requests.clear()
                self.save_history()

    def recheck(self, identifier):
        """Run current rules on a fresh copy; keep the original job and reviews."""
        from captionqa.caption_qa_review_model import AUTO_WORKFLOWS
        with self.lock:
            job=self.jobs.get(identifier)
            if job is None:raise KeyError(identifier)
            if job.status not in {'done','partial','failed'} or job.repair_exporting:
                raise ValueError('本任务仍在处理，请完成后再重新检测。')
            pending=next((j for j in self.jobs.values() if j.recheck_of==identifier
                          and j.status in {'queued','preparing','running'}),None)
            if pending:return pending.snapshot()
            if not job.source.is_file():raise ValueError('上传原文件已不存在，请重新添加文件。')
            column=job.column or (job.selected_column if self._review_input(job)==job.source else None)
            with job.source.open('rb') as stream:
                result=self.create_job(job.name,stream,job.source.stat().st_size,
                    dry_run=job.dry_run,column=column,sheet=job.sheet,
                    concurrency=job.concurrency,id_column=job.id_column,
                    review_first=job.workflow_version!='legacy',auto_fix=job.workflow_version in AUTO_WORKFLOWS)
            fresh=self.jobs[result['id']];fresh.recheck_of=identifier
            self.save_history()
            return fresh.snapshot()

    def retry(self, identifier, sheet=None, concurrency=None):
        with self.lock:
            if concurrency is not None and (type(concurrency) is not int or not 1 <= concurrency <= 10):
                raise ValueError("并行条数须在 1–10 之间。")
            job = self.jobs.get(identifier)
            if job is None:
                raise KeyError(identifier)
            job.refresh()
            if job.repair_exporting:
                raise ValueError("返修表正在导出，请完成后再重试检测。")
            if job.status not in {"failed", "partial"} and not (job.status == "done" and (job.outdated_policy or job.outdated_instruction_rows)):
                raise ValueError("此任务无需重试，或仍在处理中。")
            previous_sheet = job.sheet
            if sheet is not None:
                if sheet not in {item["name"] for item in job.sheet_choices}:
                    raise ValueError("请选择列表中的工作表。")
                if sheet != job.sheet and not job.stale_results and not job.outdated_policy and job.cache.exists() and job.cache.stat().st_size:
                    raise ValueError("此任务已有续传记录；更换工作表请重新添加文件，避免混用结果。")
                job.sheet = sheet
            if job.stale_results or job.outdated_policy or job.outdated_instruction_rows:
                if job.outdated_policy or job.outdated_instruction_rows:
                    job.dry_run = True
                old_resolutions = job.cache.with_suffix(".id_resolutions.json")
                resolutions = None
                if previous_sheet == job.sheet and old_resolutions.exists():
                    document = json.loads(old_resolutions.read_text())
                    if document.get("source_sha256") == qa._source_sha256(job.source):
                        resolutions = document
                stamp = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y%m%d_%H%M%S")
                mode = "底线问题ID名单" if job.dry_run else "辅助筛查ID名单"
                nonce = uuid.uuid4().hex[:8]
                job.output = self.output_dir / f"{Path(job.name).stem[:45]}_{mode}_{stamp}_{nonce}.txt"
                job.cache = job.source.parent / f"checkpoint_{nonce}.jsonl"
                if resolutions is not None:
                    job.cache.with_suffix(".id_resolutions.json").write_text(json.dumps(resolutions, ensure_ascii=False))
                job.rows.clear()
                job.outdated_instruction_rows.clear()
                job.field_resolution_rows.clear()
                job.timed_out_rows.clear()
                job.check_rows.clear()
                job.people_tracker = PeopleTracker()
                job.offset, job.total, job.stale_results = 0, 0, False
                job.selected_column, job.caption_location = None, ""
                job.id_summary = {}
                job.policy_version, job.outdated_policy = qa.AUTO_FIX_POLICY_VERSION, False
            if concurrency is not None:
                job.concurrency = concurrency
            # Changed checkpoints require fresh decisions; retain old approvals
            # as history rather than silently applying them to new findings.
            for suffix in ('.review.json','.review_summary.json'):
                previous = job.cache.with_suffix(suffix)
                if previous.exists():
                    previous.rename(previous.with_name(previous.stem + '_previous_' + uuid.uuid4().hex[:6] + previous.suffix))
            job.status, job.error = "queued", ""
            self.save_history()
            self.executor.submit(self.process, job)
            return job.snapshot()

    def export_repair_report(self, identifier):
        with self.lock:
            job = self.jobs.get(identifier)
            if job is None:
                raise KeyError(identifier)
            job.refresh()
            if job.status not in {"done", "partial"} or not job.cache.is_file():
                raise ValueError("检测结果尚未就绪。")
            if job.stale_results or job.outdated_policy or job.outdated_instruction_rows:
                raise ValueError("这是旧规则结果，请先按底线规则重新检测后再导出返修表。")
            if job.repair_exporting:
                raise ValueError("返修表正在导出，请稍后查看。")
            job.repair_exporting = True
        try:
            from captionqa.caption_qa_id_list import detect_id_column, _write_atomic
            from captionqa.caption_qa_split_input import PREPARED_NAME
            prepared = job.source.parent / PREPARED_NAME
            column = "id" if prepared.is_file() and not job.column else detect_id_column(job.source, sheet_name=job.sheet, column=job.id_column)
            from captionqa.caption_qa_review import WORKFLOW, AUTO_WORKFLOWS
            if job.workflow_version in {WORKFLOW, *AUTO_WORKFLOWS}:
                from captionqa.caption_qa_review_report import export_review
                result = export_review(self._review_document(job), job.output)
            else:
                result = export_repair_table(job.cache, job.output, id_column=column, dry_run=job.dry_run,
                                             complete=job.id_summary.get("complete", False))
            with self.lock:
                job.id_summary.update(result)
                job.id_summary.pop("repair_report_error", None)
                _write_atomic(job.cache.with_suffix(".summary.json"), json.dumps(job.id_summary, ensure_ascii=False, indent=2))
        except (ValueError, OSError) as exc:
            message = str(exc) if isinstance(exc, ValueError) else "返修表无法保存，请检查输出目录权限后单独重试导出。"
            with self.lock:
                job.id_summary["repair_report_error"] = message
            raise ValueError(message) from exc
        finally:
            with self.lock:
                job.repair_exporting = False
                self.save_history()
        return job.snapshot()

    def _review_input(self, job):
        from captionqa.caption_qa_split_input import PREPARED_NAME
        prepared = job.source.parent / PREPARED_NAME
        return prepared if prepared.is_file() and not job.column else job.source

    def _review_store(self, job):
        from captionqa.caption_qa_review_store import ReviewStore
        source = self._review_input(job)
        return ReviewStore(job.source, source, job.cache, job.selected_column or job.column,
                           job.id_column, job.sheet if source == job.source else None)

    def _save_review(self, job, document):
        self._review_store(job).save(document)

    def _review_document(self, job):
        return self._review_store(job).load()

    def _review_payload(self, job, document):
        return self._review_store(job).payload(document)

    def _review_job(self, identifier):
        job = self.jobs.get(identifier)
        if not job:
            raise KeyError(identifier)
        job.refresh()
        if job.outdated_instruction_rows:
            raise ValueError('旧版未识别 subject 等字段别名或 subtask_type 中的提示正文，请重新检测后审核；原文件、旧报告和审核记录保留。')
        if job.status not in {'done', 'partial'} or not job.cache.is_file() or job.stale_results:
            raise ValueError('请先完成检测，再审核与回填。')
        if job.repair_exporting:
            raise ValueError('正在导出，请稍后再操作审核结果。')
        return job

    def review(self, identifier, options=None):
        from captionqa.caption_qa_review_state import decide
        with self.lock:
            job = self._review_job(identifier)
            document = self._review_document(job)
            if options is not None:
                if not isinstance(options, dict) or set(options) - {'revision','issue_id','decision','field_id','final_text','reset'}:
                    raise ValueError('审核参数无效。')
                if bool(options.get('issue_id')) == bool(options.get('field_id')):
                    raise ValueError('每次只审核一个错误项或一个字段。')
                if 'reset' in options and type(options['reset']) is not bool:
                    raise ValueError('撤销参数无效。')
                document = decide(document, **options)
                self._save_review(job, document)
                self.save_history()
            return self._review_payload(job, document)

    def import_review(self, identifier, stream, length, name, revision):
        from captionqa.caption_qa_review_report import import_review
        if Path(name).suffix.lower() not in {'.csv','.xlsx'} or not 0 < length <= MAX_UPLOAD_BYTES:
            raise ValueError('请回传 CSV 或 XLSX 审核表，文件须小于 100 MB。')
        with self.lock:
            job = self._review_job(identifier)
            document = self._review_document(job)
            import tempfile
            with tempfile.TemporaryDirectory(prefix='caption-review-import-') as directory:
                target = Path(directory) / ('review' + Path(name).suffix.lower())
                data = stream.read(length)
                if len(data) != length:
                    raise ValueError('审核表上传未完成。')
                target.write_bytes(data)
                document = import_review(document, target, revision=revision)
            self._save_review(job, document)
            return self.review(identifier)

    def propose_review(self,identifier,revision,issue_id):
        import copy
        from captionqa.caption_qa_candidates import candidate_operations
        from captionqa.caption_qa_suggestions import prepare_proposal
        with self.lock:
            job=self._review_job(identifier);document=self._review_document(job)
            if type(revision) is not int or revision!=document['revision']:
                raise ValueError('审核版本已变化，请刷新后再生成改稿。')
            issue=next((i for i in document['issues'] if i['issue_id']==issue_id),None)
            if not issue or issue['decision'] not in {'pending','defer'}:
                raise ValueError('只为待处理项目生成改稿，已有审核结论不改变。')
            field=next(f for f in document['fields'] if f['field_id']==issue['field_id'])
            if field['field_id'] in document['field_edits'] or not issue['sample_id']:
                raise ValueError('本字段已有人工稿或缺少题目 ID，不能生成改稿。')
            issue,field=copy.deepcopy(issue),copy.deepcopy(field)
        # Network work must not hold the application-wide review lock.
        proposal=prepare_proposal(issue,field,api_key=read_saved_key(),base_url=BASE_URL,model=MODEL)
        proposed=dict(issue,suggestion_candidate=proposal)
        operations=candidate_operations(proposed,field['original'])
        if not operations:
            raise ValueError('改稿未通过定位或修改范围校验，原文和审核结论未改变；可重试或人工填写。')
        with self.lock:
            job=self._review_job(identifier);document=self._review_document(job)
            if document['revision']!=revision:
                raise ValueError('生成期间审核版本已变化，本次改稿未保存，请刷新后重试。')
            target=next(i for i in document['issues'] if i['issue_id']==issue_id)
            target.update(suggestion_candidate=proposal,operations=operations,requires_confirmation=True,
                          candidate_policy=operations[0]['candidate_policy'])
            document['revision']+=1
            document['events'].append({'revision':document['revision'],'issue_id':issue_id,'decision':'candidate_generated','model':MODEL})
            self._save_review(job,document)
            return self._review_payload(job, document)

    def apply_review(self, identifier, revision, *, automatic=False, projection_fingerprint=None):
        from captionqa.caption_qa_review_io import export_approved
        with self.lock:
            job = self._review_job(identifier)
            document = self._review_document(job)
            if revision != document['revision'] or type(revision) is not int:
                raise ValueError('审核版本已变化，请刷新后导出最新审核后版本；已保存的审核记录保留。')
            plan = self._review_store(job).plan(document).require_valid()
            if projection_fingerprint is not None and projection_fingerprint != plan.fingerprint:
                raise ValueError('审核预览已变化，请刷新后导出最新审核后版本。')
            previous = job.id_summary.get('review_output', {})
            if (not automatic and previous.get('export_kind') == 'reviewed'
                    and previous.get('review_revision') == revision
                    and previous.get('plan_fingerprint') == plan.fingerprint
                    and not previous.get('export_warnings')):
                target = Path(previous.get('corrected_path', ''))
                paths = [Path(value) for key,value in previous.items() if key.endswith('_path') and value]
                if (target.is_file() and all(p.is_file() for p in paths)
                        and previous.get('corrected_sha256') == qa._source_sha256(target)):
                    return {**previous, 'reused':True}
            job.repair_exporting = True
        try:
            result = export_approved(document, job.source, self._review_input(job), job.cache,
                                     self.output_dir, job.name, automatic=automatic, plan=plan)
            with self.lock:
                job.id_summary['review_output'] = result
                job.id_summary.pop('auto_correction_error', None)
            return result
        finally:
            with self.lock:
                job.repair_exporting = False
                self.save_history()

    def export_review(self, identifier):
        from captionqa.caption_qa_review_report import export_review
        with self.lock:
            job = self._review_job(identifier)
            document = self._review_document(job)
            job.repair_exporting = True
        try:
            result = export_review(document, job.output)
            with self.lock:
                job.id_summary.update(result)
            return result
        finally:
            with self.lock:
                job.repair_exporting = False
                self.save_history()

    def close(self):
        self.executor.shutdown(wait=True)
        self.field_jobs.close()


class Handler(BaseHTTPRequestHandler):
    server_version = "CaptionQA/1.0"

    @property
    def app(self):
        return self.server.app

    def log_message(self, format, *args):
        # Avoid logging file names, user content, or request bodies.
        pass

    def reply(self, status, data, content_type="application/json; charset=utf-8", extra=None):
        if not isinstance(data, bytes):
            data = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Security-Policy", "default-src 'self'; style-src 'self'; script-src 'self'; connect-src 'self'; img-src 'self' data:; base-uri 'none'; frame-ancestors 'none'; form-action 'self'")
        for key, value in (extra or {}).items():
            self.send_header(key, value)
        self.end_headers()
        self.wfile.write(data)

    def allowed(self, *, api=False):
        port = self.server.server_address[1]
        hosts = {f"127.0.0.1:{port}", f"localhost:{port}"}
        host = self.headers.get("Host", "")
        origin = self.headers.get("Origin")
        if host not in hosts or (origin and origin != "http://" + host):
            self.reply(403, {"error": "仅允许从本机界面访问。"})
            return False
        if api and not hmac.compare_digest(self.headers.get("X-App-Token", "").encode("utf-8"), self.app.token.encode("ascii")):
            self.reply(403, {"error": "页面已失效，请刷新后重试。"})
            return False
        return True

    def do_GET(self):
        try:
            path = urlsplit(self.path).path
            if not self.allowed(api=path.startswith("/api/")):
                return
            if path == "/api/config":
                self.reply(200, self.app.config())
            elif path == '/api/field-jobs':
                self.reply(200,{'jobs':self.app.field_jobs.snapshots()})
            elif re.fullmatch(r'/api/field-jobs/[0-9a-f]{32}/download',path):
                from urllib.parse import quote
                kind=parse_qs(urlsplit(self.path).query).get('kind',['report'])[0]
                target=self.app.field_jobs.output(path.split('/')[3],kind)
                self.reply(200,target.read_bytes(),'application/octet-stream',{'Content-Disposition':"attachment; filename*=UTF-8''"+quote(target.name)})
            elif path == "/api/jobs":
                with self.app.lock:
                    self.reply(200, {"jobs": [job.snapshot() for job in reversed(list(self.app.jobs.values()))]})
            elif re.fullmatch(r'/api/jobs/[0-9a-f]{32}/review', path):
                self.reply(200, self.app.review(path.split('/')[3]))
            elif path in {"/", "/index.html", "/app.js", "/review.js", "/fields.js", "/style.css"}:
                name = "index.html" if path in {"/", "/index.html"} else path[1:]
                data = (ROOT / "web" / name).read_bytes()
                if name == "index.html":
                    data = data.replace(b"__APP_TOKEN__", self.app.token.encode("ascii"))
                mime = {"index.html": "text/html", "app.js": "text/javascript", "review.js": "text/javascript", "fields.js":"text/javascript", "style.css": "text/css"}[name]
                self.reply(200, data, mime + "; charset=utf-8")
            else:
                self.reply(404, {"error": "页面不存在。"})
        except (BrokenPipeError, ConnectionResetError):
            pass
        except (ValueError, TypeError) as exc:
            self.reply(400, {'error': str(exc)})
        except Exception:
            self.reply(500, {"error": "读取结果失败，请稍后刷新。"})

    def do_POST(self):
        try:
            if not self.allowed(api=True):
                return
            parts = urlsplit(self.path)
            if parts.path == '/api/field-jobs':
                query=parse_qs(parts.query);name=query.get('name',[''])[0]
                options=json.loads(unquote(self.headers.get('X-Job-Options','{}')))
                if not isinstance(options,dict) or set(options)-{'column','id_column','sheet','concurrency','task_mapping'}:raise ValueError('字段规范参数无效。')
                self.reply(202,self.app.field_jobs.create(name,self.rfile,int(self.headers.get('Content-Length','0')),**options))
            elif re.fullmatch(r'/api/field-jobs/[0-9a-f]{32}/(?:retry|reveal)',parts.path):
                length=int(self.headers.get('Content-Length','0'))
                if not 0<=length<=65536:raise ValueError('字段规范参数过长。')
                options=json.loads(self.rfile.read(length)) if length else {}
                if not isinstance(options,dict):raise ValueError('字段规范参数须为对象。')
                identifier=parts.path.split('/')[3]
                if parts.path.endswith('/retry'):self.reply(202,self.app.field_jobs.retry(identifier,**options))
                else:
                    if set(options)!={'kind'}:raise ValueError('文件定位参数无效。')
                    target=self.app.field_jobs.output(identifier,options['kind'])
                    subprocess.run(['/usr/bin/open','-R',str(target)],check=True,capture_output=True)
                    self.reply(200,{'ok':True})
            elif parts.path == "/api/jobs":
                query = parse_qs(parts.query)
                name = query.get("name", [""])[0]
                options = json.loads(unquote(self.headers.get("X-Job-Options", "{}")))
                if not isinstance(options, dict) or set(options) - {"dry_run", "column", "sheet", "concurrency", "id_column", "review_first", "auto_fix"}:
                    raise ValueError("质检参数无效。")
                length = int(self.headers.get("Content-Length", "0"))
                self.reply(202, self.app.create_job(name, self.rfile, length, **options))
            elif re.fullmatch(r'/api/jobs/[0-9a-f]{32}/review(?:/(?:apply|export|import|propose))?', parts.path):
                length = int(self.headers.get('Content-Length','0'))
                if not 0 <= length <= MAX_UPLOAD_BYTES:
                    raise ValueError('审核数据过长。')
                identifier = parts.path.split('/')[3]
                if parts.path.endswith('/import'):
                    query = parse_qs(parts.query)
                    revision = int(query.get('revision',['-1'])[0])
                    self.reply(200, self.app.import_review(identifier, self.rfile, length, query.get('name',[''])[0], revision))
                elif parts.path.endswith('/export'):
                    self.reply(200, self.app.export_review(identifier))
                else:
                    if length > 1024 * 1024:
                        raise ValueError('审核参数过长。')
                    options = json.loads(self.rfile.read(length)) if length else {}
                    if not isinstance(options,dict):
                        raise ValueError('审核参数必须是对象。')
                    if parts.path.endswith('/propose'):
                        if set(options)!={'revision','issue_id'}:
                            raise ValueError('生成改稿参数无效。')
                        self.reply(200,self.app.propose_review(identifier,options['revision'],options['issue_id']))
                    elif parts.path.endswith('/apply'):
                        if 'revision' not in options or set(options) - {'revision','projection_fingerprint'}:
                            raise ValueError('生成修正副本参数无效。')
                        self.reply(200,self.app.apply_review(identifier,options['revision'],
                            projection_fingerprint=options.get('projection_fingerprint')))
                    else:
                        self.reply(200,self.app.review(identifier,options))
            elif re.fullmatch(r'/api/jobs/[0-9a-f]{32}/recheck',parts.path):
                length=int(self.headers.get('Content-Length','0'))
                if not 0<=length<=4096:raise ValueError('重新检测参数过长。')
                options=json.loads(self.rfile.read(length)) if length else {}
                if not isinstance(options,dict) or options:raise ValueError('重新检测沿用本条记录设置，无需传入参数。')
                self.reply(202,self.app.recheck(parts.path.split('/')[3]))
            elif parts.path.startswith("/api/jobs/") and parts.path.endswith("/retry"):
                length = int(self.headers.get("Content-Length", "0"))
                if not 0 <= length <= 4096:
                    raise ValueError("重试参数过长。")
                options = json.loads(self.rfile.read(length)) if length else {}
                if not isinstance(options, dict) or set(options) - {"sheet", "concurrency"}:
                    raise ValueError("重试参数无效。")
                self.reply(202, self.app.retry(parts.path.split("/")[3], **options))
            elif parts.path.startswith("/api/jobs/") and parts.path.endswith("/reveal"):
                with self.app.lock:
                    job = self.app.jobs.get(parts.path.split("/")[3])
                    if job is None or not job.snapshot()["output_ready"]:
                        raise ValueError("结果文件尚未就绪。")
                    if sys.platform != "darwin":
                        raise ValueError("请直接打开页面显示的输出目录。")
                    kind = parse_qs(parts.query).get("file", ["ids"])[0]
                    if kind not in {"ids", "explanation", "people", "repair", "review_csv", "corrected", "changes", "recheck", "highlighted", "highlight_ids"}:
                        raise ValueError("文件类型无效。")
                    target = (job.output if kind == "ids" else Path(job.snapshot()['repair_report_path']) if kind == "repair" else people_report_path(job.output) if kind == "people"
                              else job.output.with_name(job.output.stem + "_检测说明.txt"))
                    if kind == 'review_csv':
                        target = Path(job.id_summary.get('review_csv_path',''))
                    if kind in {'corrected','changes','recheck','highlighted','highlight_ids'}:
                        key = {'corrected':'corrected_path','changes':'change_log_path','recheck':'recheck_path','highlighted':'highlighted_path','highlight_ids':'highlight_ids_path'}[kind]
                        target = Path(job.id_summary.get('review_output',{}).get(key,''))
                    if target.resolve().parent != self.app.output_dir:
                        raise ValueError('结果文件位置无效。')
                    if not target.is_file():
                        raise ValueError("结果文件尚未生成。")
                    subprocess.run(["/usr/bin/open", "-R", str(target)], check=True, capture_output=True)
                self.reply(200, {"ok": True})
            elif parts.path.startswith("/api/jobs/") and parts.path.endswith("/repair-report"):
                self.reply(200, self.app.export_repair_report(parts.path.split("/")[3]))
            else:
                self.reply(404, {"error": "操作不存在。"})
        except KeyError:
            self.reply(404, {"error": "任务不存在。"})
        except (ValueError, TypeError) as exc:
            self.reply(400, {"error": str(exc)})
        except (BrokenPipeError, ConnectionResetError):
            pass
        except Exception:
            self.reply(500, {"error": "操作未完成，请检查文件权限或稍后重试。"})


def make_server(app, port=8765):
    server = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    server.app = app
    return server


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--open", action="store_true", help="启动后打开浏览器")
    args = parser.parse_args(argv)
    app = App(args.output_dir)
    try:
        server = make_server(app, args.port)
    except OSError:
        print("启动失败：端口被占用或无监听权限。可使用 --port 8766 更换端口。", file=sys.stderr)
        app.close()
        return 2
    url = f"http://127.0.0.1:{server.server_address[1]}"
    print(f"Caption QA 已启动：{url}\n输出目录：{app.output_dir}\n按 Ctrl+C 停止。", flush=True)
    if args.open:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("正在停止，请等待当前质检任务结束。", flush=True)
    finally:
        server.server_close()
        app.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
