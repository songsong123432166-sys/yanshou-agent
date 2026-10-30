"""Write only human-approved Caption changes; never overwrite source data."""
from __future__ import annotations

import csv
import json
from pathlib import Path
import re
import shutil
import tempfile
from uuid import uuid4
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import captionqa.caption_qa_agent as qa
from captionqa.caption_qa_review_store import atomic_json, validate_binding
from captionqa.caption_qa_review_state import summary
from captionqa.caption_qa_table_io import _schema, iter_table_records


def _unique_json(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("JSON 中存在重复字段，无法保证无损回填。")
        result[key] = value
    return result


def _xml_text(text):
    text = re.sub(r"_(?=x[0-9a-fA-F]{4}_)", "_x005F_", text)
    return re.sub(r"[\x00-\x08\x0b-\x1f\ufffe\uffff]", lambda m: f"_x{ord(m[0]):04X}_", text)


def patch_xlsx(source, target, cells, yellow_cells=None):
    """Patch selected text/fills only; retain unrelated ZIP members bytewise."""
    from captionqa.caption_qa_highlight import YellowStyles, _attr
    yellow_cells = yellow_cells or {}
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with ZipFile(source) as zin:
        workbook = ET.fromstring(zin.read("xl/workbook.xml"))
        relations = {r.get("Id"): r.get("Target") for r in ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))}
        sheets = {s.get("name"): relations[s.get("{" + rel_ns + "}id")] for s in workbook.find("{" + ns + "}sheets")}
        updates = {}
        styles = YellowStyles(zin.read('xl/styles.xml')) if any(yellow_cells.values()) else None
        from xml.sax.saxutils import escape
        for name in dict.fromkeys([*cells, *yellow_cells]):
            replacements = cells.get(name, {})
            yellow = yellow_cells.get(name, set())
            if name not in sheets:
                raise ValueError("原工作表已变化，无法回填。")
            member = sheets[name]
            member = member.lstrip("/") if member.startswith("/") else "xl/" + member
            import posixpath
            member = posixpath.normpath(member)
            raw = zin.read(member).decode("utf-8")
            # Optional-prefix backreferences do not match an absent group;
            # use two explicit patterns to support both common OOXML forms.
            wanted = set(replacements) | set(yellow)
            changed = set()
            styled = set()
            def replace(match):
                attrs = match.group("attrs")
                address = re.search(r'\br=["\']([^"\']+)["\']', attrs)
                if not address or address[1] not in wanted:
                    return match[0]
                if address[1] in yellow:
                    if address[1] in styled:
                        raise ValueError("工作表存在重复单元格地址。")
                    styled.add(address[1])
                    old_style = re.search(r'\bs=["\']([^"\']+)["\']', attrs)
                    attrs = _attr(attrs, 's', styles.style_id(old_style[1] if old_style else '0'))
                if address[1] not in replacements:
                    start, end = match.span('attrs')
                    return match[0][:start-match.start()] + attrs + match[0][end-match.start():]
                body = match.groupdict().get("body") or ""
                if re.search(r"<(?:\w+:)?f(?:\s|>)", body):
                    raise ValueError("目标 Caption 是公式，未改写。")
                if address[1] in changed:
                    raise ValueError("工作表存在重复单元格地址。")
                changed.add(address[1])
                prefix = match.groupdict().get("prefix") or ""
                attrs = re.sub(r'\s+t=["\'][^"\']*["\']', "", attrs)
                value = replacements[address[1]]
                if len(value) > 32767:
                    raise ValueError("修改后字段超过 Excel 单元格长度上限，未生成副本。")
                value = escape(_xml_text(value))
                return f'<{prefix}c{attrs} t="inlineStr"><{prefix}is><{prefix}t xml:space="preserve">{value}</{prefix}t></{prefix}is></{prefix}c>'
            # Qualified and unqualified cells are disjoint.
            for expression in (r'<(?P<prefix>[A-Za-z_][\w.-]*:)c\b(?P<attrs>[^>]*?)(?:/>|>(?P<body>.*?)</(?P=prefix)c>)',
                               r'<c\b(?P<attrs>[^>]*?)(?:/>|>(?P<body>.*?)</c>)'):
                raw = re.sub(expression, replace, raw, flags=re.S)
            if changed != set(replacements):
                raise ValueError("目标单元格缺失，拒绝猜测回填位置。")
            if styled != set(yellow):
                raise ValueError("标黄单元格缺失，拒绝猜测位置。")
            updates[member] = raw.encode("utf-8")
        if styles:
            updates['xl/styles.xml'] = styles.render()
        with ZipFile(target, "w") as zout:
            for info in zin.infolist():
                zout.writestr(info, updates.get(info.filename, zin.read(info)))


def _captions(document, changes, input_source):
    """Legacy adapter; final Caption construction belongs to ReviewPlan."""
    from captionqa.caption_qa_review_plan import compile_review
    from captionqa.caption_qa_review_store import source_records
    records = source_records(document, input_source)
    plan = compile_review(document, records).require_valid()
    if changes != plan.changes:
        raise ValueError('审核修改快照不一致，请刷新后重试。')
    return records, plan.captions, list(plan.audit), [r for r in plan.rechecks if r['row_index'] in plan.captions]


def _write_primary(document, source, input_source, temporary, records, captions, changes):
    """Write the complete table. Presentation and reports are separate steps."""
    suffix = source.suffix.lower()
    if suffix not in {'.csv', '.xlsx', '.json', '.jsonl'}:
        raise ValueError('不支持的回填格式。')
    if not captions:
        shutil.copyfile(source, temporary)
    elif suffix == ".csv":
        with source.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        schema = _schema(source, document["sheet"])
        column = schema.keys.index(document["caption_column"])
        for index, caption in captions.items():
            if len(rows[index]) <= column:
                raise ValueError("Caption 单元格不存在，拒绝扩展错位数据。")
            rows[index][column] = caption
        with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(rows)
    elif suffix == ".xlsx":
        from openpyxl.utils import get_column_letter
        cells = {}
        if input_source != source:
            for index, fields in changes.items():
                record = records[index]
                sheet = record.get("source_sheet")
                row = record.get("source_row")
                schema = _schema(source, sheet)
                for field in fields:
                    key = field["name"] + "v2"
                    if key not in schema.keys or type(row) is not int:
                        raise ValueError("分列 Caption 的来源映射不完整，未回填。")
                    original_record = next(r for i, r, _ in iter_table_records(source, sheet_name=sheet) if i == row - 1)
                    old_value = original_record[key] or ""
                    # Preserve a source cell's redundant same-field label.
                    match = re.match(r"^" + re.escape(field["name"]) + r"[：:][ \t]*", old_value)
                    prefix = match[0] if match and old_value[match.end():].strip() == field["original"].strip() else ""
                    value = prefix + field["after"].strip()
                    address = get_column_letter(schema.keys.index(key) + 1) + str(row)
                    cells.setdefault(sheet, {})[address] = value
        else:
            schema = _schema(source, document["sheet"])
            sheet = document["sheet"]
            if not sheet:
                from captionqa.caption_qa_table_io import _openpyxl
                book = _openpyxl().load_workbook(source, read_only=True)
                sheet = book.sheetnames[0]
                book.close()
            col = get_column_letter(schema.keys.index(document["caption_column"]) + 1)
            cells[sheet] = {col + str(index + 1): caption for index, caption in captions.items()}
        patch_xlsx(source, temporary, cells)
    elif suffix == ".jsonl":
        with source.open(encoding='utf-8-sig',newline='') as stream:
            lines = stream.read().splitlines(keepends=True)
        for index, caption in captions.items():
            value = json.loads(lines[index - 1], object_pairs_hook=_unique_json)
            value[document["caption_column"]] = caption
            ending = "\r\n" if lines[index - 1].endswith("\r\n") else "\n" if lines[index - 1].endswith("\n") else ""
            lines[index - 1] = json.dumps(value, ensure_ascii=False) + ending
        with temporary.open('w',encoding='utf-8',newline='') as stream:
            stream.write(''.join(lines))
    elif suffix == ".json":
        value = json.loads(source.read_text(encoding="utf-8-sig"), object_pairs_hook=_unique_json)
        items = value if isinstance(value, list) else [value]
        for index, caption in captions.items():
            if isinstance(items[index - 1], str):
                items[index - 1] = caption
            else:
                items[index - 1][document["caption_column"]] = caption
        temporary.write_text(json.dumps(items if isinstance(value, list) else items[0], ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        raise ValueError("不支持的回填格式。")


def _change_log(path, audit):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        columns = ['题目 ID', '字段', '修改来源', '修改前', '修改后', 'field_id']
        writer = csv.DictWriter(stream, columns, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(audit)


def _optional_file(destination, label, write, warnings):
    """A failed attachment never invalidates an already written primary table."""
    try:
        with tempfile.TemporaryDirectory(prefix='caption-attachment-', dir=destination.parent, ignore_cleanup_errors=True) as folder:
            temporary = Path(folder) / destination.name
            write(temporary)
            temporary.replace(destination)
        return str(destination)
    except Exception as exc:
        warnings.append(f'{label}未生成：{exc}；主文件已保留。')
        return ''


def export_approved(document, source, input_source, cache, output_dir, stem, *, automatic=False, plan=None):
    source, input_source, output_dir = Path(source), Path(input_source), Path(output_dir)
    validate_binding(document, source, input_source, cache)
    from captionqa.caption_qa_review_plan import compile_review, document_fingerprint
    from captionqa.caption_qa_review_store import source_records
    records = source_records(document, input_source)
    if plan is None:
        plan = compile_review(document, records)
    if plan.fingerprint != document_fingerprint(document):
        raise ValueError('审核预览已变化，请刷新后导出最新审核后版本。')
    plan.require_valid()
    for row in plan.rows:
        if row.row_index not in records or records[row.row_index].get(document['caption_column']) != row.original:
            raise ValueError('字段原文已变化，拒绝覆盖。')
    changes, captions = plan.changes, plan.captions
    marked, audit, rechecks = plan.highlights, list(plan.audit), list(plan.rechecks)
    automatic = (automatic and bool(changes) and not marked and
                 all(field['mode'] == 'automatic_rule' for fields in changes.values() for field in fields))
    label = '自动修正' if automatic else '审核后'
    name = f"{Path(stem).stem[:45]}_{label}_{uuid4().hex[:8]}"
    destination = output_dir / (name + source.suffix.lower())
    if destination.resolve() == source.resolve():
        raise ValueError('不能覆盖原文件。')
    warnings, highlight_warning, excel_details = [], '', {}
    highlighted_path = None
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='caption-approved-', dir=output_dir, ignore_cleanup_errors=True) as folder:
            temporary = Path(folder) / destination.name
            _write_primary(document, source, input_source, temporary, records, captions, changes)
            if marked:
                candidate = Path(folder) / (destination.stem+'_Caption标黄.xlsx')
                try:
                    from captionqa.caption_qa_highlight import excel_view, xlsx_cells
                    if source.suffix.lower() == '.xlsx':
                        yellow = xlsx_cells(document, records, source, input_source, marked)
                        patch_xlsx(temporary, candidate, {}, yellow)
                        candidate.replace(temporary)
                        highlighted_path = destination
                    else:
                        excel_details = excel_view(temporary, candidate, document, marked)
                        highlighted_path = output_dir / candidate.name
                except Exception as exc:
                    highlight_warning = f'标黄Excel未生成：{exc}；完整主文件仍可导出，标黄选择保留。'
                    warnings.append(highlight_warning)
                    highlighted_path = None
            # Reject stale data or invalid approved edits before publishing any
            # primary copy. Attachment failures above cannot bypass this check.
            validate_binding(document, source, input_source, cache)
            temporary.replace(destination)
            if highlighted_path and highlighted_path != destination:
                try:
                    candidate.replace(highlighted_path)
                except Exception as exc:
                    highlight_warning = f'标黄Excel未保存：{exc}；主文件已保留。'
                    warnings.append(highlight_warning)
                    highlighted_path = None
    except OSError as exc:
        raise ValueError('主数据文件未能保存，请检查输出目录权限和磁盘空间；审核记录已保留。') from exc
    log = destination.with_name(destination.stem+'_修改记录.csv')
    log_path = _optional_file(log, '修改记录', lambda p:_change_log(p, audit), warnings)
    highlight_ids_path = ''
    if marked:
        ids = list(dict.fromkeys(mark['sample_id'] for mark in marked.values()))
        path = destination.with_name(destination.stem+'_标黄ID名单.txt')
        highlight_ids_path = _optional_file(path, '标黄ID名单',
            lambda p:p.write_text(''.join(value+'\n' for value in ids), encoding='utf-8'), warnings)
    audit_path = destination.with_name(destination.stem+'_审核与复检.json')
    result = {'corrected_path':str(destination), 'corrected_sha256':qa._source_sha256(destination),
              'export_kind':'automatic' if automatic else 'reviewed', 'export_success':True,
              'change_log_path':log_path, 'highlight_ids_path':highlight_ids_path,
              'review_revision':plan.revision, 'plan_fingerprint':plan.fingerprint,
              'changed_captions':len(captions), 'changed_fields':len(audit),
              'highlighted_captions':len(marked), 'highlighted_path':str(highlighted_path) if highlighted_path else '',
              'highlight_warning':highlight_warning, 'excel_encoding':excel_details,
              'remaining_review_items':summary(document)['pending'],
              'remaining_static_issues':sum(len(r['remaining_static_errors']) for r in rechecks),
              'post_review_fields':len({w['field_id'] for w in plan.review_checks}),
              'semantic_rechecked':False, 'export_warnings':warnings}
    audit_document = {'review_revision':plan.revision, 'plan_fingerprint':plan.fingerprint, 'source_sha256':document['source_sha256'],
                      'decisions':document['issues'], 'events':document['events'], 'changes':audit,
                      'highlighted_captions':[{'row_index':index, **mark} for index,mark in marked.items()],
                      'highlight_scope':'keep_current_issue_other_changes_apply',
                      'highlighted_path':result['highlighted_path'], 'highlight_warning':highlight_warning,
                      'excel_encoding':excel_details, 'export_warnings':list(warnings),
                      'recheck':rechecks, 'model_requests':0, 'review_summary':summary(document),
                      'post_review_checks':plan.review_checks,
                      'export_kind':result['export_kind'], 'corrected_path':str(destination),
                      'corrected_sha256':result['corrected_sha256']}
    result['recheck_path'] = _optional_file(audit_path, '审核与复检说明',
        lambda p:atomic_json(p, audit_document), warnings)
    return result
