"""Portable XLSX authoring when the optional local render runtime is absent."""
from pathlib import Path
import os
import tempfile


def _book(headers,rows,title,widths,table_name):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title=title;ws.sheet_view.showGridLines=False
    for row in [headers,*rows]:
        if any(len(str(v))>32767 for v in row):raise ValueError('报告字段超过Excel单元格上限，请查看网页或CSV。')
        ws.append([str(v) for v in row])
    for row in ws:
        for cell in row:
            # IDs and captions are literal strings, including a leading equals sign.
            cell.data_type='s';cell.number_format='@'
            cell.font=Font(name='Arial',size=11,color='253247')
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='243A57')
        cell.font=Font(name='Arial',size=11,bold=True,color='FFFFFF')
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=32
    for n,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(n)].width=width
    if rows:
        table=Table(displayName=table_name,ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes='A2';return wb,ws


def _save(wb,destination,patch=None):
    destination=Path(destination);destination.parent.mkdir(parents=True,exist_ok=True)
    fd,name=tempfile.mkstemp(prefix='.caption-report-',suffix='.xlsx',dir=destination.parent);os.close(fd)
    temp=Path(name)
    try:
        wb.save(temp)
        if patch:patch(temp)
        temp.replace(destination)
    finally:
        wb.close();temp.unlink(missing_ok=True)


def review_xlsx(destination,headers,rows,highlights):
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from captionqa.caption_qa_review_report import patch_rich_text
    from captionqa.caption_qa_review_model import DECISION_LABELS
    wb,ws=_book(headers,rows,'人工审核',[28,22,18,48,85,70,22,85],'CaptionHumanReview')
    for n in range(2,len(rows)+2):
        ws.row_dimensions[n].height=120
        for col in (7,8):ws.cell(n,col).fill=PatternFill('solid',fgColor='FFF0D6')
    if rows:
        validation=DataValidation(type='list',formula1='"'+','.join(DECISION_LABELS.values())+'"')
        ws.add_data_validation(validation);validation.add(f'G2:G{len(rows)+1}')
    _save(wb,destination,lambda p:patch_rich_text(p,highlights))


def repair_xlsx(destination,entries,dry_run,complete,missing):
    wb,ws=_book(['题目 ID','错误项'],entries,'底线返修' if dry_run else '辅助复核',[27,100],'CaptionRepair')
    for n,row in enumerate(entries,2):ws.row_dimensions[n].height=min(409,max(34,(row[1].count('\n')+1)*18+14))
    if not complete or missing:
        ws.cell(len(entries)+3,2,'检测或 ID 尚不完整，本表仅含已发现且有 ID 的问题，请同时查看检测说明。')
    elif not entries:ws.cell(3,2,'当前已检测范围内无返修项。')
    _save(wb,destination)
