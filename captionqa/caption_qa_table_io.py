"""Lossless business-table input and positional QA-column export.

Original column headers may be empty, repeated, or even named ``qa_status``.
Records expose deterministic unique internal keys for inspection; exported tables
retain the original visible headers and append the five QA columns at the end.
For Excel input/output the source workbook is copied through openpyxl so other
worksheets, formulas, formatting, hyperlinks, and comments remain present.
"""
from __future__ import annotations

import csv
import json
import os
import shutil
import tempfile
from collections.abc import Iterator, Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from zipfile import ZipFile


QA_COLUMNS: list[str] = [
    "qa_status",
    "qa_missing_or_dup_subjects",
    "qa_issues_summary",
    "qa_human_review_reason",
    "qa_revised_caption",
]


@dataclass(frozen=True)
class _TableSchema:
    headers: tuple[Any, ...]
    keys: tuple[str, ...]


def _unique_keys(headers: list[Any]) -> tuple[str, ...]:
    """Generate safe dictionary keys without losing any positional columns."""
    bases = [str(value) if value is not None and str(value) else f"__column_{i}"
             for i, value in enumerate(headers, 1)]
    reserved, used = set(bases), set()
    keys: list[str] = []
    for base in bases:
        key, suffix = base, 2
        while key in used:
            key = f"{base}__{suffix}"
            while key in reserved or key in used:
                suffix += 1
                key = f"{base}__{suffix}"
        keys.append(key)
        used.add(key)
    return tuple(keys)


def _openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("Excel 支持需要 openpyxl，请运行 python -m pip install openpyxl") from exc
    return openpyxl


def _worksheet(workbook: Any, sheet_name: str | None) -> Any:
    if sheet_name is not None:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Excel 中不存在工作表 {sheet_name!r}；可用工作表: {workbook.sheetnames}")
        return workbook[sheet_name]
    if not workbook.worksheets:
        raise ValueError("Excel 文件中没有可用工作表")
    return workbook.worksheets[0]


def _streaming_worksheet(workbook: Any, sheet_name: str | None) -> Any:
    sheet = _worksheet(workbook, sheet_name)
    # Some spreadsheet exports declare dimension="A1" even with hundreds of
    # rows. Stream actual XML cells rather than trusting that cached range.
    sheet.reset_dimensions()
    return sheet


def _raise_csv_limit() -> None:
    # The stdlib default is 128 KiB, which is too small for long captions.
    csv.field_size_limit(max(csv.field_size_limit(), 2**31 - 1))


def _jsonl_records(path: Path) -> Iterator[tuple[int, dict[str, Any], str | None]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        for row_index, line in enumerate(stream, 1):
            raw = line.rstrip("\r\n")
            if not raw.strip():
                yield row_index, {}, None
                continue
            try:
                value = json.loads(raw)
            except json.JSONDecodeError as exc:
                yield row_index, {"_raw_line": raw}, f"JSONL 第 {row_index} 行解析失败: {exc}"
                continue
            if not isinstance(value, dict):
                yield row_index, {"_raw_line": raw, "_raw_value": value}, (
                    f"JSONL 第 {row_index} 行必须是 JSON 对象，实际为 {type(value).__name__}")
                continue
            yield row_index, value, None


def _json_records(path: Path) -> Iterator[tuple[int, dict[str, Any], str | None]]:
    """Retain the preexisting JSON-array/object/plain-string input interface."""
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        data = json.load(stream)
    for row_index, value in enumerate(data if isinstance(data, list) else [data], 1):
        record = (value if isinstance(value, dict) else
                  {"caption": value} if isinstance(value, str) else {"raw_input": value})
        yield row_index, record, None


def _schema(path: Path, sheet_name: str | None) -> _TableSchema:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        _raise_csv_limit()
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            headers: list[Any] = next(reader, [])
            width = len(headers)
            for row in reader:
                width = max(width, len(row))
            # Nonempty cells beyond a malformed short header still get retained.
            headers.extend([""] * (width - len(headers)))
        return _TableSchema(tuple(headers), _unique_keys(headers))
    if suffix == ".xlsx":
        workbook = _openpyxl().load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = _streaming_worksheet(workbook, sheet_name)
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows, ()))
            width = max(1, len(headers))
            for row in rows:
                width = max(width, len(row))
            headers.extend([None] * (width - len(headers)))
            return _TableSchema(tuple(headers), _unique_keys(headers))
        finally:
            workbook.close()
    if suffix in {".jsonl", ".json"}:
        columns: dict[str, None] = {}
        records = _jsonl_records(path) if suffix == ".jsonl" else _json_records(path)
        for _, record, _ in records:
            columns.update(dict.fromkeys(record))
        return _TableSchema(tuple(columns), tuple(columns))
    raise ValueError(f"不支持的输入格式 {suffix!r}；支持 .xlsx、.csv、.jsonl、.json")


def table_columns(path: str | Path, *, sheet_name: str | None = None) -> list[str]:
    """Return unique internal column names, including late JSONL keys."""
    return list(_schema(Path(path), sheet_name).keys)


def iter_table_records(
    path: str | Path, *, sheet_name: str | None = None,
) -> Iterator[tuple[int, dict[str, Any], str | None]]:
    """Read original row order, retaining blanks and duplicate/unnamed columns.

    CSV and Excel indexes count data rows from one (header excluded). JSONL
    indexes are physical line numbers. Invalid JSONL records retain their raw
    input plus an error for the caller; blank lines yield an empty dictionary.
    Excel formulas are returned as their formula text, never cached values.
    """
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".jsonl":
        yield from _jsonl_records(source)
        return
    if suffix == ".json":
        yield from _json_records(source)
        return
    schema = _schema(source, sheet_name)
    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            next(reader, None)
            for row_index, row in enumerate(reader, 1):
                values = row + [""] * (len(schema.keys) - len(row))
                yield row_index, dict(zip(schema.keys, values)), None
        return
    workbook = _openpyxl().load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = _streaming_worksheet(workbook, sheet_name)
        for row_index, row in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(schema.keys), values_only=True), 1,
        ):
            yield row_index, dict(zip(schema.keys, row)), None
    finally:
        workbook.close()


def _cell_value(value: Any) -> Any:
    """Serialize nested JSON cells without losing their contents."""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)
    return value


def _qa_values(row_index: int, results: Mapping[int, Mapping[str, Any]]) -> list[Any]:
    record = results.get(row_index, {})
    return [_cell_value(record.get(column, "")) for column in QA_COLUMNS]


def _excel_literal(cell: Any, value: Any, *, row_index: int, column: str) -> None:
    converted = _cell_value(value)
    # Excel numbers retain only 15 significant digits; keep long JSON IDs exact.
    if isinstance(converted, int) and len(str(abs(converted))) > 15:
        converted = str(converted)
    if isinstance(converted, str):
        if len(converted) > 32767:
            raise ValueError(
                f"第 {row_index} 行 {column!r} 长度 {len(converted)} 超过 Excel 单元格上限 "
                "32767，无法无损导出；请改用 CSV 或 JSONL。")
        cell.value = converted
        # openpyxl otherwise treats a leading '=' as a formula.
        cell.data_type = "s"
    else:
        cell.value = converted


def _write_csv(
    source: Path, temporary: Path, schema: _TableSchema,
    results: Mapping[int, Mapping[str, Any]], sheet_name: str | None,
) -> None:
    with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(list(schema.headers) + QA_COLUMNS)
        for row_index, record, _ in iter_table_records(source, sheet_name=sheet_name):
            values = [_cell_value(record.get(key, "")) for key in schema.keys]
            writer.writerow(values + _qa_values(row_index, results))
        stream.flush()
        os.fsync(stream.fileno())


def _preserve_xml_carriage_returns(path: Path) -> None:
    # ElementTree can serialize cell CR characters literally. XML readers then
    # normalize them to LF, silently altering CRLF text on the next open.
    with ZipFile(path) as source:
        if not any(b"\r" in source.read(item) for item in source.infolist()
                   if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml")):
            return
        with tempfile.TemporaryFile() as buffer:
            with ZipFile(buffer, "w") as target:
                for item in source.infolist():
                    data = source.read(item)
                    if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                        data = data.replace(b"\r", b"&#13;")
                    target.writestr(item, data)
            source.close()
            buffer.seek(0)
            with path.open("wb") as output:
                shutil.copyfileobj(buffer, output)


def _write_xlsx(
    source: Path, temporary: Path, schema: _TableSchema,
    results: Mapping[int, Mapping[str, Any]], sheet_name: str | None,
) -> None:
    openpyxl = _openpyxl()
    if len(schema.headers) + len(QA_COLUMNS) > 16384:
        raise ValueError("追加质检结果后超过 Excel 的 16384 列上限；请使用 CSV。")
    if source.suffix.lower() == ".xlsx":
        workbook = openpyxl.load_workbook(source, data_only=False, rich_text=True)
        try:
            sheet = _worksheet(workbook, sheet_name)
        except BaseException:
            workbook.close()
            raise
        # Do not reassign original cells; preserve their native attributes.
        indexes: Iterator[int] = iter(range(1, sheet.max_row))
    else:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        try:
            for column_index, value in enumerate(schema.headers, 1):
                _excel_literal(sheet.cell(1, column_index), value,
                               row_index=0, column=str(schema.keys[column_index - 1]))
            row_indexes: list[int] = []
            for row_index, record, _ in iter_table_records(source, sheet_name=sheet_name):
                row_indexes.append(row_index)
                for column_index, key in enumerate(schema.keys, 1):
                    _excel_literal(sheet.cell(row_index + 1, column_index), record.get(key, ""),
                                   row_index=row_index, column=key)
            indexes = iter(row_indexes)
        except BaseException:
            workbook.close()
            raise
    try:
        start_column = len(schema.headers) + 1
        for offset, column in enumerate(QA_COLUMNS):
            _excel_literal(sheet.cell(1, start_column + offset), column,
                           row_index=0, column=column)
        for row_index in indexes:
            for offset, (column, value) in enumerate(zip(QA_COLUMNS, _qa_values(row_index, results))):
                _excel_literal(sheet.cell(row_index + 1, start_column + offset), value,
                               row_index=row_index, column=column)
        workbook.save(temporary)
    finally:
        workbook.close()
    _preserve_xml_carriage_returns(temporary)


def export_qa_table(
    input_path: str | Path,
    output_path: str | Path,
    results: Mapping[int, Mapping[str, Any]],
    *,
    sheet_name: str | None = None,
) -> None:
    """Atomically export original columns plus five positional QA columns.

    ``results`` maps the original one-based row index to the five flat QA keys;
    completion order is irrelevant. The original source is never overwritten.
    Existing source columns with a QA name remain intact; new QA columns are
    always the final five columns, even when visible header names repeat.
    """
    source, target = Path(input_path), Path(output_path)
    if source.resolve() == target.resolve():
        raise ValueError("输入与输出文件路径相同；请另选结果路径，避免覆盖原表。")
    if target.suffix.lower() not in {".xlsx", ".csv"}:
        raise ValueError("表格输出格式仅支持 .xlsx 或 .csv；断点缓存请使用独立 JSONL 文件。")
    schema = _schema(source, sheet_name)
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(prefix=f".{target.stem}.", suffix=target.suffix,
                                        dir=target.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        if target.suffix.lower() == ".csv":
            _write_csv(source, temporary, schema, results, sheet_name)
        else:
            _write_xlsx(source, temporary, schema, results, sheet_name)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
