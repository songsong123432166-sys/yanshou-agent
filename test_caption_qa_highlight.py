"""Issue-local keep decisions and yellow Caption exports; no model/network calls."""
import asyncio
import csv
import json
import re
import xml.etree.ElementTree as ET
from unittest.mock import patch
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment

import caption_qa_agent as qa
from caption_qa_review import IdentificationAgent, make_review, decide, highlighted_rows, summary
from caption_qa_review_io import export_approved, patch_xlsx
from caption_qa_review_report import HEADERS, import_review, STATUS
from caption_qa_portable_exports import review_xlsx
from caption_qa_highlight import YELLOW
from caption_qa_split_input import prepare_split_input, HEADERS as SPLIT_HEADERS
from test_caption_qa_upgrades import GOOD
from test_caption_qa_split_input import GOOD as SPLIT_GOOD

RAW=GOOD.replace('<主体1>向前跑。','主体1向前跑')

class HighlightTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.cache=self.root/'cache.jsonl'

    def build(self,suffix='.csv',missing_id=False,second=True,note='=1+1'):
        self.cache.unlink(missing_ok=True)
        self.source=self.root/('source'+suffix)
        rows=[['样本ID','caption','备注'],['' if missing_id else '0001',RAW,note]]
        if second:rows.append(['0002',GOOD,'不可改'])
        if suffix=='.csv':
            with self.source.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
        elif suffix=='.xlsx':
            wb=Workbook();ws=wb.active;ws.title='源数据'
            for row in rows:ws.append(row)
            ws['B2'].font=Font(bold=True,color='FF123456');ws['B2'].alignment=Alignment(wrap_text=True)
            wb.create_sheet('其他')['A1']='=1+1';wb.save(self.source);wb.close()
        else:
            records=[dict(zip(rows[0],r)) for r in rows[1:]]
            self.source.write_text(json.dumps(records,ensure_ascii=False) if suffix=='.json' else '\n'.join(json.dumps(r,ensure_ascii=False) for r in records)+'\n')
        asyncio.run(qa.aprocess_batch(self.source,self.cache,IdentificationAgent(dry_run=True),column='caption'))
        self.input=self.source
        return make_review(self.cache,self.source,self.input,caption_column='caption',id_column='样本ID')

    def choose(self,doc,kind,decision):
        issue=next(i for i in doc['issues'] if i['kind']==kind)
        return decide(doc,revision=doc['revision'],issue_id=issue['issue_id'],decision=decision)

    def export(self,doc):return export_approved(doc,self.source,self.input,self.cache,self.root/'out','data')

    def test_only_selected_issue_kept_other_same_field_edit_applies_and_literal_csv_view(self):
        doc=self.choose(self.build(),'bare_tags','keep_highlight')
        doc=self.choose(doc,'missing_punctuation','approve');before=self.source.read_bytes()
        info=self.export(doc)
        with Path(info['corrected_path']).open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        self.assertIn('主体1向前跑。',rows[1][1]);self.assertNotIn('<主体1>向前跑',rows[1][1])
        self.assertEqual(rows[1][::2],['0001','=1+1']);self.assertEqual(info['highlighted_captions'],1)
        wb=load_workbook(info['highlighted_path']);ws=wb.active
        self.assertEqual([[c.value for c in row] for row in ws],rows)
        self.assertEqual(ws['B2'].fill.fgColor.rgb,YELLOW)
        self.assertEqual(ws['C2'].data_type,'s');self.assertEqual(ws['A2'].value,'0001')
        for c in ['A2','C2','B3']:self.assertNotEqual(ws[c].fill.fgColor.rgb,YELLOW)
        wb.close();self.assertEqual(self.source.read_bytes(),before)
        audit=json.loads(Path(info['recheck_path']).read_text());self.assertEqual(len(audit['changes']),1)
        self.assertEqual(audit['highlighted_captions'][0]['sample_id'],'0001');self.assertEqual(audit['model_requests'],0)

    def test_hidden_controls_and_literal_escape_strings_are_encoded_without_deleting_source(self):
        note='=1+1\x1e正文\x00\x0b\ufffe\uffff literal _x001E_ _x005F_\r\n\t'
        for suffix in ['.csv','.jsonl']:
            with self.subTest(suffix=suffix):
                doc=self.choose(self.build(suffix,note=note),'bare_tags','keep_highlight')
                doc=self.choose(doc,'missing_punctuation','approve');original=self.source.read_bytes()
                info=self.export(doc)
                if suffix=='.csv':
                    with Path(info['corrected_path']).open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
                    self.assertEqual(rows[1][2],note)
                else:self.assertEqual(json.loads(Path(info['corrected_path']).read_text().splitlines()[0])['备注'],note)
                with ZipFile(info['highlighted_path']) as z:
                    xml=z.read('xl/worksheets/sheet1.xml');tree=ET.fromstring(xml)
                ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                text=''.join(tree.find('.//s:c[@r="C2"]',ns).itertext())
                self.assertIn('_x001E_',text);self.assertIn('_x005F_x001E_',text)
                decoded=re.sub(r'_x([0-9a-fA-F]{4})_',lambda m:chr(int(m[1],16)),text)
                self.assertEqual(decoded,note);self.assertNotIn(b'\x1e',xml)
                wb=load_workbook(info['highlighted_path']);self.assertEqual(wb.active['B2'].fill.fgColor.rgb,YELLOW)
                self.assertEqual(wb.active['C2'].data_type,'s');wb.close()
                self.assertEqual(info['excel_encoding']['escaped_cells'][0]['cell'],'C2')
                self.assertEqual(info['highlight_warning'],'');self.assertEqual(self.source.read_bytes(),original)

    def test_excel_length_limit_does_not_block_lossless_csv_copy(self):
        note='长'*32768
        doc=self.choose(self.build(note=note),'bare_tags','keep_highlight')
        doc=self.choose(doc,'missing_punctuation','approve');original=self.source.read_bytes();info=self.export(doc)
        self.assertTrue(Path(info['corrected_path']).is_file());self.assertEqual(info['highlighted_path'],'')
        self.assertIn('C2（备注）',info['highlight_warning']);self.assertIn('32768',info['highlight_warning'])
        with Path(info['corrected_path']).open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        self.assertEqual(rows[1][2],note);self.assertIn('主体1向前跑。',rows[1][1])
        self.assertFalse(list((self.root/'out').glob('*.xlsx')))
        audit=json.loads(Path(info['recheck_path']).read_text())
        self.assertEqual(audit['highlighted_captions'][0]['sample_id'],'0001')
        self.assertEqual(audit['highlight_warning'],info['highlight_warning'])
        self.assertEqual(self.source.read_bytes(),original)

    def test_optional_excel_failure_is_reported_without_blocking_primary(self):
        doc=self.choose(self.build(),'bare_tags','keep_highlight')
        with patch('caption_qa_highlight.excel_view',side_effect=ValueError('unexpected mapping failure')):
            info=self.export(doc)
        self.assertTrue(Path(info['corrected_path']).is_file())
        self.assertIn('unexpected mapping failure',info['highlight_warning'])
        self.assertEqual(Path(info['highlight_ids_path']).read_text(),'0001\n')

    def test_color_only_export_and_reset_deduplicate_rows(self):
        doc=self.choose(self.build(),'bare_tags','keep_highlight')
        doc=self.choose(doc,'missing_punctuation','keep_highlight')
        self.assertEqual(summary(doc)['highlighted_captions'],1)
        info=self.export(doc)
        self.assertEqual(info['changed_captions'],0)
        self.assertEqual(Path(info['corrected_path']).read_bytes(),self.source.read_bytes())
        doc=self.choose(doc,'bare_tags','pending');self.assertEqual(len(highlighted_rows(doc)),1)
        doc=self.choose(doc,'missing_punctuation','keep');self.assertFalse(highlighted_rows(doc))
        info=self.export(doc);self.assertEqual(info['highlighted_captions'],0)
        self.assertEqual(Path(info['corrected_path']).read_bytes(),self.source.read_bytes())

    def test_all_keep_or_reject_still_exports_a_complete_reviewed_version(self):
        for suffix in ['.csv','.xlsx','.json','.jsonl']:
            with self.subTest(suffix=suffix):
                doc=self.build(suffix)
                for i,issue in enumerate(doc['issues']):
                    doc=decide(doc,revision=doc['revision'],issue_id=issue['issue_id'],decision='keep' if i%2 else 'reject')
                result=self.export(doc)
                self.assertTrue(result['export_success']);self.assertEqual(result['export_kind'],'reviewed')
                self.assertEqual(result['changed_captions'],0);self.assertEqual(result['remaining_review_items'],0)
                self.assertEqual(Path(result['corrected_path']).read_bytes(),self.source.read_bytes())

    def test_log_or_audit_failure_does_not_make_a_successful_primary_export_fail(self):
        for name in ['_change_log','atomic_json']:
            with self.subTest(attachment=name):
                doc=self.choose(self.build(),'bare_tags','keep_highlight')
                doc=self.choose(doc,'missing_punctuation','approve')
                with patch('caption_qa_review_io.'+name,side_effect=OSError('synthetic attachment failure')):
                    result=self.export(doc)
                self.assertTrue(result['export_success']);self.assertTrue(Path(result['corrected_path']).is_file())
                self.assertTrue(Path(result['highlighted_path']).is_file());self.assertTrue(result['export_warnings'])
                key='change_log_path' if name=='_change_log' else 'recheck_path'
                self.assertEqual(result[key],'')
                with Path(result['corrected_path']).open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
                self.assertIn('主体1向前跑。',rows[1][1]);self.assertEqual(rows[1][2],'=1+1')

    def test_xlsx_highlight_failure_keeps_approved_text_and_exports_marked_id_list(self):
        doc=self.choose(self.build('.xlsx'),'bare_tags','keep_highlight')
        doc=self.choose(doc,'missing_punctuation','approve')
        with patch('caption_qa_highlight.xlsx_cells',side_effect=ValueError('synthetic marking failure')):
            info=self.export(doc)
        self.assertTrue(info['export_success']);self.assertEqual(info['highlighted_path'],'')
        self.assertIn('synthetic marking failure',info['highlight_warning'])
        self.assertEqual(Path(info['highlight_ids_path']).read_text(),'0001\n')
        wb=load_workbook(info['corrected_path']);self.assertIn('主体1向前跑。',wb['源数据']['B2'].value)
        self.assertNotEqual(wb['源数据']['B2'].fill.fgColor.rgb,YELLOW);wb.close()

    def test_primary_write_error_is_distinct_and_does_not_publish_a_partial_table(self):
        doc=self.choose(self.build(),'missing_punctuation','approve')
        with patch('caption_qa_review_io._write_primary',side_effect=PermissionError('synthetic primary failure')):
            with self.assertRaisesRegex(ValueError,'主数据文件未能保存'):self.export(doc)
        self.assertFalse(list((self.root/'out').glob('*.csv')))

    def test_xlsx_only_target_cell_fill_and_text_preserving_other_members_and_styles(self):
        doc=self.choose(self.build('.xlsx'),'bare_tags','keep_highlight')
        doc=self.choose(doc,'missing_punctuation','approve');before=self.source.read_bytes()
        info=self.export(doc);self.assertEqual(info['highlighted_path'],info['corrected_path'])
        wb=load_workbook(info['corrected_path']);ws=wb['源数据']
        self.assertIn('主体1向前跑。',ws['B2'].value);self.assertEqual(ws['B2'].fill.fgColor.rgb,YELLOW)
        self.assertTrue(ws['B2'].font.bold);self.assertTrue(ws['B2'].alignment.wrap_text)
        self.assertEqual(ws['B2'].font.color.rgb,'FF123456');self.assertEqual(wb['其他']['A1'].value,'=1+1')
        wb.close()
        with ZipFile(self.source) as a,ZipFile(info['corrected_path']) as b:
            self.assertEqual({n for n in a.namelist() if a.read(n)!=b.read(n)}, {'xl/styles.xml','xl/worksheets/sheet1.xml'})
        self.assertEqual(self.source.read_bytes(),before)

    def test_automatic_and_manual_changes_continue_with_mark(self):
        from caption_qa_review import SafeAutoFixAgent
        doc=self.build();self.cache.unlink()
        # Missing subject stays human; its same-field final period stays automatic.
        raw=GOOD.replace('<主体1>向前跑。','男子向前跑')
        with self.source.open('w',encoding='utf-8-sig',newline='') as f:
            csv.writer(f).writerows([['样本ID','caption','备注'],['0001',raw,'note']])
        asyncio.run(qa.aprocess_batch(self.source,self.cache,SafeAutoFixAgent(dry_run=True),column='caption'))
        doc=make_review(self.cache,self.source,self.source,caption_column='caption',id_column='样本ID')
        doc=self.choose(doc,'missing_subjects','keep_highlight');info=self.export(doc)
        with Path(info['corrected_path']).open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        self.assertIn('男子向前跑。',rows[1][1]);self.assertEqual(info['highlighted_captions'],1)
        issue=next(i for i in doc['issues'] if i['kind']=='missing_subjects')
        field=next(f for f in doc['fields'] if f['field_id']==issue['field_id'])
        doc=decide(doc,revision=doc['revision'],field_id=field['field_id'],final_text='男子向前走。')
        self.assertEqual(next(i for i in doc['issues'] if i['issue_id']==issue['issue_id'])['decision'],'keep_highlight')
        info=self.export(doc)
        with Path(info['corrected_path']).open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        self.assertIn('男子向前走。',rows[1][1]);self.assertEqual(info['highlighted_captions'],1)

    def test_highlight_only_xlsx_preserves_formula_and_reuses_original_style(self):
        self.build('.xlsx');target=self.root/'marked.xlsx'
        patch_xlsx(self.source,target,{}, {'其他':{'A1'}})
        wb=load_workbook(target);self.assertEqual(wb['其他']['A1'].value,'=1+1')
        self.assertEqual(wb['其他']['A1'].fill.fgColor.rgb,YELLOW);wb.close()
        with self.assertRaisesRegex(ValueError,'缺失'):patch_xlsx(self.source,self.root/'bad.xlsx',{}, {'其他':{'Z10'}})

    def test_source_revision_and_missing_id_guards(self):
        doc=self.build(missing_id=True)
        with self.assertRaisesRegex(ValueError,'真实题目ID'):self.choose(doc,'bare_tags','keep_highlight')
        doc=self.choose(self.build(),'bare_tags','keep_highlight')
        with self.assertRaisesRegex(ValueError,'版本'):decide(doc,revision=0,issue_id=doc['issues'][0]['issue_id'],decision='pending')
        self.source.write_bytes(self.source.read_bytes()+b'\n')
        with self.assertRaisesRegex(ValueError,'原文或检测版本'):self.export(doc)
        self.assertFalse((self.root/'out').exists())

    def test_review_csv_import_and_excel_choice_roundtrip(self):
        doc=self.choose(self.build(),'bare_tags','keep_highlight')
        rows=[[i['issue_id'],i['sample_id'],i['field'],i['label'],'\n'.join(i['excerpts']),i['subject'],STATUS[i['decision']],''] for i in doc['issues']]
        target=self.root/'review.csv'
        with target.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.writer(f);w.writerow(HEADERS);w.writerows(rows)
        restored=import_review(doc,target,revision=doc['revision']);self.assertEqual(highlighted_rows(restored),highlighted_rows(doc))
        xlsx=self.root/'review.xlsx';review_xlsx(xlsx,HEADERS,rows,{})
        wb=load_workbook(xlsx);self.assertIn('保留原文并标黄',list(wb.active.data_validations.dataValidation)[0].formula1);wb.close()
        restored=import_review(doc,xlsx,revision=doc['revision']);self.assertEqual(highlighted_rows(restored),highlighted_rows(doc))

    def test_json_and_jsonl_companion_preserves_ids_and_columns(self):
        for suffix in ['.json','.jsonl']:
            with self.subTest(suffix=suffix):
                doc=self.choose(self.build(suffix),'bare_tags','keep_highlight');info=self.export(doc)
                self.assertEqual(Path(info['corrected_path']).read_bytes(),self.source.read_bytes())
                wb=load_workbook(info['highlighted_path']);ws=wb.active
                self.assertEqual(ws['A2'].value,'0001');self.assertEqual(ws['C2'].value,'=1+1')
                self.assertEqual(ws['B2'].fill.fgColor.rgb,YELLOW);wb.close()

    def test_split_xlsx_marks_caption_fields_only_and_keeps_other_changes(self):
        self.source=self.root/'split.xlsx';wb=Workbook();ws=wb.active;ws.title='数据'
        fields=list(SPLIT_GOOD);fields[-1]='0-1秒，主体1站立在<场景1>中'
        ws.append([*SPLIT_HEADERS,'备注']);ws.append(['001',*fields,'不改'])
        wb.save(self.source);wb.close();prepared=prepare_split_input(self.source);self.input=prepared['path']
        asyncio.run(qa.aprocess_batch(self.input,self.cache,IdentificationAgent(dry_run=True),column='caption'))
        doc=make_review(self.cache,self.source,self.input,caption_column='caption',id_column='id')
        doc=self.choose(doc,'bare_tags','keep_highlight');doc=self.choose(doc,'missing_punctuation','approve');info=self.export(doc)
        wb=load_workbook(info['corrected_path']);ws=wb['数据']
        self.assertEqual(ws['G2'].value,fields[-1]+'。')
        for col in 'BCDEFG':self.assertEqual(ws[col+'2'].fill.fgColor.rgb,YELLOW)
        for col in 'AH':self.assertNotEqual(ws[col+'2'].fill.fgColor.rgb,YELLOW)
        self.assertEqual(ws['H2'].value,'不改');wb.close()

if __name__=='__main__':unittest.main()
