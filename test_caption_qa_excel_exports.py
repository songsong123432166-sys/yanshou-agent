"""Spreadsheet-export regressions: incorrect A1 dimensions and cover sheets."""
from __future__ import annotations

import io
import json
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

import caption_qa_agent as qa
import caption_qa_web as web
from caption_qa_table_io import iter_table_records, table_columns
from caption_qa_table_io import export_qa_table
from test_caption_qa_upgrades import GOOD


def break_dimensions(path):
    output = io.BytesIO()
    with ZipFile(path) as source, ZipFile(output, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', data)
            target.writestr(item, data)
    path.write_bytes(output.getvalue())


def make_workbook(path, *, multiple=False, cover=True, reference=True):
    book = Workbook()
    sheet = book.active
    if cover:
        sheet.title = "运镜说明"
        sheet.append(["类型", "说明"])
        sheet.append(["固定镜头", "说明页，不是 Caption"])
    else:
        book.remove(sheet)
    if reference:
        sheet = book.create_sheet("batch1_参考")
        sheet.append(["id", "Caption"])
        for i in range(4):
            sheet.append([f"ref{i}", GOOD])
    for name in (["S1_200条", "S2_200条"] if multiple else ["S1_200条"]):
        sheet = book.create_sheet(name)
        sheet.append(["样本ID", "Caption", "原始备注"])
        sheet.append(["0001", GOOD, "保留原列"])
        sheet.append(["0002", GOOD.replace("身穿白袍。", "身穿白袍"), "后段单元格"])
    book.save(path)
    book.close()
    break_dimensions(path)


class ExcelExportCompatibilityTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)
        self.source = self.root / "export.xlsx"
        make_workbook(self.source)
        self.app = web.App(self.root / "Downloads", self.root / "work")

    def tearDown(self):
        self.app.close()
        self.folder.cleanup()

    def finish(self):
        self.app.executor.submit(lambda: None).result(timeout=10)

    def upload(self, **options):
        data = self.source.read_bytes()
        item = self.app.create_job("业务表.xlsx", io.BytesIO(data), len(data), review_first=False, dry_run=True, **options)
        self.finish()
        return self.app.jobs[item["id"]]

    def test_incorrect_dimension_does_not_truncate_headers_rows_or_cells(self):
        self.assertEqual(table_columns(self.source, sheet_name="S1_200条"), ["样本ID", "Caption", "原始备注"])
        rows = list(iter_table_records(self.source, sheet_name="S1_200条"))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1]["样本ID"], "0002")
        self.assertEqual(rows[1][1]["原始备注"], "后段单元格")
        self.assertTrue(qa.is_human_caption(rows[0][1]["Caption"]))

    def test_real_pipeline_selects_data_sheet_and_preserves_other_sheets(self):
        job = self.upload()
        self.assertEqual(job.status, "done", job.error)
        self.assertEqual(job.sheet, "S1_200条")
        self.assertEqual(job.snapshot()["processed"], 2)
        self.assertEqual(job.output.suffix, ".txt")
        self.assertEqual(job.output.read_text().splitlines(), ["0002"])
        repairs = list(self.app.output_dir.glob("*_返修表.xlsx"))
        self.assertEqual(len(repairs), 1)
        repair = load_workbook(repairs[0])
        try:
            self.assertEqual(repair.active.max_column, 2)
            self.assertEqual(repair.active["A2"].value, "0002")
            self.assertEqual([cell.value for cell in repair.active[1]], ["题目 ID", "错误项"])
        finally:
            repair.close()
        book = load_workbook(self.source)
        self.assertEqual(book["运镜说明"]["B2"].value, "说明页，不是 Caption")
        self.assertEqual(book["S1_200条"]["A2"].value, "0001")
        self.assertEqual(book["S1_200条"].max_column, 3)
        self.assertEqual(book["batch1_参考"].max_column, 2)
        book.close()

    def test_multiple_data_sheets_require_choice_then_retry_without_reupload(self):
        make_workbook(self.source, multiple=True)
        job = self.upload()
        self.assertEqual(job.status, "failed")
        self.assertIn("多个", job.error)
        self.assertEqual(len(job.sheet_choices), 3)
        self.assertFalse(job.cache.exists())
        self.app.retry(job.id, sheet="S2_200条")
        self.finish()
        self.assertEqual(job.status, "done", job.error)
        self.assertEqual(job.sheet, "S2_200条")

    def test_explicit_sheet_is_not_silently_changed(self):
        job = self.upload(sheet="batch1_参考")
        self.assertEqual(job.status, "done", job.error)
        self.assertEqual(job.sheet, "batch1_参考")
        self.assertEqual(job.snapshot()["processed"], 4)

    def test_empty_workbook_has_clear_error_before_keychain_or_api(self):
        book = Workbook()
        book.save(self.source)
        book.close()
        with patch.object(web, "read_saved_key", side_effect=AssertionError("must not read")):
            data = self.source.read_bytes()
            item = self.app.create_job("empty.xlsx", io.BytesIO(data), len(data), review_first=False)
            self.finish()
        job = self.app.jobs[item["id"]]
        self.assertEqual(job.status, "failed")
        self.assertIn("未找到", job.error)

    def test_history_survives_restart_without_sending_requests(self):
        job = self.upload()
        with patch.object(web, "read_saved_key", side_effect=AssertionError("must not read")):
            other = web.App(self.app.output_dir, self.app.work_dir)
            try:
                restored = other.jobs[job.id]
                self.assertEqual(restored.status, "done")
                self.assertEqual(restored.sheet, "S1_200条")
                self.assertEqual(restored.snapshot()["processed"], 2)
            finally:
                other.close()
        saved = json.loads((self.app.work_dir / "session.json").read_text())
        self.assertNotIn("api_key", saved["jobs"][0])

    def test_changing_sheet_with_existing_cache_cannot_mix_results(self):
        job = self.upload()
        job.status = "partial"
        with self.assertRaisesRegex(ValueError, "续传记录"):
            self.app.retry(job.id, sheet="batch1_参考")

    def test_empty_header_can_still_retain_columns_from_later_rows(self):
        book = Workbook()
        sheet = book.active
        sheet.append([None, None, None])
        sheet.append(["0001", GOOD, "末列数据"])
        book.save(self.source)
        book.close()
        break_dimensions(self.source)
        columns = table_columns(self.source)
        row = next(iter_table_records(self.source))[1]
        self.assertEqual(len(columns), 3)
        self.assertEqual(list(row.values()), ["0001", GOOD, "末列数据"])

    def test_original_and_qa_crlf_text_survive_excel_roundtrip(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["id", "Caption"])
        sheet.append(["0001", "第一行\r\n第二行\r第三行"])
        book.save(self.source)
        book.close()
        # Build a valid CR-bearing XML fixture, as used by the business export.
        buffer = io.BytesIO()
        with ZipFile(self.source) as source, ZipFile(buffer, "w") as target:
            for item in source.infolist():
                data = source.read(item)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"\r", b"&#xD;")
                target.writestr(item, data)
        self.source.write_bytes(buffer.getvalue())
        output = self.root / "qa.xlsx"
        revised = "修订第一行\r\n修订第二行"
        export_qa_table(self.source, output, {1: {"qa_revised_caption": revised}})
        book = load_workbook(output)
        self.assertEqual(book.active["B2"].value, "第一行\r\n第二行\r第三行")
        self.assertEqual(book.active["G2"].value, revised)
        book.close()


if __name__ == "__main__":
    unittest.main()
