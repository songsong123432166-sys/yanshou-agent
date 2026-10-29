"""Review-first invariants: explicit consent, source binding and exact scope."""
import asyncio
import copy
import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import caption_qa_agent as qa
from caption_qa_review import IdentificationAgent, make_review, decide, approved_changes, summary
from caption_qa_review_io import export_approved, patch_xlsx
from caption_qa_review_report import HEADERS, import_review, export_review
from test_caption_qa_upgrades import GOOD, FakeSyncClient, FakeAsyncClient, semantic_reply, semantic_issue
from test_caption_qa_split_input import workbook, HEADERS as SPLIT_HEADERS, GOOD as SPLIT_GOOD
from caption_qa_split_input import prepare_split_input
from openpyxl import load_workbook
import caption_qa_web as web
import test_caption_qa_web as web_test_helpers


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.source=self.root/'source.csv';self.cache=self.root/'cache.jsonl'

    def build(self,caption=GOOD,records=None):
        self.cache.unlink(missing_ok=True)
        with self.source.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream);writer.writerow(['样本ID','caption','备注'])
            writer.writerows(records or [('0001',caption,'不得修改')])
        asyncio.run(qa.aprocess_batch(self.source,self.cache,IdentificationAgent(dry_run=True),column='caption'))
        return make_review(self.cache,self.source,self.source,caption_column='caption',id_column='样本ID')

    def approve(self,doc,issue):
        return decide(doc,revision=doc['revision'],issue_id=issue['issue_id'],decision='approve')

    def test_detector_returns_original_and_keeps_suggestions_separate_sync_and_async(self):
        raw=GOOD.replace('身穿白袍。','身穿白袍')
        sync=IdentificationAgent(dry_run=True).check(raw)
        self.assertEqual(sync['revised_caption'],raw)
        self.assertNotEqual(sync['proposed_caption'],raw)
        self.assertEqual(sync['status'],'NEED_HUMAN_REVIEW')
        self.assertFalse(any(i['resolved'] for i in sync['issues']))
        self.build(raw)
        cached=json.loads(self.cache.read_text())
        self.assertEqual(cached['qa_revised_caption'],raw)
        self.assertEqual(cached['qa_checkpoint']['workflow_version'],'review-first-v1')

    def test_one_approval_does_not_fix_another_issue_in_same_field(self):
        raw=GOOD.replace('<主体1>向前跑。','主体1向前跑')
        doc=self.build(raw)
        issue=next(i for i in doc['issues'] if i['kind']=='missing_punctuation')
        self.assertTrue(issue['operations'])
        approved=self.approve(doc,issue)
        info=export_approved(approved,self.source,self.source,self.cache,self.root,'data.csv')
        with Path(info['corrected_path']).open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.reader(stream))
        self.assertIn('主体1向前跑。',rows[1][1]);self.assertNotIn('<主体1>向前跑',rows[1][1])
        self.assertEqual(rows[1][0],'0001');self.assertEqual(rows[1][2],'不得修改')
        self.assertGreater(info['remaining_review_items'],0)
        self.assertGreater(info['remaining_static_issues'],0)
        self.assertIn('主体1向前跑\n',self.source.read_text())

    def test_existing_tag_format_can_be_approved_independently(self):
        doc=self.build(GOOD.replace('<主体1>向前跑。','主体1向前跑。'))
        issue=next(i for i in doc['issues'] if i['kind']=='bare_tags')
        doc=self.approve(doc,issue)
        changes=approved_changes(doc)
        self.assertIn('<主体1>向前跑。',next(iter(changes.values()))[0]['after'])

    def test_no_unreviewed_changes_and_missing_id_not_fabricated(self):
        raw=GOOD.replace('身穿白袍。','身穿白袍')
        doc=self.build(raw)
        result=export_approved(doc,self.source,self.source,self.cache,self.root,'data')
        self.assertEqual(Path(result['corrected_path']).read_bytes(),self.source.read_bytes())
        self.assertEqual(result['changed_captions'],0);self.assertGreater(result['remaining_review_items'],0)
        doc=self.build(records=[('',raw,'note')])
        self.assertEqual(doc['issues'][0]['sample_id'],'')
        with self.assertRaisesRegex(ValueError,'真实题目 ID'):
            self.approve(doc,doc['issues'][0])

    def test_source_changes_or_stale_revision_stop_apply_and_decisions(self):
        doc=self.build(GOOD.replace('身穿白袍。','身穿白袍'))
        approved=self.approve(doc,doc['issues'][0])
        with self.assertRaisesRegex(ValueError,'版本已变化'):
            decide(approved,revision=0,issue_id=approved['issues'][0]['issue_id'],decision='reject')
        self.source.write_text(self.source.read_text()+'\n')
        with self.assertRaisesRegex(ValueError,'原文或检测版本'):
            export_approved(approved,self.source,self.source,self.cache,self.root,'data')

    def test_manual_grammar_draft_is_written_exactly_and_audited(self):
        doc=self.build(GOOD.replace('身穿白袍。','身穿白袍'))
        field=doc['fields'][0];final='<主体1>：穿白袍的男子，站在门边。'
        doc=decide(doc,revision=0,field_id=field['field_id'],final_text=final)
        info=export_approved(doc,self.source,self.source,self.cache,self.root,'data')
        with Path(info['corrected_path']).open(encoding='utf-8-sig',newline='') as stream:
            rows=list(csv.reader(stream))
        self.assertEqual(qa.parse_caption_sections(rows[1][1])['主体'].strip(),final)
        log=json.loads(Path(info['recheck_path']).read_text())
        self.assertEqual(log['changes'][0]['修改后'].strip(),final)
        self.assertEqual(log['changes'][0]['修改来源'],'human')
        self.assertEqual(log['model_requests'],0)
        self.assertFalse(info['semantic_rechecked'])

    def test_manual_draft_cannot_smuggle_additional_caption_field(self):
        doc=self.build(GOOD.replace('身穿白袍。','身穿白袍'))
        field=doc['fields'][0]
        doc=decide(doc,revision=0,field_id=field['field_id'],final_text='<主体1>。\n## 其他\n偷偷新增。')
        with self.assertRaisesRegex(ValueError,'标题或边界'):
            export_approved(doc,self.source,self.source,self.cache,self.root,'data')

    def test_audio_internal_punctuation_has_no_auto_operation_outer_period_does(self):
        raw=GOOD.replace('0-4秒，第0秒，<主体1>向前跑。','0-4秒，<主体1>说道{你好。}')
        doc=self.build(raw)
        ending=next(i for i in doc['issues'] if i['kind']=='missing_punctuation')
        self.assertTrue(ending['operations'])
        approved=self.approve(doc,ending)
        value=next(iter(approved_changes(approved).values()))[0]['after']
        self.assertTrue(value.strip().endswith('{你好。}。'))
        from caption_qa_review import touches_audio
        for text in ['（背景音乐轻快,节奏稳定）','<传来脚步声,敲门声>','{你好,再见}']:
            at=text.index(',');self.assertTrue(touches_audio(text,at,at+1))
            self.assertFalse(touches_audio(text,len(text),len(text)))

    def test_csv_import_is_atomic_conflicting_drafts_and_altered_ids_rejected(self):
        doc=self.build(GOOD.replace('<主体1>向前跑。','主体1向前跑'))
        target=self.root/'review.csv'
        def write(rows):
            with target.open('w',encoding='utf-8-sig',newline='') as f:
                w=csv.writer(f);w.writerow(HEADERS);w.writerows(rows)
        rows=[[i['issue_id'],i['sample_id'],i['field'],i['label'],'\n'.join(i['excerpts']),i['subject'],'待审核',''] for i in doc['issues']]
        rows[0][6]='确认按建议改';write(rows)
        result=import_review(doc,target,revision=0)
        self.assertEqual(result['issues'][0]['decision'],'approve')
        self.assertEqual(doc['issues'][0]['decision'],'pending')
        rows[0][1]='伪造ID';write(rows)
        with self.assertRaisesRegex(ValueError,'题目 ID'):
            import_review(doc,target,revision=0)
        rows[0][1]='0001'
        for n,row in enumerate(rows):row[6]='已人工修改';row[7]=f'不同人工稿{n}。'
        write(rows)
        with self.assertRaisesRegex(ValueError,'不同人工稿'):
            import_review(doc,target,revision=0)

    def test_report_has_editable_columns_red_excerpt_and_full_subject(self):
        doc=self.build(GOOD.replace('<主体1>向前跑。','男子向前跑。'))
        info=export_review(doc,self.root/'ids.txt')
        book=load_workbook(info['repair_report_path'],rich_text=True)
        self.addCleanup(book.close)
        sheet=book.active
        self.assertEqual([c.value for c in sheet[1]],HEADERS)
        self.assertEqual(sheet['B2'].value,'0001')
        self.assertIn('<主体1>',str(sheet['E2'].value))
        self.assertIn('<主体1>',str(sheet['F2'].value))
        from openpyxl.cell.rich_text import TextBlock
        self.assertTrue(any(isinstance(run,TextBlock) and run.font.color.rgb=='FFCC0000' for run in sheet['E2'].value))
        self.assertEqual(sheet['G2'].value,'待审核')
        self.assertTrue(sheet.data_validations.dataValidation)

    def test_split_xlsx_backfill_preserves_other_cells_and_archive_members(self):
        source=self.root/'split.xlsx'
        bad=list(SPLIT_GOOD);bad[2]='<主体1>：一个人'
        workbook(source,[('数据',[list(SPLIT_HEADERS)+['备注'],['0001',*bad,'不能改变']]),('其他',[['内容'],['原样保留']])])
        prepared=prepare_split_input(source,sheet_name='数据')
        asyncio.run(qa.aprocess_batch(prepared['path'],self.cache,IdentificationAgent(dry_run=True),column='caption'))
        doc=make_review(self.cache,source,prepared['path'],caption_column='caption',id_column='id')
        issue=next(i for i in doc['issues'] if i['kind']=='missing_punctuation')
        doc=self.approve(doc,issue)
        before=source.read_bytes()
        info=export_approved(doc,source,prepared['path'],self.cache,self.root,'source')
        book=load_workbook(info['corrected_path']);self.addCleanup(book.close)
        self.assertEqual(book['数据']['D2'].value,'<主体1>：一个人。')
        self.assertEqual(book['数据']['H2'].value,'不能改变')
        self.assertEqual(book['其他']['A2'].value,'原样保留')
        with ZipFile(source) as original,ZipFile(info['corrected_path']) as updated:
            changed=[n for n in original.namelist() if original.read(n)!=updated.read(n)]
            self.assertEqual(changed,['xl/worksheets/sheet1.xml'])
        self.assertEqual(source.read_bytes(),before)


class ReviewWebTests(unittest.TestCase):
    request=web_test_helpers.WebTests.request

    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        root=Path(self.temp.name)
        self.app=web.App(root/'out',root/'jobs');self.addCleanup(self.app.close)
        self.report_patch=patch('caption_qa_review_report.export_review',return_value={})
        self.report_patch.start();self.addCleanup(self.report_patch.stop)

    def job(self,caption=None):
        data=(json.dumps({'样本ID':'0009','caption':caption or GOOD.replace('身穿白袍。','身穿白袍'),'note':'unchanged'},ensure_ascii=False)+'\n').encode()
        item=self.app.create_job('data.jsonl',io.BytesIO(data),len(data),dry_run=True,auto_fix=False)
        self.app.executor.submit(lambda:None).result(timeout=10)
        self.assertEqual(self.app.jobs[item['id']].status,'done',self.app.jobs[item['id']].error)
        return item['id']

    def test_default_web_identifies_then_approval_apply_and_restart_keeps_review(self):
        identifier=self.job();job=self.app.jobs[identifier]
        self.assertEqual(job.workflow_version,'review-first-v1')
        row=json.loads(job.cache.read_text());self.assertEqual(row['qa_revised_caption'],row['caption'])
        status,body=self.request('GET',f'/api/jobs/{identifier}/review')
        self.assertEqual(status,200,body)
        doc=json.loads(body)['review'];issue=doc['issues'][0]
        status,body=self.request('POST',f'/api/jobs/{identifier}/review',body=json.dumps({'revision':doc['revision'],'issue_id':issue['issue_id'],'decision':'approve'}).encode())
        self.assertEqual(status,200,body);revision=json.loads(body)['review']['revision']
        status,body=self.request('POST',f'/api/jobs/{identifier}/review/apply',body=json.dumps({'revision':revision}).encode())
        self.assertEqual(status,200,body)
        result=json.loads(body)
        self.assertTrue(Path(result['corrected_path']).is_file())
        self.assertEqual(json.loads(Path(result['corrected_path']).read_text())['note'],'unchanged')
        self.app.save_history()
        restored=web.App(self.app.output_dir,self.app.work_dir)
        try:
            self.assertEqual(restored.jobs[identifier].workflow_version,'review-first-v1')
            self.assertEqual(restored.review(identifier)['summary']['decisions']['approve'],1)
        finally:restored.close()

    def test_review_endpoints_require_token_and_reject_arbitrary_fields(self):
        identifier=self.job()
        for method,path in [('GET',f'/api/jobs/{identifier}/review'),('POST',f'/api/jobs/{identifier}/review/apply')]:
            status,_=self.request(method,path,token=False);self.assertEqual(status,403)
        status,body=self.request('POST',f'/api/jobs/{identifier}/review',body=json.dumps({'revision':0,'field_id':'fake','final_text':'别的字段。'}).encode())
        self.assertEqual(status,400,body)
        status,body=self.request('POST',f'/api/jobs/{identifier}/review/apply',body=json.dumps({'revision':0}).encode())
        self.assertEqual(status,200,body)
        result=json.loads(body);self.assertEqual(result['changed_captions'],0)
        self.assertEqual(Path(result['corrected_path']).read_bytes(),self.app.jobs[identifier].source.read_bytes())

    def test_repeat_export_reuses_valid_file_but_new_review_or_edited_output_creates_fresh_copy(self):
        identifier=self.job();document=self.app.review(identifier)['review'];issue=document['issues'][0]
        self.app.review(identifier,{'revision':0,'issue_id':issue['issue_id'],'decision':'approve'})
        first=self.app.apply_review(identifier,1);again=self.app.apply_review(identifier,1)
        self.assertEqual(first['corrected_path'],again['corrected_path']);self.assertTrue(again['reused'])
        self.app.review(identifier,{'revision':1,'issue_id':issue['issue_id'],'decision':'keep'})
        second=self.app.apply_review(identifier,2)
        self.assertNotEqual(first['corrected_path'],second['corrected_path'])
        self.assertEqual(Path(second['corrected_path']).read_bytes(),self.app.jobs[identifier].source.read_bytes())
        Path(second['corrected_path']).write_text('externally edited')
        third=self.app.apply_review(identifier,2)
        self.assertNotEqual(second['corrected_path'],third['corrected_path'])
        self.assertEqual(Path(third['corrected_path']).read_bytes(),self.app.jobs[identifier].source.read_bytes())

    def test_auto_copy_is_not_reused_as_a_human_reviewed_export(self):
        data=(json.dumps({'id':'001','caption':GOOD.replace('身穿白袍。','身穿白袍')},ensure_ascii=False)+'\n').encode()
        item=self.app.create_job('automatic.jsonl',io.BytesIO(data),len(data),dry_run=True)
        self.app.executor.submit(lambda:None).result(timeout=10)
        before=self.app.jobs[item['id']].id_summary['review_output']
        self.assertEqual(before['export_kind'],'automatic')
        doc=self.app.review(item['id'])['review'];after=self.app.apply_review(item['id'],doc['revision'])
        self.assertEqual(after['export_kind'],'reviewed')
        self.assertNotEqual(before['corrected_path'],after['corrected_path'])

    def test_model_grammar_requires_human_text_not_auto_approval(self):
        original='0-4秒，第0秒，<主体1>站立着进行。'
        raw=GOOD.replace('0-4秒，第0秒，<主体1>向前跑。',original)
        fake=FakeAsyncClient(semantic_reply(issues=[dict(semantic_issue('自然指令','站立着进行','句式不完整','语法病句与语句不通顺'),can_auto_fix=False)],need_human_review=True,human_review_reason='语句成分不完整'))
        real=IdentificationAgent
        data=(json.dumps({'id':'001','caption':raw},ensure_ascii=False)+'\n').encode()
        with patch.object(web,'read_saved_key',return_value='tp-test'),patch('caption_qa_review.IdentificationAgent',side_effect=lambda **kwargs:real(async_client=fake,**kwargs)):
            item=self.app.create_job('grammar.jsonl',io.BytesIO(data),len(data),dry_run=False,auto_fix=False)
            self.app.executor.submit(lambda:None).result(timeout=10)
        self.assertEqual(len(fake.calls),1)
        document=self.app.review(item['id'])['review'];issue=document['issues'][0]
        self.assertFalse(issue['operations'])
        with self.assertRaisesRegex(ValueError,'人工填写'):
            self.app.review(item['id'],{'revision':0,'issue_id':issue['issue_id'],'decision':'approve'})
        field=document['fields'][0]
        self.app.review(item['id'],{'revision':0,'field_id':field['field_id'],'final_text':'0-4秒，第0秒，<主体1>站立。'})
        result=self.app.apply_review(item['id'],1)
        self.assertIn('<主体1>站立。',Path(result['corrected_path']).read_text())


if __name__=='__main__':unittest.main()
