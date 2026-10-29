"""Yellow Caption cells, preserving XLSX members and exporting literal CSV views."""
from pathlib import Path
import csv
import json
import re
import xml.etree.ElementTree as ET

YELLOW='FFFFFF00'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
XML_CONTROLS=re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]')
EXCEL_ESCAPE=re.compile(r'_x[0-9a-fA-F]{4}_')


class ExcelViewError(ValueError):
    """A companion cannot represent the data; the original-format copy is valid."""


def _attr(attrs,name,value):
    pattern=r'\s+'+re.escape(name)+r'=["\'][^"\']*["\']'
    replacement=f' {name}="{value}"'
    return re.sub(pattern,replacement,attrs) if re.search(pattern,attrs) else attrs+replacement


class YellowStyles:
    """Append fills/XFs to the original XML; do not reserialize existing styles."""
    def __init__(self,raw):
        self.raw=raw.decode('utf-8');tree=ET.fromstring(raw);tag=lambda n:'{'+NS+'}'+n
        fills=tree.find(tag('fills'));xfs=tree.find(tag('cellXfs'))
        if fills is None or xfs is None:raise ValueError('原工作簿缺少可安全扩展的样式表，无法标黄。')
        self.fill_id=len(fills);self.add_fill=True
        for n,fill in enumerate(fills):
            pattern=fill.find(tag('patternFill'));color=None if pattern is None else pattern.find(tag('fgColor'))
            if pattern is not None and pattern.get('patternType')=='solid' and color is not None and color.get('rgb','').upper()==YELLOW:
                self.fill_id=n;self.add_fill=False;break
        self.fill_count=len(fills);self.xf_count=len(xfs);self.mapping={};self.added=[]
        root=re.search(r'<(?P<prefix>[A-Za-z_][\w.-]*:)?styleSheet\b',self.raw)
        if not root:raise ValueError('无法确认工作簿样式命名空间。')
        self.prefix=root['prefix'] or ''
        block=self._block('cellXfs')
        p=re.escape(self.prefix)
        self.original_xfs=re.findall(r'<'+p+r'xf\b[^>]*?(?:/>|>.*?</'+p+r'xf>)',block['body'],flags=re.S)
        if len(self.original_xfs)!=self.xf_count:raise ValueError('工作簿样式结构不明确，未标黄。')

    def _block(self,name):
        p=re.escape(self.prefix)
        match=re.search(r'<'+p+name+r'\b(?P<attrs>[^>]*?)>(?P<body>.*?)</'+p+name+r'>',self.raw,re.S)
        if not match:raise ValueError('无法安全定位工作簿样式节点。')
        return match

    def style_id(self,old):
        if not old.isdigit() or not 0<=int(old)<self.xf_count:raise ValueError('Caption原样式索引无效，未标黄。')
        if old not in self.mapping:
            raw=self.original_xfs[int(old)];end=raw.find('>');opening=raw[:end]
            tail='/' if opening.endswith('/') else ''
            if tail:opening=opening[:-1]
            opening=_attr(_attr(opening,'fillId',self.fill_id),'applyFill',1)
            self.mapping[old]=self.xf_count+len(self.added);self.added.append(opening+tail+raw[end:])
        return str(self.mapping[old])

    def render(self):
        if not self.added:return self.raw.encode('utf-8')
        if self.add_fill:
            block=self._block('fills');p=self.prefix
            extra=f'<{p}fill><{p}patternFill patternType="solid"><{p}fgColor rgb="{YELLOW}"/><{p}bgColor indexed="64"/></{p}patternFill></{p}fill>'
            self.raw=self.raw[:block.start()]+f'<{p}fills'+_attr(block['attrs'],'count',self.fill_count+1)+'>'+block['body']+extra+f'</{p}fills>'+self.raw[block.end():]
        block=self._block('cellXfs');p=self.prefix
        self.raw=self.raw[:block.start()]+f'<{p}cellXfs'+_attr(block['attrs'],'count',self.xf_count+len(self.added))+'>'+block['body']+''.join(self.added)+f'</{p}cellXfs>'+self.raw[block.end():]
        return self.raw.encode('utf-8')


def xlsx_cells(document,records,source,input_source,marked):
    from openpyxl.utils import get_column_letter
    from caption_qa_table_io import _schema,_openpyxl,iter_table_records
    source=Path(source);result={}
    if Path(input_source)!=source:
        from caption_qa_split_input import FIELDS
        for index in marked:
            record=records[index];sheet=record.get('source_sheet');number=record.get('source_row')
            if not sheet or type(number) is not int:raise ValueError('分列Caption的标黄位置缺少来源映射。')
            schema=_schema(source,sheet)
            original=next((r for i,r,error in iter_table_records(source,sheet_name=sheet) if i==number-1 and not error),None)
            if original is None:raise ValueError('分列Caption来源行不存在，未标黄。')
            for field in FIELDS:
                key=field+'v2'
                if key not in schema.keys:raise ValueError('分列Caption字段映射不完整，未标黄。')
                if original.get(key) not in (None,''):
                    result.setdefault(sheet,set()).add(get_column_letter(schema.keys.index(key)+1)+str(number))
    else:
        sheet=document['sheet'];schema=_schema(source,sheet)
        if not sheet:
            book=_openpyxl().load_workbook(source,read_only=True)
            try:sheet=book.sheetnames[0]
            finally:book.close()
        column=get_column_letter(schema.keys.index(document['caption_column'])+1)
        result[sheet]={column+str(index+1) for index in marked}
    return result


def excel_view(source,target,document,marked):
    """A companion only; the original-format copy stays authoritative."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    from caption_qa_table_io import _schema,iter_table_records
    source=Path(source);wb=Workbook();ws=wb.active;ws.title='数据'
    if source.suffix.lower()=='.csv':
        with source.open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.reader(stream))
        column=_schema(source,document['sheet']).keys.index(document['caption_column'])+1
        marked_excel={index+1 for index in marked}
    else:
        records=list(iter_table_records(source));headers=list(_schema(source,document['sheet']).keys)
        rows=[headers];marked_excel=set();column=headers.index(document['caption_column'])+1
        for index,record,error in records:
            if error:raise ValueError('记录无法解析，无法生成标黄Excel副本。')
            rows.append([json.dumps(record.get(h),ensure_ascii=False) if isinstance(record.get(h),(dict,list)) else '' if record.get(h) is None else str(record[h]) for h in headers])
            if index in marked:marked_excel.add(len(rows))
    encoded={};escaped_cells=[];patched=Path(target).with_suffix('.encoded.xlsx')
    try:
        for number,row in enumerate(rows,1):
            values=[]
            for col,value in enumerate(row,1):
                address=get_column_letter(col)+str(number)
                if len(value)>32767:
                    header=rows[0][col-1] if col<=len(rows[0]) else ''
                    raise ExcelViewError(f'第 {number} 行 {address}（{header}）有 {len(value)} 个字符，超过Excel单元格32767字符上限；未截断。')
                controls=sorted({f'U+{ord(ch):04X}' for ch in XML_CONTROLS.findall(value)})
                if controls or EXCEL_ESCAPE.search(value):
                    # Write legal placeholder XML, then use the existing OOXML
                    # text encoder. Preserve controls as _xHHHH_ and protect
                    # literal escape-looking strings as _x005F_xHHHH_.
                    encoded[address]=value;values.append('Excel encoded text')
                    if controls:escaped_cells.append({'cell':address,'column':rows[0][col-1] if col<=len(rows[0]) else '', 'controls':controls})
                else:values.append(value)
            ws.append(values)
        for row in ws:
            for cell in row:
                cell.data_type='s';cell.number_format='@';cell.alignment=Alignment(vertical='top',wrap_text=cell.column==column)
        for cell in ws[1]:cell.font=Font(bold=True)
        for number in marked_excel:ws.cell(number,column).fill=PatternFill('solid',fgColor=YELLOW)
        for n in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(n)].width=90 if n==column else 22
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        wb.save(target)
        if encoded:
            from caption_qa_review_io import patch_xlsx
            patch_xlsx(target,patched,{ws.title:encoded})
            patched.replace(target)
        return {'escaped_cells':escaped_cells}
    finally:
        wb.close();patched.unlink(missing_ok=True)
