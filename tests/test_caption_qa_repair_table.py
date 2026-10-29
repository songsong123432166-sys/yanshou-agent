"""Repair handoff preserves ID identity and original-source finding coverage."""
import asyncio
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

import captionqa.caption_qa_agent as qa
from captionqa.caption_qa_repair_table import export_repair_table, read_repair_entries
from tests.test_caption_qa_upgrades import GOOD


class RepairTableTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.jsonl'
        self.cache = self.root / 'cache.jsonl'

    def run_records(self, records):
        self.source.write_text(''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in records))
        asyncio.run(qa.aprocess_batch(self.source, self.cache, qa.CaptionQAAgent(dry_run=True)))
        return self.entries()

    def entries(self):
        return read_repair_entries(self.cache, id_column='id', dry_run=True)

    def test_one_id_one_row_includes_source_format_errors_without_revised_body(self):
        raw = GOOD.replace('身穿白袍。', '这里是不要导出的正文').replace('向前跑。', '向前跑')
        entries, missing = self.run_records([{'id': '00001', 'caption': raw}, {'id': 'pass', 'caption': GOOD},
                                              {'id': 'skip', 'caption': ''}])
        self.assertEqual(missing, [])
        self.assertEqual([r[0] for r in entries], ['00001'])
        self.assertIn('【主体】第1段缺少句末标点', entries[0][1])
        self.assertIn('【自然指令】第1段缺少句末标点', entries[0][1])
        self.assertNotIn('不要导出的正文', entries[0][1])

    def test_latest_retry_replaces_old_errors_not_appends_and_disabled_items_ignored(self):
        self.run_records([{'id': '0001', 'caption': GOOD}])
        row = json.loads(self.cache.read_text())
        row['qa_report']['issues'] = [{'source': 'static', 'issue_type': 'missing_punctuation',
            'section': '主体', 'original_snippet': '旧文本', 'resolved': True}]
        old = json.dumps(row) + '\n'
        row['qa_report']['issues'] = [{'source': 'llm', 'issue_type': '衍字/叠词冗余'}]
        self.cache.write_text(old + json.dumps(row) + '\n')
        self.assertEqual(self.entries(), ([], []))

    def test_tag_format_note_names_exact_defect_and_tag(self):
        raw = GOOD.replace('<主体1>：', '< 主体1 >：').replace('<主体1>向前跑', '主体1向前跑')
        entries, _ = self.run_records([{'id': 'tag-example', 'caption': raw}])
        self.assertIn('【主体】标签内部多余空格（主体1）', entries[0][1])
        self.assertIn('【自然指令】编号标签缺少尖括号（主体1）', entries[0][1])

    def test_missing_ids_are_not_replaced_with_row_numbers_and_resolution_is_hash_bound(self):
        self.run_records([{'id': '', 'caption': GOOD.replace('身穿白袍。', '身穿白袍')}])
        self.assertEqual(self.entries(), ([], [2]))
        row = json.loads(self.cache.read_text())
        sidecar = self.cache.with_suffix('.id_resolutions.json')
        document = {'source_sha256': row['qa_checkpoint']['source_sha256'], 'rows': {'1': {'id': '00009'}}}
        sidecar.write_text(json.dumps(document))
        self.assertEqual(self.entries()[0][0][0], '00009')
        document['source_sha256'] = 'different-source'
        sidecar.write_text(json.dumps(document))
        with self.assertRaisesRegex(ValueError, '不匹配'):
            self.entries()

    def test_baseline_rejects_semantic_result_instead_of_mislabeling_it(self):
        self.run_records([{'id': '1', 'caption': GOOD}])
        row = json.loads(self.cache.read_text())
        row['qa_scope'] = 'static_and_semantic'
        self.cache.write_text(json.dumps(row) + '\n')
        with self.assertRaisesRegex(ValueError, '不能混入'):
            self.entries()

    def test_saved_xlsx_has_exact_two_columns_literal_ids_filter_and_frozen_header(self):
        bad = GOOD.replace('身穿白袍。', '身穿白袍')
        self.run_records([{'id': '00001', 'caption': bad}, {'id': '=1+1', 'caption': bad},
                          {'id': 'ordinary', 'caption': GOOD.replace('<主体1>向前跑。', '男子向前跑。')}])
        before = self.cache.read_bytes(), self.source.read_bytes()
        info = export_repair_table(self.cache, self.root / 'ids.txt', id_column='id', dry_run=True, complete=True)
        book = load_workbook(info['repair_report_path'])
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet.title, '底线返修')
        self.assertEqual((sheet.max_row, sheet.max_column), (4, 2))
        self.assertEqual([c.value for c in sheet[1]], ['题目 ID', '错误项'])
        self.assertEqual(sheet['A2'].value, '00001')
        self.assertEqual(sheet['A3'].value, '=1+1')
        self.assertNotEqual(sheet['A3'].data_type, 'f')
        self.assertIn('<主体1>', sheet['B4'].value)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertTrue(sheet['B2'].alignment.wrap_text)
        self.assertTrue(sheet.tables['CaptionRepair'].autoFilter)
        self.assertEqual((self.cache.read_bytes(), self.source.read_bytes()), before)


if __name__ == '__main__':
    unittest.main()
