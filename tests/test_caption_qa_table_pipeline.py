"""Full-table acceptance checks; uses static QA and never contacts an API."""
from __future__ import annotations

import asyncio
import csv
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import captionqa.caption_qa_agent as qa
from captionqa.caption_qa_table_io import QA_COLUMNS


HUMAN = "## 任务类型\n人物参考\n## 主体\n<主体1>：白袍女子。\n## 自然指令\n<主体1>向前奔跑。\n"


class TablePipelineTests(unittest.TestCase):
    def test_explicit_cli_table_export_remains_compatible(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder)/"business.csv"
            originals = [
                ["001", "审核甲", HUMAN, ""],
                ["002", "审核乙", "ordinary machine caption" * 3, ""],
                ["003", "", "", HUMAN.replace("奔跑。", "奔跑")],
                ["004", "", "", ""],
            ]
            with source.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["video_id", "审核员", "人标", "备用"])
                writer.writerows(originals)
            result = subprocess.run([sys.executable, str(Path(qa.__file__).resolve().parents[1]/'agent.py'), 'check', "--dry-run", "--input", str(source),
                                     "--output", str(source.with_name("business_qa_result.csv"))],
                                    capture_output=True, encoding="utf-8", timeout=30)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("[Auto-Detect]", result.stdout)
            target, cache = source.with_name("business_qa_result.csv"), source.with_name("business_qa_result.jsonl")
            self.assertTrue(cache.exists())
            with target.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))
            self.assertEqual(rows[0], ["video_id", "审核员", "人标", "备用"] + QA_COLUMNS)
            self.assertEqual([row[:4] for row in rows[1:]], originals)
            self.assertEqual([row[4] for row in rows[1:]],
                             ["PASS", "SKIPPED_NOT_HUMAN_CAPTION", "SKIPPED_NOT_HUMAN_CAPTION", "SKIPPED_NOT_HUMAN_CAPTION"])
            self.assertEqual(len([json.loads(line) for line in cache.read_text(encoding="utf-8").splitlines()]), 4)

    def test_native_excel_export_preserves_dates_formulas_qa_collisions_and_other_sheets(self) -> None:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font
        with tempfile.TemporaryDirectory() as folder:
            source, target = Path(folder)/"input.xlsx", Path(folder)/"result.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "业务"
            sheet.append(["video_id", "日期", "公式", "qa_status", "人标"])
            sheet.append(["000001", datetime(2026, 10, 22), "=1+2", "old-status", HUMAN])
            sheet.append(["000002", None, "=3+4", "old-status-2", "machine text"])
            sheet["A2"].font = Font(bold=True)
            other = book.create_sheet("其他工作表")
            other["A1"] = "原始备注"
            book.save(source)
            book.close()
            self.assertEqual(qa.process_batch(source, target, qa.CaptionQAAgent(dry_run=True)), 2)
            restored = load_workbook(target, data_only=False)
            self.addCleanup(restored.close)
            self.assertEqual(restored.sheetnames, ["业务", "其他工作表"])
            out = restored["业务"]
            self.assertEqual(out["A2"].value, "000001")
            self.assertEqual(out["B2"].value, datetime(2026, 10, 22))
            self.assertEqual(out["C2"].value, "=1+2")
            self.assertEqual(out["C2"].data_type, "f")
            self.assertEqual(out["D2"].value, "old-status")
            self.assertTrue(out["A2"].font.bold)
            self.assertEqual([c.value for c in out[1]][-5:], QA_COLUMNS)
            self.assertEqual(out.cell(2, 6).value, "PASS")
            self.assertEqual(out.cell(3, 6).value, "SKIPPED_NOT_HUMAN_CAPTION")
            self.assertEqual(restored["其他工作表"]["A1"].value, "原始备注")
            self.assertEqual(qa.process_batch(source, target, qa.CaptionQAAgent(dry_run=True)), 0)

    def test_no_human_file_exits_before_static_and_export(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            source, target = Path(folder)/"input.csv", Path(folder)/"result.csv"
            source.write_text("video_id,机标\n001,ordinary machine text\n", encoding="utf-8-sig")
            with patch.object(qa, "run_static_rules", side_effect=AssertionError("must not run static")):
                with self.assertRaisesRegex(ValueError, "人标"):
                    qa.process_batch(source, target, qa.CaptionQAAgent(dry_run=True))
            self.assertFalse(target.exists())
            self.assertFalse(target.with_suffix(".jsonl").exists())

    def test_explicit_column_can_mark_entire_table_skipped_without_static(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            source, target = Path(folder)/"input.csv", Path(folder)/"result.csv"
            source.write_text("id,caption\n1,plain text\n2,\n", encoding="utf-8-sig")
            with patch.object(qa, "run_static_rules", side_effect=AssertionError("must not run static")):
                self.assertEqual(qa.process_batch(source, target, qa.CaptionQAAgent(dry_run=True), column="caption"), 2)
            with target.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))
            self.assertTrue(all(row[-5] == "SKIPPED_NOT_HUMAN_CAPTION" for row in rows[1:]))


if __name__ == "__main__":
    unittest.main()
