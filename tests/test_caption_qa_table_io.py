"""Offline source-table fidelity and QA export acceptance tests."""
from __future__ import annotations

import csv
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from typing import Any

import captionqa.caption_qa_table_io as table_io


class TableIOTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def csv_source(self, rows: list[list[Any]]) -> Path:
        path = self.root / "input.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(rows)
        return path

    def test_csv_preserves_duplicate_empty_headers_short_rows_and_extra_cells(self) -> None:
        path = self.csv_source([["id", "caption", "caption", ""],
                                ["001", "first", "second", "tail"], [],
                                ["002", "short"], ["003", "a", "b", "c", "extra"]])
        rows = list(table_io.iter_table_records(path))
        self.assertEqual([index for index, _, _ in rows], [1, 2, 3, 4])
        self.assertEqual(list(rows[0][1].values()), ["001", "first", "second", "tail", ""])
        self.assertEqual(len(set(rows[0][1])), 5)
        self.assertEqual(list(rows[1][1].values()), [""] * 5)
        self.assertEqual(list(rows[2][1].values()), ["002", "short", "", "", ""])
        self.assertEqual(list(rows[3][1].values())[-1], "extra")
        output = self.root / "output.csv"
        table_io.export_qa_table(path, output, {4: {"qa_status": "PASS"}})
        with output.open(encoding="utf-8-sig", newline="") as stream:
            exported = list(csv.reader(stream))
        self.assertEqual(exported[0][:4], ["id", "caption", "caption", ""])
        self.assertEqual(exported[0][-5:], table_io.QA_COLUMNS)
        self.assertEqual(exported[4][-5], "PASS")
        self.assertEqual(exported[4][4], "extra")

    def test_qa_header_collisions_retain_original_columns_then_append_new_qa_columns(self) -> None:
        source = self.csv_source([["qa_status", "id"], ["original", "0007"]])
        output = self.root / "qa.csv"
        table_io.export_qa_table(source, output, {1: {"qa_status": "AUTO_FIXED"}})
        with output.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[0], ["qa_status", "id"] + table_io.QA_COLUMNS)
        self.assertEqual(rows[1][:3], ["original", "0007", "AUTO_FIXED"])

    def test_jsonl_physical_indexes_include_blank_malformed_and_nonobject_rows(self) -> None:
        path = self.root / "input.jsonl"
        path.write_text('{"id":"0001","caption":"line1\\nline2"}\n\n{bad\n["raw"]\n'
                        '{"id":"0005","later":{"nested":true}}\n', encoding="utf-8")
        rows = list(table_io.iter_table_records(path))
        self.assertEqual([row[0] for row in rows], [1, 2, 3, 4, 5])
        self.assertEqual(rows[1][1], {})
        self.assertIsNone(rows[1][2])
        self.assertTrue(rows[2][2])
        self.assertIn("{bad", rows[2][1].values())
        self.assertTrue(rows[3][2])
        output = self.root / "jsonl.csv"
        table_io.export_qa_table(path, output, {5: {"qa_status": "PASS"}})
        with output.open(encoding="utf-8-sig", newline="") as stream:
            exported = list(csv.DictReader(stream))
        self.assertEqual(exported[0]["id"], "0001")
        self.assertEqual(json.loads(exported[4]["later"]), {"nested": True})
        self.assertEqual(exported[4]["qa_status"], "PASS")

    def test_same_input_output_is_rejected_without_changing_source(self) -> None:
        path = self.csv_source([["id"], ["001"]])
        original = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "(?i)(same|相同|覆盖)"):
            table_io.export_qa_table(path, path, {})
        self.assertEqual(path.read_bytes(), original)

    def test_csv_multiline_and_formula_like_text_remain_exact(self) -> None:
        source = self.csv_source([["id", "caption", "url"], ["0001", "## 主体\n行1\n行2", "=literal"]])
        output = self.root / "out.csv"
        table_io.export_qa_table(source, output, {1: {"qa_revised_caption": "=SUM(A1:A2)\ncaption"}})
        with output.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[1][:3], ["0001", "## 主体\n行1\n行2", "=literal"])
        self.assertEqual(rows[1][-1], "=SUM(A1:A2)\ncaption")

    def test_unsupported_extensions_raise_clear_error(self) -> None:
        path = self.root / "input.txt"
        path.write_text("data", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "(?i)(格式|extension|supported)"):
            list(table_io.iter_table_records(path))

    def test_duplicate_header_unique_keys_do_not_collide_with_existing_headers(self) -> None:
        source = self.csv_source([["id", "id__2", "id", ""], ["a", "b", "c", "d"]])
        row = next(table_io.iter_table_records(source))[1]
        self.assertEqual(len(row), 4)
        self.assertEqual(list(row.values()), ["a", "b", "c", "d"])
        self.assertEqual(table_io.table_columns(source), list(row))

    def test_legacy_json_array_string_and_scalar_inputs_remain_supported(self) -> None:
        source = self.root / "legacy.json"
        source.write_text(json.dumps([{"id": "001", "text": "caption"}, "standalone", None]),
                          encoding="utf-8")
        records = list(table_io.iter_table_records(source))
        self.assertEqual(records, [(1, {"id": "001", "text": "caption"}, None),
                                   (2, {"caption": "standalone"}, None),
                                   (3, {"raw_input": None}, None)])
        output = self.root / "legacy.csv"
        table_io.export_qa_table(source, output, {2: {"qa_status": "PASS"}})
        with output.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(rows[0]["id"], "001")
        self.assertEqual(rows[1]["caption"], "standalone")
        self.assertEqual(rows[1]["qa_status"], "PASS")
        source.write_text('{"id":"object"}', encoding="utf-8")
        self.assertEqual(list(table_io.iter_table_records(source)), [(1, {"id": "object"}, None)])

    def test_long_csv_source_and_qa_cells_are_not_truncated(self) -> None:
        original, revised = "原" * 150000, "订正" * 100000
        source = self.csv_source([["caption"], [original]])
        output = self.root / "long.csv"
        table_io.export_qa_table(source, output, {1: {"qa_revised_caption": revised}})
        with output.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[1][0], original)
        self.assertEqual(rows[1][-1], revised)


@unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is not installed")
class ExcelTableIOTests(TableIOTests):
    def test_excel_to_excel_preserves_formulas_sheets_formatting_and_literal_qa(self) -> None:
        import openpyxl
        from openpyxl.comments import Comment
        from openpyxl.styles import Font

        source, output = self.root / "source.xlsx", self.root / "result.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "业务"
        sheet.append(["id", "caption", "formula"])
        sheet.append(["00001", "人标", "=1+2"])
        sheet.append([None, None, None])
        sheet.append(["00003", "第三行", None])
        sheet["A2"].font = Font(bold=True, color="FF0000")
        sheet["A2"].comment = Comment("note", "author")
        sheet["B2"].hyperlink = "https://example.invalid/caption"
        book.create_sheet("other")["A1"] = "keep me"
        book.save(source)
        original = source.read_bytes()
        rows = list(table_io.iter_table_records(source))
        self.assertEqual([index for index, _, _ in rows], [1, 2, 3])
        self.assertEqual(rows[0][1]["formula"], "=1+2")
        table_io.export_qa_table(source, output, {
            3: {"qa_status": "AUTO_FIXED"}, 1: {"qa_revised_caption": "=literal caption"}})
        self.assertEqual(source.read_bytes(), original)
        restored = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(restored.close)
        self.assertEqual(restored.sheetnames, ["业务", "other"])
        sheet = restored["业务"]
        self.assertEqual(sheet["A2"].value, "00001")
        self.assertEqual(sheet["C2"].value, "=1+2")
        self.assertTrue(sheet["A2"].font.bold)
        self.assertEqual(sheet["A2"].comment.text, "note")
        self.assertEqual(sheet["B2"].hyperlink.target, "https://example.invalid/caption")
        self.assertEqual(sheet["H2"].value, "=literal caption")
        self.assertEqual(sheet["H2"].data_type, "s")
        self.assertEqual(sheet["D4"].value, "AUTO_FIXED")
        self.assertEqual(restored["other"]["A1"].value, "keep me")

    def test_csv_to_excel_does_not_coerce_ids_or_formula_like_strings(self) -> None:
        import openpyxl

        source = self.csv_source([["id", "caption"], ["0001", "=literal source"]])
        output = self.root / "converted.xlsx"
        table_io.export_qa_table(source, output, {1: {"qa_status": "PASS"}})
        book = openpyxl.load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.active["A2"].value, "0001")
        self.assertEqual(book.active["B2"].value, "=literal source")
        self.assertEqual(book.active["B2"].data_type, "s")

    def test_excel_overlong_qa_raises_and_leaves_existing_output_untouched(self) -> None:
        source = self.csv_source([["caption"], ["source"]])
        output = self.root / "output.xlsx"
        output.write_bytes(b"old result")
        with self.assertRaisesRegex(ValueError, "(?i)(32767|CSV|JSONL)"):
            table_io.export_qa_table(source, output, {1: {"qa_revised_caption": "长" * 32768}})
        self.assertEqual(output.read_bytes(), b"old result")

    def test_explicit_worksheet_is_selected_and_missing_sheet_has_clear_error(self) -> None:
        import openpyxl

        source = self.root / "input.xlsx"
        book = openpyxl.Workbook()
        book.active.append(["first"])
        sheet = book.create_sheet("target")
        sheet.append(["caption", "id"])
        sheet.append(["正文", "0002"])
        book.save(source)
        rows = list(table_io.iter_table_records(source, sheet_name="target"))
        self.assertEqual(rows[0][1], {"caption": "正文", "id": "0002"})
        with self.assertRaisesRegex(ValueError, "missing"):
            list(table_io.iter_table_records(source, sheet_name="missing"))
        output = self.root / "sheet.csv"
        table_io.export_qa_table(source, output, {}, sheet_name="target")
        with output.open(encoding="utf-8-sig", newline="") as stream:
            rows_out = list(csv.reader(stream))
        self.assertEqual(rows_out[1][:2], ["正文", "0002"])

    def test_jsonl_excel_export_preserves_large_integer_identifiers(self) -> None:
        import openpyxl

        source, output = self.root / "large_id.jsonl", self.root / "large_id.xlsx"
        source.write_text('{"id":123456789012345678901,"caption":"正文"}\n', encoding="utf-8")
        table_io.export_qa_table(source, output, {})
        book = openpyxl.load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.active["A2"].value, "123456789012345678901")

    def test_native_excel_rich_text_runs_are_retained(self) -> None:
        import openpyxl
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont

        source, output = self.root / "rich.xlsx", self.root / "rich_qa.xlsx"
        book = openpyxl.Workbook()
        book.active.append(["caption"])
        book.active["A2"] = CellRichText("plain", TextBlock(InlineFont(b=True), "bold"))
        book.save(source)
        table_io.export_qa_table(source, output, {1: {"qa_status": "PASS"}})
        restored = openpyxl.load_workbook(output, rich_text=True)
        self.addCleanup(restored.close)
        self.assertIsInstance(restored.active["A2"].value, CellRichText)
        self.assertTrue(restored.active["A2"].value[1].font.b)


if __name__ == "__main__":
    unittest.main()
