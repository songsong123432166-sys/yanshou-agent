"""Field migration tests use source-preserving plans and no real API/key."""
import asyncio
import copy
import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from captionqa.caption_field_schema import VERSION, FIELDS, parse, render, normalize, apply_plan, task_value, is_legacy_reference
from captionqa.caption_field_jobs import FieldJobs, detect_column, select_id_column, write_copy, missing_fields
from tests.test_caption_qa_upgrades import sdk_response

MAPPING={'参考 / 剪辑/卡点迁移':['运镜参考']}
OLD='''## task_type
参考
## subtask_type
剪辑/卡点迁移
## similarity_label
Yes
## similarity_tag
镜头/运镜-相同|实体-相同
## audio_similarity_tag
台词内容-相同|环境音-相同
## brief_captions
- 图片1: 白色床单上的青年。
- 视频1: 镜头从电脑摇向玩偶。
## subject
- 主体1: 青年男子。
- 主体2: 玩偶。
## natural_instruction
参考<视频1>的运镜，生成一段互动视频。

室内环境柔和明亮，整体色调温馨。

0-1秒，<持续的室内底噪>镜头俯拍电脑，屏幕显示【你好】，画外传来台词{Now, hello.}。
第2秒，<主体1>抚摸<主体2>。
'''
STANDARD=render({name:('人物参考' if name=='任务类型' else 'Yes' if name=='similarity_label' else
                       '<主体1>：青年。' if name=='主体' else '<主体1>站立。' if name=='自然指令' else '') for name in FIELDS})


class PlanClient:
    def __init__(self,mutate=None):self.chat=self;self.completions=self;self.calls=[];self.closed=False;self.mutate=mutate
    async def create(self,**kw):
        self.calls.append(kw);blocks=json.loads(kw['messages'][1]['content'])['blocks'];assignments=[]
        for block in blocks:
            field=block['field']
            if block['alias']=='natural_instruction':
                field='提示' if block['text'].startswith('参考') else '场景' if block['text'].startswith('室内') else '自然指令'
            assignments.append({'id':block['id'],'field':field})
        plan={'ready':True,'assignments':assignments,'derived':[],'notes':[]}
        if self.mutate:self.mutate(plan,blocks)
        await asyncio.sleep(0)
        return sdk_response(plan)
    async def close(self):self.closed=True


class SchemaTests(unittest.IsolatedAsyncioTestCase):
    async def test_field_blank_lines_are_repaired_locally_and_idempotently(self):
        original=STANDARD.replace('\n\n## ', '\n## ').rstrip('\n')
        client=PlanClient()
        result=await normalize(original,client=client)
        self.assertEqual(result['caption'],STANDARD)
        self.assertEqual(result['tags'],['字段间空行'])
        self.assertEqual(client.calls,[])
        self.assertEqual((await normalize(result['caption']))['status'],'already_standard')

    async def test_field_spacing_preserves_crlf_internal_paragraphs_and_literals(self):
        body='  首句。  \n\n第二句{hello\n## subject\nworld}，画面有【## 提示】。  '
        expected=STANDARD.replace('<主体1>站立。',body).replace('\n','\r\n')
        original=expected.replace('\r\n\r\n## ', '\r\n\r\n \t\r\n## ').rstrip('\r\n')
        result=await normalize(original)
        self.assertEqual(result['caption'],expected)
        self.assertEqual(result['model_requests'],0)

    async def test_old_titles_with_new_child_labels_and_misplaced_prompt_are_repaired_locally(self):
        prompt='以<图片1>的机甲UI界面为基础，去除外层大框UI，生成一段包含多视角镜头切换的科幻动画视频。'
        expected=STANDARD.replace('人物参考','人物参考｜视觉风格参考').replace('## 提示\n\n','## 提示\n'+prompt+'\n\n')
        original=expected.replace('## 任务类型','## task_type').replace('## 提示','## subtask_type').replace(
            '## video_similarity_tag','## similarity_tag').replace('## 输入描述','## brief_captions').replace('## 主体','## subject')
        client=PlanClient();result=await normalize(original,client=client)
        self.assertEqual(result['status'],'converted');self.assertEqual(result['caption'],expected)
        self.assertIn('生成要求归入提示',result['tags']);self.assertEqual(client.calls,[])
        self.assertEqual(missing_fields(original),['提示']);self.assertEqual(missing_fields(result['caption']),[])
        again=await normalize(result['caption'],client=client)
        self.assertEqual(again['status'],'already_standard');self.assertEqual(again['caption'],expected)

    async def test_misplaced_request_preserves_crlf_spaces_audio_and_existing_scene_numbers(self):
        prompt='  以<视频1>为参考，生成视频，保留台词{Now, hello.}与文字【生成## 提示】。  '
        expected=STANDARD.replace('## 提示\n\n','## 提示\n'+prompt+'\n\n').replace(
            '## 场景\n\n','## 场景\n<场景2>：原来的编号。\n\n').replace('\n','\r\n')
        original=expected.replace('## 任务类型','task_type').replace('## 提示\r\n','subtask_type\r\n')
        result=await normalize(original)
        self.assertEqual(result['caption'],expected);self.assertEqual(result['model_requests'],0)

    async def test_known_task_taxonomy_under_old_title_does_not_need_a_new_mapping(self):
        from captionqa.caption_field_schema import TASK_LABELS
        self.assertEqual(len(TASK_LABELS),38)
        for label in TASK_LABELS:
            expected=STANDARD.replace('人物参考',label)
            result=await normalize(expected.replace('## 任务类型','## task_type'))
            self.assertEqual(result['caption'],expected,label);self.assertEqual(result['status'],'converted')

    async def test_unknown_tasks_roots_and_real_old_subtypes_do_not_get_guessed_or_partly_changed(self):
        prompt='以<图片1>为参考，生成一段短片。'
        raw=STANDARD.replace('## 任务类型','## task_type').replace('## 提示\n\n','## subtask_type\n'+prompt+'\n\n')
        client=PlanClient()
        for text in [raw.replace('人物参考','参考'),raw.replace('人物参考','人物参考｜未批准任务'),
                     raw.replace(prompt,'剪辑/卡点迁移')]:
            result=await normalize(text,client=client)
            self.assertEqual(result['status'],'needs_confirmation');self.assertEqual(result['caption'],text)
        self.assertEqual(client.calls,[])

    async def test_current_prompt_takes_precedence_over_misplaced_old_prompt_without_merge(self):
        raw=STANDARD.replace('## 任务类型','## task_type').replace('## 提示\n\n','## 提示\n生成室内视频。\n\n')
        raw+='\n## subtask_type\n以<图片1>为参考，生成室外视频。\n'
        client=PlanClient();result=await normalize(raw,client=client)
        self.assertEqual(result['status'],'converted')
        self.assertIn('生成室内视频。',result['caption']);self.assertNotIn('生成室外视频。',result['caption'])
        self.assertEqual(result['original'],raw);self.assertIn('新版字段优先',result['tags'])
        self.assertTrue(result['notes']);self.assertEqual(client.calls,[])

    async def test_ai_cannot_drop_a_request_recovered_from_old_metadata(self):
        prompt='以<图片1>为基础，生成短片。'
        raw=STANDARD.replace('## 任务类型','## task_type').replace('## 提示\n\n','## subtask_type\n'+prompt+'\n\n').replace('## 场景\n\n','')
        def drop_prompt(plan,blocks):
            identifiers={b['id'] for b in blocks if b['alias']=='subtask_type'}
            plan['assignments']=[a for a in plan['assignments'] if a['id'] not in identifiers]
        client=PlanClient(drop_prompt);result=await normalize(raw,client=client)
        self.assertEqual(result['status'],'needs_confirmation');self.assertEqual(result['caption'],raw)
        self.assertEqual(len(client.calls),1);self.assertIn('遗漏',result['notes'][0])
        result=await normalize(raw,client=PlanClient())
        self.assertEqual(result['status'],'converted');self.assertIn(prompt,result['caption'])

    async def test_natural_description_alias_changes_only_title_and_is_not_a_missing_field(self):
        raw=STANDARD.replace('## 自然指令','## 自然描述')
        client=PlanClient();result=await normalize(raw,client=client)
        self.assertEqual(result['caption'],STANDARD);self.assertEqual(result['status'],'converted')
        self.assertEqual(client.calls,[]);self.assertEqual(missing_fields(raw),[])

    async def test_current_visual_header_and_metadata_are_preserved(self):
        text=STANDARD.replace('人物参考','人物参考|场景参考').replace('Yes','yes').replace(
            '## video_similarity_tag\n','## video_similarity_tag\n实体-相同|场景-相同\n').replace(
            '## audio_similarity_tag\n','## audio_similarity_tag\n无\n')
        result=await normalize(text,client=PlanClient())
        self.assertEqual(result['status'],'already_standard')
        self.assertEqual(result['caption'],text);self.assertEqual(result['model_requests'],0)
        self.assertNotIn('## similarity_tag\n',result['caption'])

    async def test_missing_space_after_hashes_is_fixed_but_body_and_header_trailing_spaces_are_preserved(self):
        body='  0-6秒，<主体1>念道{腰下剑...}。第1秒，镜头前推。  \n\n第3秒，画面有【## 提示】。\n第4秒，<雷声>响起。  '
        original=STANDARD.replace('<主体1>站立。',body).replace('## 任务类型','##任务类型',1)
        expected=original.replace('##任务类型','## 任务类型',1)
        original=original.replace('## video_similarity_tag\n','## video_similarity_tag  \n')
        expected=expected.replace('## video_similarity_tag\n','## video_similarity_tag  \n')
        result=await normalize(original)
        self.assertEqual(result['status'],'converted');self.assertEqual(result['caption'],expected)
        self.assertEqual(result['tags'],['标题规范']);self.assertEqual(result['model_requests'],0)

    async def test_exactly_one_ascii_space_is_required_after_hashes(self):
        for prefix in ['##','##  ','##\t','##\u3000']:
            original=STANDARD.replace('## 任务类型',prefix+'任务类型',1)
            result=await normalize(original)
            self.assertEqual(result['status'],'converted');self.assertEqual(result['caption'],STANDARD)
            self.assertEqual(result['model_requests'],0)

    async def test_scene_definitions_and_grid_task_value_are_preserved_during_header_repair(self):
        scenes='<场景1>：云中看台。\n<场景2>：云朵前的看台。\n<场景2>：浅蓝色大厅。'
        expected=STANDARD.replace('人物参考','九宫格/少宫格').replace('## 场景\n','## 场景\n'+scenes+'\n').replace(
            '<主体1>站立。','镜头切换，全景拍摄<场景3>。')
        client=PlanClient()
        result=await normalize(expected.replace('## 场景','##场景'),client=client)
        self.assertEqual(result['caption'],expected);self.assertEqual(result['status'],'converted')
        self.assertEqual(client.calls,[]);self.assertIn(scenes,result['caption'])

    async def test_picture_elements_is_required_and_missing_it_is_reported_after_ai_adds_empty_heading(self):
        original=STANDARD.replace('## 画面元素\n\n','')
        self.assertEqual(missing_fields(original),['画面元素'])
        result=await normalize(original,client=PlanClient())
        self.assertEqual(result['status'],'converted');self.assertIn('## 画面元素\n\n',result['caption'])
        self.assertNotIn('## 画面元素\n无',result['caption'])

    async def test_missing_hashes_do_not_rewrap_chinese_sentences_crlf_or_audio(self):
        text=STANDARD.replace('<主体1>站立。','首句。第二句。  \n\n说话{你好\n## subject\n再见}。').replace('\n','\r\n')
        original=text.replace('## 提示\r\n','提示\r\n')
        result=await normalize(original)
        self.assertEqual(result['caption'],text);self.assertEqual(result['status'],'converted')

    async def test_legacy_visual_alias_changes_only_its_header_line(self):
        expected=STANDARD.replace('<主体1>站立。','  第1秒，镜头左摇。第2秒，切换。\n\n{Hello, world.}。 ')
        original=expected.replace('## video_similarity_tag','## similarity_tag')
        result=await normalize(original)
        self.assertEqual(result['caption'],expected);self.assertEqual(result['status'],'converted')
        self.assertIn('视觉标签字段更名',result['tags'])

    async def test_layout_alone_does_not_make_a_record_need_conversion(self):
        original='\n'+STANDARD.replace('## 自然指令\n','## 自然指令  \n\n').replace('<主体1>站立。','  首句。第二句。  \n\n末句。  ')
        result=await normalize(original)
        self.assertEqual(result['status'],'already_standard');self.assertEqual(result['caption'],original)

    async def test_reordering_keeps_each_field_body_exactly(self):
        pieces=['## '+field+'\n  '+field+'首句。第二句。 \n\n' for field in FIELDS]
        pieces[0]='## 任务类型\n人物参考|场景参考\n\n'
        expected=''.join(pieces);pieces[5],pieces[6]=pieces[6],pieces[5]
        result=await normalize(''.join(pieces))
        self.assertEqual(result['caption'],expected);self.assertIn('字段顺序整理',result['tags'])

    async def test_missing_metadata_is_added_empty_without_reformatting_other_bodies(self):
        original=STANDARD.replace('## similarity_label\nYes\n\n','').replace('<主体1>站立。','首句。第二句。\n\n末句。')
        result=await normalize(original)
        self.assertEqual(result['status'],'converted');self.assertIn('## similarity_label\n\n',result['caption'])
        self.assertEqual(result['caption'].replace('## similarity_label\n\n',''),original)

    async def test_standard_record_is_byte_identical_and_does_not_call_model(self):
        client=PlanClient();result=await normalize(STANDARD,client=client)
        self.assertEqual(result['status'],'already_standard');self.assertEqual(result['caption'],STANDARD);self.assertEqual(client.calls,[])

    async def test_complete_mixed_header_format_is_fixed_locally(self):
        text=STANDARD.replace('## 输入描述','brief_captions：').replace('## 主体','subject').replace('## 自然指令','自然指令:')
        result=await normalize(text)
        self.assertEqual(result['status'],'converted');self.assertEqual([h['field'] for h in parse(result['caption'])['headers']],list(FIELDS))
        self.assertEqual(result['model_requests'],0);self.assertIn('<主体1>站立。',result['caption'])

    async def test_old_format_is_partitioned_without_losing_or_rewriting_words_audio_or_numbers(self):
        client=PlanClient();result=await normalize(OLD,client=client,task_mapping=MAPPING)
        self.assertEqual(result['status'],'converted');self.assertEqual(len(client.calls),1)
        new=parse(result['caption']);self.assertEqual([h['field'] for h in new['headers']],list(FIELDS))
        for snippet in ['<持续的室内底噪>','【你好】','{Now, hello.}','0-1秒','第2秒','<主体1>抚摸<主体2>。','室内环境柔和明亮，整体色调温馨。']:
            self.assertIn(snippet,result['caption'])
        self.assertIn('<图片1>：白色床单上的青年。',result['caption']);self.assertIn('<主体1>：青年男子。',result['caption'])
        self.assertNotIn('全程无声',result['caption']);self.assertNotIn('<主体3>',result['caption'])
        self.assertIn('## 任务类型\n运镜参考',result['caption'])

    async def test_unmapped_old_task_is_preserved_without_calling_api(self):
        client=PlanClient();result=await normalize(OLD,client=client)
        self.assertEqual(result['status'],'needs_confirmation');self.assertEqual(result['caption'],OLD);self.assertEqual(client.calls,[])
        self.assertIn('映射',result['notes'][0])

    async def test_reference_marker_matches_only_old_task_type_exact_value(self):
        for text in ['## task_type\n参考\n## subject\n青年。','task_type：参考\r\nsubject\r\n青年。']:
            self.assertTrue(is_legacy_reference(text))
        for text in [STANDARD,'## task_type\n人物参考','## task_type\n不参考','## subtask_type\n参考','## 自然指令\n参考<视频1>。']:
            self.assertFalse(is_legacy_reference(text))

    async def test_unknown_headers_and_outside_text_are_not_silently_dropped(self):
        for text in ['说明正文\n'+STANDARD,STANDARD+'## unknown_field\n额外资料。']:
            result=await normalize(text,client=PlanClient());self.assertEqual(result['status'],'needs_confirmation');self.assertEqual(result['caption'],text)

    async def test_mixed_new_and_old_task_metadata_cannot_silently_override_each_other(self):
        text=STANDARD+'\n## task_type\n参考\n## subtask_type\n剪辑/卡点迁移\n'
        for mapping in (None,MAPPING):
            result=await normalize(text,client=PlanClient(),task_mapping=mapping)
            self.assertEqual(result['status'],'needs_confirmation');self.assertEqual(result['caption'],text)

    async def test_missing_duplicate_unknown_block_and_illegal_field_plans_fail_closed(self):
        mutations=[lambda p,b:p['assignments'].pop(),lambda p,b:p['assignments'].append(p['assignments'][0]),
                   lambda p,b:p['assignments'][0].update(id=999),lambda p,b:p['assignments'][0].update(field='不存在'),
                   lambda p,b:p['assignments'][0].update(field='主体')]
        for mutate in mutations:
            result=await normalize(OLD,client=PlanClient(mutate),task_mapping=MAPPING)
            self.assertEqual(result['status'],'needs_confirmation');self.assertEqual(result['caption'],OLD)

    async def test_ungrounded_derived_content_is_left_empty_with_report(self):
        def mutate(plan,blocks):plan['derived']=[{'field':'基调','text':'豪华宫廷风格。','evidence':[blocks[0]['id']]}]
        result=await normalize(OLD,client=PlanClient(mutate),task_mapping=MAPPING)
        self.assertEqual(result['status'],'converted');self.assertNotIn('豪华宫廷',result['caption'])
        self.assertTrue(any('缺少直接原文依据' in note for note in result['notes']))

    async def test_grounded_summary_can_be_copied_and_empty_fields_are_not_filled_with_none(self):
        def mutate(plan,blocks):
            block=next(b for b in blocks if b['text'].startswith('参考'))
            plan['derived']=[{'field':'总结','text':block['text'],'evidence':[block['id']]}]
        result=await normalize(OLD,client=PlanClient(mutate),task_mapping=MAPPING)
        self.assertEqual(result['status'],'converted');self.assertIn('## 总结\n参考<视频1>的运镜，生成一段互动视频。',result['caption'])
        self.assertIn('## 画面元素\n\n',result['caption']);self.assertNotIn('## 画面元素\n无',result['caption'])

    async def test_duplicate_nonidentical_content_is_not_overwritten_and_model_can_refuse_conflict(self):
        text=STANDARD+'\n## 主体\n<主体2>：另一人。'
        result=await normalize(text,client=PlanClient())
        self.assertEqual(result['status'],'converted');self.assertIn('<主体1>：青年。',result['caption']);self.assertIn('<主体2>：另一人。',result['caption'])
        result=await normalize(text,client=PlanClient(lambda p,b:p.update(ready=False,notes=['两份主体定义互相矛盾。'])))
        self.assertEqual(result['status'],'needs_confirmation');self.assertEqual(result['caption'],text)

    async def test_title_like_lines_inside_speech_and_screen_text_are_preserved(self):
        text=STANDARD.replace('<主体1>站立。','<主体1>说道{你好\n## subject\n再见}。画面有【## task_type】。')
        result=await normalize(text);self.assertEqual(result['status'],'already_standard');self.assertEqual(result['caption'],text)


class JobTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        self.client=PlanClient();self.manager=FieldJobs(self.root/'jobs',self.root/'out',lambda:self.client);self.addCleanup(self.manager.close)
    def create(self,records,**kw):
        data=''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in records).encode()
        job=self.manager.create('data.jsonl',io.BytesIO(data),len(data),**kw)
        self.manager.executor.submit(lambda:None).result(timeout=10)
        return self.manager.jobs[job['id']],data

    def test_misplaced_prompt_job_writes_fixed_and_missing_id_reports_without_api(self):
        prompt='以<图片1>为参考，生成一段科幻动画视频。'
        expected=STANDARD.replace('人物参考','人物参考｜视觉风格参考').replace('## 提示\n\n','## 提示\n'+prompt+'\n\n')
        raw=expected.replace('## 任务类型','## task_type').replace('## 提示','## subtask_type')
        job,data=self.create([{'id':'cmdcc4c1d5','caption':raw,'备注':'不改','机标caption':OLD}])
        self.assertEqual(job['status'],'done');self.assertEqual(job['counts'],{'converted':1})
        output=json.loads(Path(job['outputs']['corrected']).read_text())
        self.assertEqual(output,{'id':'cmdcc4c1d5','caption':expected,'备注':'不改','机标caption':OLD})
        self.assertEqual(Path(job['outputs']['fixed_ids']).read_text(),'cmdcc4c1d5\n')
        self.assertEqual(Path(job['outputs']['review_ids']).read_text(),'')
        self.assertEqual(Path(job['outputs']['missing_field_ids']).read_text(),'cmdcc4c1d5\n')
        with Path(job['outputs']['report']).open(encoding='utf-8-sig',newline='') as f:report=list(csv.DictReader(f))
        self.assertIn('生成要求归入提示',report[0]['修复标签'])
        self.assertEqual(report[0]['原始缺少字段'],'提示');self.assertEqual(report[0]['整理后仍缺少字段'],'')
        self.assertEqual(self.client.calls,[]);self.assertEqual(job['model_requests'],0)
        self.assertEqual((self.manager.root/job['id']/'input.jsonl').read_bytes(),data)

    def test_job_outputs_real_id_reports_and_changes_only_selected_caption(self):
        rows=[{'id':'0007','caption':OLD,'机标caption':'不得选用','note':'原样保留'},
              {'id':'0008','caption':STANDARD,'机标caption':OLD,'note':'原样保留'},
              {'id':'0009','caption':'','机标caption':OLD,'note':'不跨列补充'}]
        job,data=self.create(rows,task_mapping=MAPPING)
        self.assertEqual(job['status'],'done');self.assertEqual(job['counts'],{'converted':1,'already_standard':1,'skipped':1})
        converted=[json.loads(line) for line in Path(job['outputs']['corrected']).read_text().splitlines()]
        self.assertNotEqual(converted[0]['caption'],OLD);self.assertEqual(converted[1:],rows[1:])
        self.assertEqual(converted[0]['note'],rows[0]['note']);self.assertEqual(converted[0]['机标caption'],rows[0]['机标caption'])
        self.assertEqual(Path(job['outputs']['fixed_ids']).read_text(),'0007\n')
        self.assertEqual((self.manager.root/job['id']/'input.jsonl').read_bytes(),data)
        with Path(job['outputs']['report']).open(encoding='utf-8-sig',newline='') as stream:report=list(csv.DictReader(stream))
        self.assertEqual(report[0]['题目 ID'],'0007');self.assertIn('旧字段转换',report[0]['修复标签'])

    def test_report_counts_real_header_changes_and_keeps_valid_visual_header(self):
        spaced=STANDARD.replace('## 任务类型','##任务类型')
        rows=[{'id':'correct','caption':STANDARD},{'id':'spaces','caption':spaced}]
        job,_=self.create(rows)
        self.assertEqual(job['counts'],{'already_standard':1,'converted':1})
        self.assertEqual(Path(job['outputs']['fixed_ids']).read_text(),'spaces\n')
        output=[json.loads(line) for line in Path(job['outputs']['corrected']).read_text().splitlines()]
        self.assertEqual([r['caption'] for r in output],[STANDARD,STANDARD])
        self.assertEqual(self.client.calls,[])

    def test_missing_headers_keep_source_id_marker_after_successful_conversion(self):
        incomplete=STANDARD.replace('## similarity_label\nYes\n\n','')
        job,_=self.create([{'id':'0001','caption':incomplete},{'id':'0002','caption':STANDARD}])
        self.assertEqual(job['missing_field_count'],1)
        self.assertEqual(Path(job['outputs']['missing_field_ids']).read_text(),'0001\n')
        with Path(job['outputs']['missing_field_report']).open(encoding='utf-8-sig',newline='') as stream:detail=list(csv.DictReader(stream))
        self.assertEqual(len(detail),1);self.assertEqual(detail[0]['题目 ID'],'0001')
        self.assertEqual(detail[0]['原始缺少字段'],'similarity_label');self.assertEqual(detail[0]['整理后仍缺少字段'],'')
        with Path(job['outputs']['report']).open(encoding='utf-8-sig',newline='') as stream:report=list(csv.DictReader(stream))
        self.assertIn('原始缺少字段',report[0]['修复标签']);self.assertEqual(report[0]['原始缺少字段'],'similarity_label')
        self.assertEqual(report[1]['原始缺少字段'],'');self.assertEqual(self.client.calls,[])

    def test_missing_headers_include_deferred_failed_and_missing_id_rows_without_fake_ids(self):
        incomplete=STANDARD.replace('## 画面元素\n\n','')
        job,_=self.create([{'id':'0001','caption':OLD},{'id':'','caption':incomplete},{'id':'0003','caption':STANDARD}])
        self.assertEqual(job['missing_field_count'],2)
        self.assertEqual(Path(job['outputs']['missing_field_ids']).read_text(),'0001\n')
        with Path(job['outputs']['missing_field_report']).open(encoding='utf-8-sig',newline='') as stream:detail=list(csv.DictReader(stream))
        self.assertEqual(detail[1]['题目 ID'],'');self.assertEqual(detail[1]['原始缺少字段'],'画面元素')
        self.assertEqual(detail[1]['整理后仍缺少字段'],'画面元素')
        self.assertNotIn('任务类型',detail[0]['原始缺少字段']);self.assertEqual(self.client.calls,[])

    def test_missing_header_detection_recognizes_aliases_and_ignores_empty_values_and_non_captions(self):
        alias=STANDARD.replace('video_similarity_tag','similarity_tag').replace('## 主体','## subject')
        for text in [alias,STANDARD,'','Caption','普通文字','=1+2']:
            self.assertEqual(missing_fields(text),[])
        self.assertEqual(missing_fields(STANDARD.replace('## 主体\n<主体1>：青年。\n\n','')),['主体'])
        self.assertEqual(missing_fields('## 任务类型\n人物参考'),[f for f in FIELDS if f!='任务类型'])
        self.assertEqual(missing_fields(render({})),[])

    def test_failed_model_conversion_still_lists_original_missing_fields(self):
        class Failing(PlanClient):
            async def create(self,**kw):raise TimeoutError('offline')
        self.manager.client_factory=Failing
        job,_=self.create([{'id':'0007','caption':OLD}],task_mapping=MAPPING)
        self.assertEqual(job['counts'],{'failed':1});self.assertEqual(job['missing_field_count'],1)
        self.assertEqual(Path(job['outputs']['missing_field_ids']).read_text(),'0007\n')
        with Path(job['outputs']['missing_field_report']).open(encoding='utf-8-sig',newline='') as stream:report=list(csv.DictReader(stream))
        self.assertIn('场景',report[0]['原始缺少字段']);self.assertEqual(report[0]['原始缺少字段'],report[0]['整理后仍缺少字段'])

    def test_history_supplements_missing_reports_without_model_or_overwriting_data(self):
        incomplete=STANDARD.replace('## similarity_label\nYes\n\n','')
        job,_=self.create([{'id':'0001','caption':incomplete}])
        checkpoint=self.manager.root/job['id']/'checkpoint.json'
        old_cache=checkpoint.read_bytes();corrected=Path(job['outputs']['corrected']).read_bytes()
        for key in ('missing_field_ids','missing_field_report'):Path(job['outputs'].pop(key)).unlink()
        job.pop('missing_field_count',None);self.manager.save()
        restored=FieldJobs(self.manager.root,self.manager.output_dir,lambda:self.client);self.addCleanup(restored.close)
        backfilled=restored.snapshots()[0]
        self.assertEqual(backfilled['missing_field_count'],1)
        self.assertEqual(Path(backfilled['outputs']['missing_field_ids']).read_text(),'0001\n')
        self.assertEqual(checkpoint.read_bytes(),old_cache)
        self.assertEqual(Path(backfilled['outputs']['corrected']).read_bytes(),corrected);self.assertEqual(self.client.calls,[])

    def test_old_rule_history_is_marked_and_retry_invalidates_its_checkpoint(self):
        job,_=self.create([{'id':'correct','caption':STANDARD}])
        checkpoint=self.manager.root/job['id']/'checkpoint.json'
        saved=json.loads(checkpoint.read_text());saved['fingerprint']='old-rules'
        saved['results']['1'].update(status='converted',caption=STANDARD.replace('video_similarity_tag','similarity_tag'))
        checkpoint.write_text(json.dumps(saved));job.pop('schema_version',None)
        self.assertTrue(self.manager.snapshots()[0]['outdated_schema'])
        self.manager.retry(job['id']);self.manager.executor.submit(lambda:None).result(timeout=10)
        self.assertEqual(job['schema_version'],VERSION);self.assertFalse(self.manager.snapshots()[0]['outdated_schema'])
        self.assertEqual(job['counts'],{'already_standard':1});self.assertEqual(self.client.calls,[])
        output=json.loads(Path(job['outputs']['corrected']).read_text())
        self.assertEqual(output['caption'],STANDARD)

    def test_missing_and_duplicate_ids_are_preserved_and_reported_before_api(self):
        job,_=self.create([{'id':'','caption':OLD},{'id':'same','caption':OLD},{'id':'same','caption':OLD}],task_mapping=MAPPING)
        self.assertEqual(job['counts']['needs_confirmation'],3);self.assertEqual(self.client.calls,[])
        self.assertEqual(Path(job['outputs']['review_ids']).read_text(),'same\n')

    def test_reference_ids_are_reported_without_mapping_api_or_data_change(self):
        rows=[{'id':'0007','caption':OLD},{'id':'0008','caption':STANDARD},{'id':'','caption':OLD}]
        job,_=self.create(rows)
        self.assertEqual(job['legacy_reference_count'],2);self.assertEqual(self.client.calls,[])
        self.assertEqual(Path(job['outputs']['legacy_reference_ids']).read_text(),'0007\n')
        self.assertEqual([json.loads(line) for line in Path(job['outputs']['corrected']).read_text().splitlines()],rows)
        with Path(job['outputs']['report']).open(encoding='utf-8-sig',newline='') as stream:report=list(csv.DictReader(stream))
        self.assertIn('旧任务类型为参考',report[0]['修复标签']);self.assertEqual(report[2]['题目 ID'],'')
        self.assertIn('缺少真实题目ID',report[2]['说明'])

    def test_explicitly_mapped_reference_record_still_appears_in_original_value_id_list(self):
        job,_=self.create([{'id':'0007','caption':OLD}],task_mapping=MAPPING)
        self.assertEqual(job['counts']['converted'],1)
        self.assertEqual(Path(job['outputs']['legacy_reference_ids']).read_text(),'0007\n')
        self.assertEqual(job['legacy_reference_count'],1)

    def test_old_only_csv_is_detected_by_column_name_without_human_caption_filter(self):
        source=self.root/'legacy.csv'
        with source.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream);writer.writerow(['id','caption','备注']);writer.writerow(['00001',OLD,STANDARD])
        self.assertEqual(detect_column(source),'caption')
        job=self.manager.create(source.name,io.BytesIO(source.read_bytes()),source.stat().st_size,task_mapping=MAPPING)
        self.manager.executor.submit(lambda:None).result(timeout=10);job=self.manager.jobs[job['id']]
        with Path(job['outputs']['corrected']).open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.reader(stream))
        self.assertEqual(rows[1][0],'00001');self.assertEqual(rows[1][2],STANDARD);self.assertIn('## 任务类型',rows[1][1])

    def test_real_question_id_header_is_detected_without_using_row_numbers(self):
        source=self.root/'questions.csv'
        with source.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream);writer.writerow(['题目 ID','caption']);writer.writerow(['00001',OLD])
        self.assertEqual(select_id_column(source),'题目 ID')

    def test_checkpoint_restart_keeps_success_and_retries_failed_request(self):
        class Failing(PlanClient):
            async def create(self,**kw):raise TimeoutError('offline')
        self.manager.client_factory=Failing
        job,_=self.create([{'id':'one','caption':OLD},{'id':'two','caption':STANDARD}],task_mapping=MAPPING)
        self.assertEqual(job['counts']['failed'],1)
        self.manager.client_factory=lambda:self.client
        self.manager.retry(job['id']);self.manager.executor.submit(lambda:None).result(timeout=10)
        self.assertEqual(job['status'],'done');self.assertEqual(len(self.client.calls),1)
        restored=FieldJobs(self.manager.root,self.manager.output_dir,lambda:self.client);self.addCleanup(restored.close)
        self.assertEqual(restored.snapshots()[0]['outputs'],job['outputs'])

    def test_task_mapping_can_be_added_on_retry(self):
        job,_=self.create([{'id':'one','caption':OLD}]);self.assertEqual(job['status'],'partial')
        self.manager.retry(job['id'],task_mapping=MAPPING);self.manager.executor.submit(lambda:None).result(timeout=10)
        self.assertEqual(job['status'],'done');self.assertEqual(job['counts']['converted'],1)

    def test_xlsx_export_changes_only_target_sheet_xml_and_keeps_formula_style_and_reference_sheet(self):
        import openpyxl
        source=self.root/'input.xlsx';book=openpyxl.Workbook();sheet=book.active;sheet.title='数据'
        sheet.append(['id','caption','公式']);sheet.append(['0007',OLD,'=1+2']);sheet['A2'].number_format='@'
        book.create_sheet('说明').append(['必须保留']);book.save(source);book.close()
        result=asyncio.run(normalize(OLD,client=PlanClient(),task_mapping=MAPPING));target=self.root/'fixed.xlsx'
        write_copy(source,target,'caption',{'1':result},'数据')
        with ZipFile(source) as old,ZipFile(target) as new:
            self.assertEqual(old.namelist(),new.namelist())
            self.assertEqual([n for n in old.namelist() if old.read(n)!=new.read(n)],['xl/worksheets/sheet1.xml'])
        book=openpyxl.load_workbook(target);self.assertEqual(book['数据']['A2'].value,'0007');self.assertEqual(book['数据']['C2'].value,'=1+2')
        self.assertEqual(book['说明']['A1'].value,'必须保留');book.close()

    def test_invalid_mapping_and_duplicate_json_keys_fail_without_overwriting_input(self):
        with self.assertRaises(ValueError):self.manager.create('x.jsonl',io.BytesIO(b'{}'),2,task_mapping={'old':7})
        data=b'{"id":"x","caption":"one","caption":"two"}\n'
        job=self.manager.create('x.jsonl',io.BytesIO(data),len(data));self.manager.executor.submit(lambda:None).result(timeout=10)
        saved=self.manager.jobs[job['id']];self.assertEqual(saved['status'],'failed');self.assertEqual(saved['outputs'],{})
        self.assertEqual((self.manager.root/job['id']/'input.jsonl').read_bytes(),data)


if __name__=='__main__':unittest.main()
