"""A colleague without the author's optional Node runtime can export XLSX."""
import asyncio
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
import captionqa.caption_qa_agent as qa
import captionqa.caption_qa_web as web
from captionqa.caption_qa_review import make_review
from captionqa.caption_qa_review_report import export_review, import_review
from captionqa.caption_qa_repair_table import export_repair_table
from tests.test_caption_qa_upgrades import GOOD


class PortableExportsTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        self.source=self.root/'input.jsonl';self.cache=self.root/'cache.jsonl'

    def prepare(self):
        raw=GOOD.replace('<主体1>向前跑。','男子向前跑。')
        self.source.write_text(json.dumps({'id':'00001','caption':raw},ensure_ascii=False)+'\n')
        asyncio.run(qa.aprocess_batch(self.source,self.cache,qa.CaptionQAAgent(dry_run=True),column='caption'))
        return make_review(self.cache,self.source,self.source,caption_column='caption',id_column='id')

    def test_review_export_is_styled_rich_literal_and_can_be_imported_without_node(self):
        doc=self.prepare()
        with patch('captionqa.caption_qa_review_report._artifact_runtime',side_effect=ValueError('not installed')):
            result=export_review(doc,self.root/'ids.txt')
        path=Path(result['repair_report_path']);book=load_workbook(path)
        try:
            sheet=book.active;self.assertEqual(sheet['B2'].value,'00001');self.assertEqual(sheet.freeze_panes,'A2')
            self.assertEqual(sheet.max_column,8);self.assertTrue(sheet.tables)
            self.assertTrue(sheet.data_validations.dataValidation);self.assertEqual(sheet['G2'].fill.fgColor.rgb[-6:],'FFF0D6')
            self.assertTrue(sheet['E2'].alignment.wrap_text)
        finally:book.close()
        with ZipFile(path) as z:self.assertIn(b'FFCC0000',z.read('xl/worksheets/sheet1.xml'))
        imported=import_review(doc,path,revision=doc['revision']);self.assertEqual(imported['source_sha256'],doc['source_sha256'])

    def test_repair_export_keeps_leading_zeros_and_formula_like_ids_as_text(self):
        raw=GOOD.replace('身穿白袍。','身穿白袍')
        self.source.write_text(''.join(json.dumps({'id':v,'caption':raw},ensure_ascii=False)+'\n' for v in ('00001','=1+1')))
        asyncio.run(qa.aprocess_batch(self.source,self.cache,qa.CaptionQAAgent(dry_run=True),column='caption'))
        with patch('captionqa.caption_qa_repair_table._artifact_runtime',side_effect=ValueError('not installed')):
            result=export_repair_table(self.cache,self.root/'ids.txt',id_column='id',dry_run=True,complete=True)
        book=load_workbook(result['repair_report_path'])
        try:
            sheet=book.active;self.assertEqual(sheet.max_column,2);self.assertEqual(sheet['A2'].value,'00001')
            self.assertEqual(sheet['A3'].value,'=1+1');self.assertEqual(sheet['A3'].data_type,'s')
            self.assertEqual(sheet.freeze_panes,'A2');self.assertTrue(sheet.tables['CaptionRepair'].autoFilter)
        finally:book.close()

    def test_web_dry_run_produces_corrected_copy_and_review_with_no_runtime_or_key(self):
        app=web.App(self.root/'out',self.root/'work');self.addCleanup(app.close)
        raw=GOOD.replace('<主体1>向前跑。','男子向前跑')
        payload=(json.dumps({'id':'00001','caption':raw,'备注':'保留'},ensure_ascii=False)+'\n').encode()
        with patch('captionqa.caption_qa_review_report._artifact_runtime',side_effect=ValueError('not installed')),patch('captionqa.caption_qa_repair_table._artifact_runtime',side_effect=ValueError('not installed')),patch.object(web,'read_saved_key',side_effect=AssertionError('no key read')):
            item=app.create_job('data.jsonl',io.BytesIO(payload),len(payload),dry_run=True)
            app.executor.submit(lambda:None).result(timeout=10)
        job=app.jobs[item['id']];self.assertEqual(job.status,'done',job.error)
        result=job.snapshot()['review_output'];output=json.loads(Path(result['corrected_path']).read_text())
        self.assertEqual(output['id'],'00001');self.assertEqual(output['备注'],'保留')
        self.assertIn('男子向前跑。',output['caption']);self.assertTrue(list(app.output_dir.glob('*人工审核*.xlsx')))

    def test_default_downloads_is_the_current_users_home(self):
        self.assertEqual(web.DEFAULT_OUTPUT_DIR,Path.home()/'Downloads')


if __name__=='__main__':unittest.main()
