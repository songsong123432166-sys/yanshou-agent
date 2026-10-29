"""Local UI backend acceptance; no external API or real Keychain access."""
from __future__ import annotations

import csv
import io
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from types import SimpleNamespace
from urllib.parse import quote
from unittest.mock import patch

import captionqa.caption_qa_agent as qa
import captionqa.caption_qa_web as web
from tests.test_caption_qa_upgrades import FakeAsyncClient, GOOD, semantic_reply


def csv_bytes(records):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["id", "人标 Caption"])
    writer.writerows(records)
    return stream.getvalue().encode("utf-8-sig")


class Socket:
    def __init__(self, request):
        self.input = io.BytesIO(request)
        self.output = bytearray()

    def makefile(self, *args, **kwargs):
        return self.input

    def sendall(self, value):
        self.output.extend(value)


class WebTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        root = Path(self.directory.name)
        self.app = web.App(root / "Downloads", root / "work")

    def tearDown(self):
        self.app.close()
        self.directory.cleanup()

    def finish(self):
        # Web completion now includes the separate Excel authoring process.
        self.app.executor.submit(lambda: None).result(timeout=30)

    def request(self, method, path, *, body=b"", headers=None, token=True):
        values = {"Host": "127.0.0.1:8765", "Content-Length": str(len(body))}
        if token:
            values["X-App-Token"] = self.app.token
        values.update(headers or {})
        request = f"{method} {path} HTTP/1.0\r\n" + "".join(f"{k}: {v}\r\n" for k, v in values.items()) + "\r\n"
        sock = Socket(request.encode("ascii") + body)
        web.Handler(sock, ("127.0.0.1", 12345), SimpleNamespace(app=self.app, server_address=("127.0.0.1", 8765)))
        head, data = bytes(sock.output).split(b"\r\n\r\n", 1)
        return int(head.split()[1]), data

    def test_upload_process_and_export_preserve_original_rows(self):
        data = csv_bytes([("001", GOOD), ("002", GOOD.replace("身穿白袍。", "身穿白袍")), ("003", "普通描述")])
        item = self.app.create_job("业务样例.csv", io.BytesIO(data), len(data), review_first=False, dry_run=True)
        self.finish()
        job = self.app.jobs[item["id"]]
        result = job.snapshot()
        self.assertEqual(result["status"], "done")
        self.assertEqual((result["processed"], result["total"]), (3, 3))
        self.assertEqual(result["counts"], {"PASS": 1, "AUTO_FIXED": 1, "SKIPPED_NOT_HUMAN_CAPTION": 1})
        self.assertTrue(result["output_ready"])
        self.assertEqual(job.output.parent, self.app.output_dir)
        self.assertEqual(job.source.read_bytes(), data)
        self.assertEqual(job.output.suffix, ".txt")
        self.assertEqual(job.output.read_text().splitlines(), ["002"])
        self.assertFalse(list(self.app.output_dir.glob("*.csv")))
        self.assertEqual(result["id_summary"]["problem_ids"], 1)
        book = load_workbook(result["repair_report_path"])
        try:
            self.assertEqual([cell.value for cell in book.active[1]], ["题目 ID", "错误项"])
            self.assertEqual(book.active.max_column, 2)
            self.assertEqual(book.active["A2"].value, "002")
            self.assertIn("缺少句末标点", book.active["B2"].value)
        finally:
            book.close()

    def test_repair_export_failure_keeps_qa_done_and_can_retry_without_key_or_model(self):
        data = csv_bytes([("001", GOOD.replace("身穿白袍。", "身穿白袍"))])
        with patch.object(web, "export_repair_table", side_effect=ValueError("导出组件暂不可用")):
            item = self.app.create_job("repair.csv", io.BytesIO(data), len(data), review_first=False, dry_run=True)
            self.finish()
        job = self.app.jobs[item["id"]]
        self.assertEqual(job.status, "done")
        self.assertIn("暂不可用", job.snapshot()["repair_report_error"])
        # A stale artifact from a prior export must not be presented as current.
        web.repair_report_path(job.output).write_bytes(b"previous export")
        self.assertEqual(job.snapshot()["repair_report_path"], "")
        original = job.cache.read_bytes(), job.source.read_bytes(), job.output.read_bytes()
        with patch.object(web, "read_saved_key", side_effect=AssertionError("no API")):
            status, result = self.request("POST", f"/api/jobs/{job.id}/repair-report", token=False)
            self.assertEqual(status, 403)
            status, result = self.request("POST", f"/api/jobs/{job.id}/repair-report")
        self.assertEqual(status, 200, result)
        self.assertTrue(Path(json.loads(result)["repair_report_path"]).is_file())
        self.assertNotIn("repair_report_error", job.id_summary)
        self.assertEqual((job.cache.read_bytes(), job.source.read_bytes(), job.output.read_bytes()), original)
        job.outdated_policy = True
        status, result = self.request("POST", f"/api/jobs/{job.id}/repair-report")
        self.assertEqual(status, 400)

    def test_same_name_does_not_overwrite_existing_result(self):
        data = csv_bytes([("001", GOOD)])
        first = self.app.create_job("same.csv", io.BytesIO(data), len(data), review_first=False, dry_run=True)
        second = self.app.create_job("same.csv", io.BytesIO(data), len(data), review_first=False, dry_run=True)
        self.finish()
        self.assertNotEqual(first["output_path"], second["output_path"])
        self.assertTrue(Path(first["output_path"]).is_file())
        self.assertTrue(Path(second["output_path"]).is_file())

    def test_partial_semantic_failure_retries_only_unfinished_row(self):
        data = csv_bytes([("001", GOOD), ("002", GOOD)])
        real_agent = qa.CaptionQAAgent
        clients = [FakeAsyncClient(semantic_reply(), ValueError("offline invalid response")), FakeAsyncClient(semantic_reply())]

        def factory(**kwargs):
            return real_agent(async_client=clients.pop(0), model_name=web.MODEL, base_url=web.BASE_URL,
                              max_retries=0, retry_delay=0)

        with patch.object(web, "read_saved_key", return_value="tp-offline"), patch.object(web.qa, "CaptionQAAgent", side_effect=factory):
            item = self.app.create_job("retry.csv", io.BytesIO(data), len(data), review_first=False)
            self.finish()
            job = self.app.jobs[item["id"]]
            self.assertEqual(job.snapshot()["status"], "partial")
            self.assertEqual(job.snapshot()["pending_retry"], 1)
            self.app.retry(job.id)
            self.finish()
        result = job.snapshot()
        self.assertEqual(result["status"], "done")
        self.assertEqual(result["pending_retry"], 0)
        self.assertEqual(result["counts"], {"PASS": 2})
        self.assertEqual(len(job.cache.read_text().splitlines()), 3)

    def test_dry_run_never_reads_keychain(self):
        data = csv_bytes([("001", GOOD)])
        with patch.object(web, "read_saved_key", side_effect=AssertionError("must not read")):
            item = self.app.create_job("static.csv", io.BytesIO(data), len(data), review_first=False, dry_run=True)
            self.finish()
        self.assertEqual(self.app.jobs[item["id"]].status, "done")

    def test_reveal_person_report_targets_separate_generated_file(self):
        data = (json.dumps({"id": "001", "caption": GOOD, "标注员": "甲"}) + "\n").encode()
        item = self.app.create_job("people.jsonl", io.BytesIO(data), len(data), review_first=False, dry_run=True)
        self.finish()
        snapshot = self.app.jobs[item["id"]].snapshot()
        self.assertTrue(snapshot["people_report_path"])
        with patch.object(web.sys, "platform", "darwin"), patch.object(web.subprocess, "run") as opened:
            status, result = self.request("POST", f"/api/jobs/{item['id']}/reveal?file=people")
        self.assertEqual(status, 200, result)
        self.assertEqual(opened.call_args.args[0], ["/usr/bin/open", "-R", snapshot["people_report_path"]])

    def test_default_parallelism_and_retry_endpoint_change_without_repeating_completed_rows(self):
        data = csv_bytes([("001", GOOD), ("002", GOOD)])
        with patch.object(web, "read_saved_key", side_effect=AssertionError("no API")):
            item = self.app.create_job("parallel.csv", io.BytesIO(data), len(data), review_first=False, dry_run=True)
            self.finish()
            job = self.app.jobs[item["id"]]
            self.assertEqual(job.concurrency, 8)
            original = job.cache.read_bytes()
            job.status = "failed"
            status, result = self.request("POST", f"/api/jobs/{job.id}/retry", body=b'{"concurrency":5}')
            self.assertEqual(status, 202, result)
            self.finish()
            self.assertEqual(job.concurrency, 5)
            self.assertEqual(job.status, "done", job.error)
            self.assertEqual(job.cache.read_bytes(), original)
        for value in (0, 11, True, 1.5, "8"):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "1–10"):
                self.app.retry(job.id, concurrency=value)
            self.assertEqual(job.concurrency, 5)

    def test_invalid_paths_extensions_and_upload_limits_are_rejected(self):
        for name, length in (("../escape.csv", 1), ("a\\b.csv", 1), ("x.csv\n", 1),
                             ("file.exe", 1), ("empty.csv", 0), ("large.csv", web.MAX_UPLOAD_BYTES + 1)):
            with self.subTest(name=name, length=length):
                with self.assertRaises(ValueError):
                    self.app.create_job(name, io.BytesIO(b"x"), length, review_first=False)
        self.assertFalse(self.app.jobs)

    def test_incomplete_upload_never_starts_job(self):
        with self.assertRaisesRegex(ValueError, "上传未完成"):
            self.app.create_job("broken.csv", io.BytesIO(b"x"), 100, review_first=False)
        self.assertFalse(self.app.jobs)

    def test_auth_host_and_origin_are_checked_before_upload(self):
        for headers, token in (({}, False), ({"Host": "other.example:8765"}, True),
                               ({"Origin": "https://other.example"}, True)):
            with self.subTest(headers=headers, token=token):
                status, _ = self.request("POST", "/api/jobs?name=test.csv", body=b"x", headers=headers, token=token)
                self.assertEqual(status, 403)
        self.assertFalse(self.app.jobs)

    def test_page_has_drop_target_and_same_session_token(self):
        status, page = self.request("GET", "/", token=False)
        self.assertEqual(status, 200)
        self.assertIn(b'id="dropzone"', page)
        self.assertIn(self.app.token.encode(), page)
        self.assertNotIn(b"__APP_TOKEN__", page)

    def test_config_does_not_send_key_to_browser(self):
        with patch.object(web, "read_saved_key", return_value="tp-DO-NOT-DISPLAY"):
            status, data = self.request("GET", "/api/config")
        self.assertEqual(status, 200)
        self.assertTrue(json.loads(data)["key_ready"])
        self.assertNotIn(b"tp-DO-NOT-DISPLAY", data)

    def test_http_upload_with_chinese_column_and_filename(self):
        data = csv_bytes([("001", GOOD)])
        options = quote(json.dumps({"dry_run": True, "column": "人标 Caption"}, ensure_ascii=False))
        status, response = self.request("POST", "/api/jobs?name=" + quote("中文样例.csv"), body=data,
                                        headers={"X-Job-Options": options})
        self.assertEqual(status, 202, response)
        self.finish()
        status, response = self.request("GET", "/api/jobs")
        self.assertEqual(status, 200)
        job = json.loads(response)["jobs"][0]
        self.assertEqual(job["status"], "done")
        self.assertEqual(job["name"], "中文样例.csv")


if __name__ == "__main__":
    unittest.main()
