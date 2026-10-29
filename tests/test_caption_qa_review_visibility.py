"""Automatic provenance is excluded from human lists and round-trip tables."""
import asyncio
import copy
import csv
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
import captionqa.caption_qa_agent as qa
from captionqa.caption_qa_review import SafeAutoFixAgent, make_review, summary, human_review_issues
from captionqa.caption_qa_review_report import HEADERS, export_review, import_review
from captionqa.caption_qa_review_io import export_approved
from tests.test_caption_qa_upgrades import GOOD


class ReviewVisibilityTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.source=self.root/'input.jsonl';self.cache=self.root/'cache.jsonl'

    def build(self,only_auto=False):
        raw=GOOD.replace('<主体1>向前跑。','<主体1>向前跑' if only_auto else '男子向前跑')
        self.source.write_text(json.dumps({'id':'0007','caption':raw,'note':'原样保留'})+'\n')
        asyncio.run(qa.aprocess_batch(self.source,self.cache,SafeAutoFixAgent(dry_run=True),column='caption'))
        return make_review(self.cache,self.source,self.source,caption_column='caption',id_column='id')

    def test_human_summary_omits_auto_but_keeps_execution_and_audit_data(self):
        document=self.build();unchanged=copy.deepcopy(document)
        report=summary(document)
        self.assertEqual(report['issues'],2)
        self.assertEqual(report['human_issues'],1)
        self.assertEqual(report['pending'],1)
        self.assertEqual(len(human_review_issues(document)),1)
        self.assertEqual(human_review_issues(document)[0]['kind'],'missing_subjects')
        result=export_approved(document,self.source,self.source,self.cache,self.root,'data')
        self.assertIn('男子向前跑。',json.loads(Path(result['corrected_path']).read_text())['caption'])
        audit=json.loads(Path(result['recheck_path']).read_text())
        self.assertTrue(any(i['decision']=='auto' for i in audit['decisions']))
        self.assertEqual(document,unchanged)

    def test_xlsx_and_csv_only_export_human_rows(self):
        document=self.build();output=export_review(document,self.root/'ids.txt')
        with Path(output['review_csv_path']).open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.reader(stream))
        automatic={i['issue_id'] for i in document['issues'] if i['decision']=='auto'}
        self.assertEqual(output['repair_report_rows'],1)
        self.assertEqual(len(rows),2)
        self.assertFalse(automatic & {row[0] for row in rows[1:]})
        book=load_workbook(output['repair_report_path']);self.addCleanup(book.close)
        self.assertEqual(book.active.max_row,2)
        self.assertEqual(book.active['A2'].value,human_review_issues(document)[0]['issue_id'])

    def test_roundtrip_human_draft_can_omit_hidden_auto_in_same_field(self):
        document=self.build();issue=human_review_issues(document)[0]
        final='0-4秒，第0秒，<主体1>向前跑。'
        path=self.root/'review.csv'
        with path.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream);writer.writerow(HEADERS)
            writer.writerow([issue['issue_id'],issue['sample_id'],issue['field'],issue['label'],'\n'.join(issue['excerpts']),issue['subject'],'已人工修改',final])
        imported=import_review(document,path,revision=0)
        self.assertEqual(len(human_review_issues(imported)),1)
        self.assertEqual(summary(imported)['pending'],0)
        result=export_approved(imported,self.source,self.source,self.cache,self.root,'data')
        self.assertIn(final,json.loads(Path(result['corrected_path']).read_text())['caption'])
        again=export_review(imported,self.root/'next.txt')
        self.assertEqual(again['repair_report_rows'],1)

    def test_only_automatic_results_have_no_human_rows(self):
        document=self.build(only_auto=True)
        self.assertEqual(human_review_issues(document),[])
        self.assertEqual(summary(document)['human_issues'],0)
        output=export_review(document,self.root/'ids.txt')
        self.assertEqual(output['repair_report_rows'],0)
        with Path(output['review_csv_path']).open(encoding='utf-8-sig',newline='') as stream:
            self.assertEqual(list(csv.reader(stream)),[HEADERS])
        book=load_workbook(output['repair_report_path']);self.addCleanup(book.close)
        self.assertEqual(book.active.max_row,1)


if __name__=='__main__':unittest.main()
