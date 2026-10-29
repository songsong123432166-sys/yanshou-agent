"""The distributable skill must run isolated and alter only chosen Caption cells."""
import csv
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from tests.test_caption_qa_upgrades import GOOD


SKILL = Path(__file__).resolve().parents[1] / 'skills/caption-safe-fix'
SCRIPT = SKILL / 'scripts/fix_caption.py'
BAD = GOOD.replace('身穿白袍。', '身穿白袍').replace('<主体1>向前跑', '主体1向前跑')


class CaptionFixSkillTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'business.jsonl'
        self.output = self.root / 'out'

    def run_supported(self, *args, source=None):
        # No third-party site packages, project cwd, API keys or PYTHONPATH.
        source = source or self.source
        original = source.read_bytes()
        import os
        env = {k: v for k, v in os.environ.items() if k not in {'PYTHONPATH', 'OPENAI_API_KEY', 'OPENAI_BASE_URL'}}
        result = subprocess.run([sys.executable, '-S', str(SCRIPT), str(source), '--output-dir', str(self.output), *args],
                                cwd=self.root, env=env, capture_output=True, text=True)
        self.assertEqual(source.read_bytes(), original)
        return result

    def finish(self, result):
        self.assertEqual(result.returncode, 0, result.stderr)
        paths = json.loads(result.stdout)
        return paths, json.loads(Path(paths['report']).read_text())

    def test_jsonl_isolated_run_keeps_other_fields_ids_and_literal_content(self):
        literal = '说道{我是主体3，眼晴 3<5}，标牌写着【主体9 < 音频2 >】。'
        raw = BAD.replace('0-4秒，第0秒，主体1向前跑。', '0-4秒，第0秒，主体1向前跑。' + literal)
        record = {'id': '00001', 'caption': raw, '备注': BAD, '原始标注员': '原值'}
        self.source.write_text(json.dumps(record, ensure_ascii=False) + '\n')
        paths, summary = self.finish(self.run_supported('--column', 'caption', '--id-column', 'id'))
        actual = json.loads(Path(paths['output']).read_text())
        self.assertEqual(actual['id'], '00001')
        self.assertEqual(actual['备注'], BAD)
        self.assertEqual(actual['原始标注员'], '原值')
        self.assertIn(literal, actual['caption'])
        self.assertIn('身穿白袍。', actual['caption'])
        self.assertIn('<主体1>向前跑', actual['caption'])
        self.assertEqual(summary['counts']['changed'], 1)
        self.assertEqual(summary['counts']['review'], 0)
        self.assertNotIn('我是主体3', Path(paths['report']).read_text())

    def test_csv_secondary_header_missing_cells_and_note_column_are_preserved(self):
        self.source = self.source.with_suffix('.csv')
        with self.source.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerows([['样本ID', '人工标注', '', '备注'], ['', '是否运镜', 'Caption', ''],
                              ['00001', '是', BAD, BAD], ['00002']])
        paths, summary = self.finish(self.run_supported('--column', '3', '--id-column', '1'))
        with Path(paths['output']).open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[2][3], BAD)
        self.assertEqual(rows[2][0], '00001')
        self.assertEqual(rows[3], ['00002'])
        self.assertEqual(summary['counts']['total'], 4)
        self.assertEqual(summary['counts']['skipped'], 3)

    def test_ambiguous_caption_columns_abort_instead_of_falling_back(self):
        self.source = self.source.with_suffix('.csv')
        self.source.write_text('id,Caption,Caption\n1,a,b\n')
        result = self.run_supported()
        self.assertEqual(result.returncode, 2)
        self.assertFalse(self.output.exists())

    def make_xlsx(self):
        self.source = self.source.with_suffix('.xlsx')
        book = Workbook()
        sheet = book.active
        sheet.title = '数据'
        sheet.append(['ID', 'Caption', '备注'])
        sheet.append(['00001', BAD, BAD])
        sheet['B2'].font = Font(name='Arial', bold=True, color='FF0000')
        sheet['B3'] = '=B2'
        sheet['C3'] = '=1+2'
        book.create_sheet('说明')['A1'] = BAD
        book.save(self.source)
        book.close()
        # Unrelated archive parts should survive byte-for-byte, including parts
        # an ordinary high-level Excel writer might discard.
        with zipfile.ZipFile(self.source, 'a') as archive:
            archive.writestr('customXml/preserve-test.xml', '<original>keep me</original>')

    def test_xlsx_only_target_cell_changes_and_all_other_parts_survive(self):
        self.make_xlsx()
        paths, summary = self.finish(self.run_supported('--sheet', '数据', '--column', 'B', '--id-column', 'A'))
        output = Path(paths['output'])
        with zipfile.ZipFile(self.source) as before, zipfile.ZipFile(output) as after:
            self.assertEqual(before.namelist(), after.namelist())
            modified_parts = [name for name in before.namelist() if before.read(name) != after.read(name)]
            self.assertEqual(modified_parts, ['xl/worksheets/sheet1.xml'])
        book = load_workbook(output)
        try:
            sheet = book['数据']
            self.assertEqual(sheet['A2'].value, '00001')
            self.assertEqual(sheet['C2'].value, BAD)
            self.assertIn('身穿白袍。', sheet['B2'].value)
            self.assertEqual(sheet['B3'].value, '=B2')
            self.assertEqual(sheet['C3'].value, '=1+2')
            self.assertEqual(book['说明']['A1'].value, BAD)
            self.assertTrue(sheet['B2'].font.bold)
            self.assertEqual(sheet['B2'].font.color.rgb, '00FF0000')
        finally:
            book.close()
        self.assertEqual(summary['counts']['changed'], 1)
        self.assertTrue(any('公式' in row['reason'] for row in summary['skipped']))

    def test_xlsx_multiple_sheets_require_explicit_selection(self):
        self.make_xlsx()
        result = self.run_supported('--column', 'B')
        self.assertEqual(result.returncode, 2)
        self.assertIn('--sheet', result.stderr)
        self.assertFalse(self.output.exists())

    def test_xlsx_shared_strings_prefixed_cells_and_bad_dimensions_are_supported(self):
        from xml.sax.saxutils import escape
        import re
        self.make_xlsx()
        with zipfile.ZipFile(self.source) as archive:
            parts = {entry.filename: (entry, archive.read(entry.filename)) for entry in archive.infolist()}
        info, raw = parts['xl/worksheets/sheet1.xml']
        text = raw.decode().replace('<worksheet ', '<worksheet xmlns:m="http://schemas.openxmlformats.org/spreadsheetml/2006/main" ', 1)
        text = re.sub(r'<dimension[^>]*/>', '<dimension ref="A1"/>', text)
        text = re.sub(r'<c\b[^>]*r="B2"[^>]*>.*?</c>', '<m:c r="B2" s="1" t="s"><v>0</v></m:c>', text)
        parts['xl/worksheets/sheet1.xml'] = info, text.encode()
        shared = ('<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>' + escape(BAD) + '</t></si></sst>').encode()
        with zipfile.ZipFile(self.source, 'w') as archive:
            for info, content in parts.values():
                archive.writestr(info, content)
            archive.writestr('xl/sharedStrings.xml', shared)
        paths, summary = self.finish(self.run_supported('--sheet', '数据', '--column', 'B', '--id-column', 'A'))
        with zipfile.ZipFile(paths['output']) as archive:
            import xml.etree.ElementTree as ET
            root = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            cell = next(element for element in root.iter() if element.get('r') == 'B2')
            value = ''.join(element.text or '' for element in cell.iter() if element.tag.endswith('}t'))
            self.assertIn('身穿白袍。', value)
            self.assertIn('<主体1>向前跑', value)
            self.assertEqual(archive.read('xl/sharedStrings.xml'), shared)
        self.assertEqual(summary['counts']['changed'], 1)

    def test_manual_findings_are_not_fixed_or_labeled_as_complete(self):
        raw = BAD.replace('## 总结', '<主体2>：女孩。\n## 总结')
        self.source.write_text(json.dumps({'id': '00001', 'caption': raw}) + '\n')
        paths, report = self.finish(self.run_supported())
        revised = json.loads(Path(paths['output']).read_text())['caption']
        self.assertNotIn('<主体2>向', revised)
        self.assertEqual(report['counts']['review'], 1)
        with Path(paths['review']).open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[0], ['题目 ID', '错误项'])
        self.assertEqual(rows[1][0], '00001')
        self.assertIn('<主体2>', rows[1][1])
        self.assertNotIn('身穿白袍', rows[1][1])

    def test_metadata_and_disabled_items_remain_unchanged(self):
        metadata = '## 任务类型\n人物参考、动态参考\n## video_similarity_tag\n实体/IP｜动态\n'
        raw = metadata + GOOD.replace('<主体1>：身穿白袍。', '<主体1>：<主体1>：手手抬起，眼晴明亮。')
        self.source.write_text(json.dumps({'id': '001', 'caption': raw}) + '\n')
        paths, report = self.finish(self.run_supported())
        self.assertEqual(json.loads(Path(paths['output']).read_text())['caption'], raw)
        self.assertEqual(report['counts']['changed'], 0)
        self.assertEqual(report['counts']['review'], 1)

    def test_missing_id_stays_missing_and_is_not_replaced_by_row_number(self):
        raw = GOOD.replace('<主体1>向前跑。', '男子向前跑。')
        self.source.write_text(json.dumps({'caption': raw}) + '\n')
        paths, report = self.finish(self.run_supported())
        self.assertEqual(len(report['missing_review_ids']), 1)
        with Path(paths['review']).open(encoding='utf-8-sig', newline='') as stream:
            self.assertEqual(list(csv.reader(stream)), [['题目 ID', '错误项']])

    def test_duplicate_json_fields_abort_before_output(self):
        self.source.write_text('{"caption":"a","caption":"b"}\n')
        result = self.run_supported()
        self.assertEqual(result.returncode, 2)
        self.assertIn('重复字段', result.stderr)
        self.assertFalse(self.output.exists())

    def test_second_run_is_idempotent_and_never_overwrites_the_first_copy(self):
        self.source.write_text(json.dumps({'id': '001', 'caption': BAD}) + '\n')
        paths, _ = self.finish(self.run_supported())
        first = Path(paths['output'])
        original = first.read_bytes()
        next_paths, summary = self.finish(self.run_supported(source=first))
        self.assertEqual(summary['counts']['changed'], 0)
        self.assertNotEqual(first, Path(next_paths['output']))
        self.assertEqual(first.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
